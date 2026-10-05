#!/usr/bin/env python3
"""Build the five-seed evaluation update as HTML and an editable PPTX deck."""

from __future__ import annotations

import base64
import csv
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "cxr_real_experiment" / "evaluation_update_deck_20260712"

ARCHIVE_ROOT = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "results_mobilenetv3_noes50_clean_3seed/20260707_123320"
)
PERFORMANCE_DIR = ARCHIVE_ROOT / "evaluation_20260711_remove_vs_refine_5seed_ars"
QUALITY_DIR = ARCHIVE_ROOT / "evaluation_20260711_data_quality_5seed_ars"

PERFORMANCE_FIGURE = PERFORMANCE_DIR / "five_seed_method_curve.png"
DQS_FIGURE = QUALITY_DIR / "dqs_and_coverage_adjusted_dqs.png"

PLUGIN_ROOT = Path(
    os.environ.get(
        "DATA_ANALYTICS_PLUGIN_ROOT",
        "/homes/yz3522/.codex/plugins/cache/openai-curated-remote/"
        "data-analytics/0.2.8-13ceeea1f599",
    )
)
SLIDES_HELPER = (
    PLUGIN_ROOT
    / "skills/build-report/report-to-google-slides/scripts/report_to_google_slides.py"
)

DECK_TITLE = "Evaluation update: can we trust the label-cleaning results?"
DECK_SUBTITLE = (
    "Early removal helps, but repeated removal loses coverage; fixed LLM refinement "
    "has the stronger late quality-coverage point estimate."
)

FINAL_TITLES = [
    DECK_TITLE,
    "Five paired seeds test training robustness under a controlled protocol",
    "The two methods show different trajectories across five seeds",
    "The trend is promising, not confirmatory",
    "DQS measures estimated consistency on the currently covered entries",
    "Refinement improves estimated label health while preserving coverage",
    "Proper scoring rules expose late-removal degradation, but expert validation is still missing",
    "What the evidence supports, and what must be tested next",
    "Backup: exact reproducibility protocol",
    "Backup: full five-seed loop table",
    "Backup: rare-label and distribution sensitivity",
    "Backup: evidence modes and wording controls",
    "Backup: how Cleanlab obtains M_t in a binary mini example",
]

SPEAKER_NOTES = [
    (
        "Opening\n\n"
        "Last time, the main concern was not whether we could add another cleaning method, "
        "but whether the current results were actually reliable. So this week I focused on "
        "three questions: whether the trend repeats across seeds, whether DQS is inflated by "
        "deleting data, and whether we have evidence beyond confident learning itself.\n\n"
        "The main message is that early removal does help, but repeated removal progressively "
        "loses coverage. Fixed LLM refinement shows a stronger late-loop quality-coverage "
        "trade-off. However, I will treat this as promising rather than conclusive."
    ),
    (
        "I repeated the comparison with five paired training seeds. For each seed, I trained "
        "the no-clean baseline, five removal loops, and five fixed-refinement loops, giving 55 "
        "models in total. All models were evaluated on exactly the same 605 expert-labelled "
        "studies.\n\n"
        "The statistical repeat is the five paired seeds, not ten independent method-seed "
        "experiments. One important limitation is that these experiments test training-seed "
        "robustness, but not full-pipeline reproducibility. Removal reruns OOF detection for "
        "every seed, whereas refinement uses fixed cumulative tables originally generated "
        "with XRV seed 13."
    ),
    (
        "The trajectories are quite different. Simple removal gives the strongest improvement "
        "early, around Loop 1, but the gain declines as more samples are removed. Refinement "
        "starts more modestly, then improves later and reaches the highest Loop 5 point estimate "
        "while retaining all training samples.\n\n"
        "Remove Loop 1 improves AUROC by about 0.028. By Loop 5, that falls to about 0.008. "
        "Refine Loop 5 improves by about 0.034 and is positive in all five seeds. Loop 5 was "
        "identified after inspecting the trajectory, so it should still be considered exploratory."
    ),
    (
        "The five seeds make the direction more credible, but they do not make it confirmatory. "
        "Once both training-seed and study-level variation are included, all hierarchical "
        "confidence intervals still include zero.\n\n"
        "The apparent Refine-versus-Remove advantage at Loop 5 is 0.026 across all labels, but "
        "it falls to 0.0068 after excluding labels with very small minority support. With only "
        "five paired seeds, even five out of five positive differences cannot produce a "
        "two-sided exact sign-flip p-value below 0.05. So I would not claim universal superiority."
    ),
    (
        "This addresses the denominator question directly. Raw DQS is one minus the estimated "
        "number of label issues divided by the number of currently valid entries. Therefore, "
        "after removal, DQS is evaluated only on the surviving subset.\n\n"
        "To make coverage loss explicit, I also report coverage-adjusted DQS, which measures "
        "estimated healthy-entry mass relative to the original denominator. This adjusted "
        "score is our diagnostic guardrail, not an official Cleanlab metric.\n\n"
        "I also found that the previous 1.77 percent figure used all possible label slots as "
        "the denominator. The corresponding valid-entry issue rate is actually 9.63 percent."
    ),
    (
        "This changes the interpretation quite clearly. Remove Loop 5 has the highest raw DQS, "
        "0.9745, but it retains only about 83 percent of the original valid entries. After "
        "adjusting for coverage, its score falls to 0.8097.\n\n"
        "Refinement retains about 99.6 percent coverage, and both its raw and adjusted DQS are "
        "above baseline. Therefore, its improvement cannot be explained simply by shrinking "
        "the denominator. But this is still frozen-OOF model-consistency evidence. It is not "
        "direct evidence that the new labels are clinically correct."
    ),
    (
        "I then checked held-out performance using AUROC, average precision, Brier score, and "
        "NLL. Refinement has the best AUROC, average precision, and Brier score at Loop 5. Its "
        "NLL is close to baseline and much better than late removal, whose NLL degrades substantially.\n\n"
        "The initial OOF rankings are quite stable across seeds, but the exact removed sets "
        "have only moderate overlap. This suggests that the general notion of which examples "
        "are difficult is stable, while threshold-level selection is more variable. These "
        "checks support the quality-coverage explanation, but none of them directly verifies "
        "label correctness. That still requires expert adjudication."
    ),
    (
        "The strongest conclusion supported by the current evidence is that simple removal "
        "provides an early gain, but repeated removal trades away coverage and changes the "
        "training distribution. Fixed refinement provides a promising coverage-preserving "
        "alternative with a stronger late-loop point estimate.\n\n"
        "I would not yet claim that LLM refinement is universally superior, that all proposed "
        "relabels are clinically correct, or that five training seeds reproduce the full LLM "
        "pipeline. The next useful validation is not another internal metric. It is to regenerate "
        "refinement tables from independent detection seeds, lock the primary endpoint before "
        "confirmation, and perform an expert audit covering both flagged and randomly sampled "
        "unflagged entries.\n\n"
        "Closing line: The result is more reliable than before, but the correct claim is a "
        "promising quality-coverage trade-off, not confirmed universal superiority."
    ),
    (
        "Q&A use only.\n\n"
        "Use this slide if the exact implementation settings or statistical repeat are questioned. "
        "Emphasise that method and loop are paired within five seeds, the same test studies are "
        "used throughout, and fixed refinement tables do not constitute five independent LLM pipelines."
    ),
    (
        "Q&A use only.\n\n"
        "Use this slide to show that the conclusion comes from the full five-loop trajectory, "
        "not a selected best result. Early removal is stronger; refinement overtakes it only "
        "at Loops 4 and 5. Loop 5 remains exploratory until a primary endpoint is locked."
    ),
    (
        "Q&A use only.\n\n"
        "Use this slide if the Loop 5 advantage is challenged by rare-label support or "
        "distribution shift. Large effects for labels with one or two minority studies are "
        "unstable. The support-filtered direct advantage is only 0.0068, while removal also "
        "causes materially larger coverage and prevalence shifts."
    ),
    (
        "Q&A use only.\n\n"
        "Use this slide to keep evidence claims separate. Frozen OOF evidence compares actions, "
        "dynamic OOF evidence diagnoses the remaining data, held-out metrics measure downstream "
        "utility, and only expert adjudication can directly establish label correctness."
    ),
    (
        "Q&A use only.\n\n"
        "This example starts with eight valid binary entries. Cleanlab computes one self-confidence "
        "threshold for each observed class: t0 is the mean probability of class zero among entries "
        "labelled zero, and t1 is the mean probability of class one among entries labelled one. "
        "Entries D and H are confident off-diagonal disagreements, so M_t equals two and raw DQS is "
        "one minus two over eight, or 0.75.\n\n"
        "If D and H are deleted, raw DQS becomes 1.00 on the six survivors. Coverage is now 0.75, "
        "so adjusted DQS remains 0.75. The survivor set looks perfectly clean, but the estimated "
        "healthy-entry mass relative to the original denominator has not increased."
    ),
]

SOURCE_FOOTERS = [
    "Sources: five-seed ARS validation and data-quality reports; generated 12 July 2026.",
    "Source: evaluation_update_slides_20260712_Chinese.md; five paired seeds 7, 13, 42, 97, 123.",
    "Source: score_summary.csv and coverage_summary.csv; five-seed mean +/- SD.",
    "Source: paired_bootstrap_statistics.csv; hierarchical seed-and-study bootstrap, 1,000 iterations.",
    "Source: Cleanlab overall_label_health_score implementation and denominator audit in the data-quality report.",
    "Source: oof_quality_summary.csv; frozen initial OOF post-action evidence.",
    "Sources: full-precision primary AUROC; heldout_supplementary_metrics_summary.csv; OOF stability CSVs.",
    "Source: five-seed ARS validation; wording limited to evidence supported by the current design.",
    "Source: experiment scripts and saved evaluation metadata.",
    "Source: score_summary.csv; primary AUROC uses full-precision per-label summaries.",
    "Sources: per_label_deltas.csv, test_label_support.csv, and label_distribution_shift.csv.",
    "Sources: five-seed ARS reports and saved evaluation artifacts; modes must not be conflated.",
    "Source: constructed example following the local Cleanlab 2.9.0 num_label_issues implementation.",
]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def select_one(rows: list[dict[str, str]], **criteria: Any) -> dict[str, str]:
    matches = [
        row
        for row in rows
        if all(str(row.get(key)) == str(value) for key, value in criteria.items())
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one row for {criteria}, found {len(matches)}")
    return matches[0]


def checked_number(actual: float, expected: float, tolerance: float = 5e-5) -> float:
    if abs(actual - expected) > tolerance:
        raise ValueError(f"Evidence changed: expected {expected}, got {actual}")
    return actual


def source(
    source_id: str,
    label: str,
    path: str,
    description: str,
    *,
    sql: str | None = None,
    executed_at: str | None = None,
    tables_used: list[str] | None = None,
) -> dict[str, Any]:
    item: dict[str, Any] = {
        "id": source_id,
        "label": label,
        "path": path,
        "description": description,
    }
    if sql is not None:
        item["query"] = {
            "engine": "sqlite_snapshot",
            "sql": sql,
            "description": description,
            "executed_at": executed_at,
            "language": "sql",
            "tables_used": tables_used or [],
        }
    return item


def sql_literal(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return repr(value)
    return "'" + str(value).replace("'", "''") + "'"


def rows_to_snapshot_sql(rows: list[dict[str, Any]]) -> str:
    if not rows:
        raise ValueError("Snapshot SQL requires at least one row")
    fields = list(rows[0])
    selects = []
    for row_index, row in enumerate(rows):
        values = []
        for field in fields:
            value = sql_literal(row.get(field))
            values.append(f'{value} AS "{field}"' if row_index == 0 else value)
        selects.append("SELECT " + ", ".join(values))
    return "\nUNION ALL\n".join(selects)


def execute_snapshot_sql(sql: str) -> list[dict[str, Any]]:
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row
    try:
        return [dict(row) for row in connection.execute(sql).fetchall()]
    finally:
        connection.close()


def table_columns(fields: list[tuple[str, str]]) -> list[dict[str, str]]:
    return [{"field": field, "label": label, "type": "text"} for field, label in fields]


def build_evidence() -> dict[str, Any]:
    scores = read_csv(PERFORMANCE_DIR / "score_summary.csv")
    coverage = read_csv(PERFORMANCE_DIR / "coverage_summary.csv")
    bootstrap = read_csv(PERFORMANCE_DIR / "paired_bootstrap_statistics.csv")
    quality = read_csv(QUALITY_DIR / "oof_quality_summary.csv")
    supplementary = read_csv(QUALITY_DIR / "heldout_supplementary_metrics_summary.csv")

    baseline = select_one(scores, scope="all_labels", method="baseline", loop="0")
    baseline_mean = checked_number(float(baseline["mean_weighted_auroc"]), 0.7230864)

    performance_rows: list[dict[str, Any]] = []
    for loop in range(1, 6):
        performance_rows.append(
            {
                "loop": f"Loop {loop}",
                "method": "No-clean baseline",
                "mean_auroc": baseline_mean,
                "sd": float(baseline["sd_weighted_auroc"]),
            }
        )
        for method, label in (("remove", "Simple remove"), ("refine", "Fixed LLM refine")):
            row = select_one(scores, scope="all_labels", method=method, loop=str(loop))
            performance_rows.append(
                {
                    "loop": f"Loop {loop}",
                    "method": label,
                    "mean_auroc": float(row["mean_weighted_auroc"]),
                    "sd": float(row["sd_weighted_auroc"]),
                }
            )

    contrast_specs = [
        ("all_labels", "remove_vs_baseline", "3", "Remove L3 - baseline", "5/5 positive"),
        ("all_labels", "refine_vs_baseline", "5", "Refine L5 - baseline", "5/5 positive"),
        ("all_labels", "refine_vs_remove", "5", "Refine L5 - Remove L5", "5/5 positive"),
        (
            "support_ge_5",
            "refine_vs_remove",
            "5",
            "Direct contrast, support >= 5",
            "5/5 positive",
        ),
    ]
    uncertainty_rows: list[dict[str, Any]] = []
    for scope, comparison, loop, label, direction in contrast_specs:
        row = select_one(
            bootstrap,
            scope=scope,
            comparison=comparison,
            loop=loop,
            uncertainty="hierarchical_seed_and_study",
        )
        uncertainty_rows.append(
            {
                "contrast": label,
                "mean_delta": float(row["observed_mean_delta"]),
                "ci_low": float(row["ci_low_2p5"]),
                "ci_high": float(row["ci_high_97p5"]),
                "seed_direction": direction,
            }
        )

    frozen_rows = [
        row for row in quality if row["evidence_mode"] == "frozen_initial_oof_post_action"
    ]
    baseline_quality = select_one(frozen_rows, method="baseline", loop="0")
    checked_number(float(baseline_quality["dqs_flattened_mean"]), 0.9277273)
    dqs_chart_rows: list[dict[str, Any]] = []
    for loop in range(1, 6):
        for method, label in (("remove", "Remove"), ("refine", "Refine")):
            row = select_one(frozen_rows, method=method, loop=str(loop))
            dqs_chart_rows.extend(
                [
                    {
                        "loop": f"Loop {loop}",
                        "series": f"{label} raw DQS",
                        "score": float(row["dqs_flattened_mean"]),
                    },
                    {
                        "loop": f"Loop {loop}",
                        "series": f"{label} adjusted DQS",
                        "score": float(row["coverage_adjusted_dqs_mean"]),
                    },
                ]
            )

    summary_by_key = {
        (row["method"], int(row["loop"])): row
        for row in scores
        if row["scope"] == "all_labels"
    }
    coverage_by_loop = {int(row["loop"]): row for row in coverage}
    full_loop_rows = []
    for loop in range(1, 6):
        remove = summary_by_key[("remove", loop)]
        refine = summary_by_key[("refine", loop)]
        full_loop_rows.append(
            {
                "loop": str(loop),
                "remove_auroc": f"{float(remove['mean_weighted_auroc']):.4f}",
                "remove_delta": f"{float(remove['mean_delta_vs_baseline']):+.4f}",
                "remove_positive": f"{remove['positive_delta_seeds']}/5",
                "refine_auroc": f"{float(refine['mean_weighted_auroc']):.4f}",
                "refine_delta": f"{float(refine['mean_delta_vs_baseline']):+.4f}",
                "refine_positive": f"{refine['positive_delta_seeds']}/5",
                "direct": (
                    f"{float(refine['mean_weighted_auroc']) - float(remove['mean_weighted_auroc']):+.4f}"
                ),
            }
        )

    anchor_rows = []
    for method, loop in (("baseline", 0), ("remove", 5), ("refine", 5)):
        row = select_one(supplementary, scope="all_labels", method=method, loop=str(loop))
        anchor_rows.append(row)

    heldout = {(row["method"], int(row["loop"])): row for row in anchor_rows}
    primary_l5 = {
        "baseline": summary_by_key[("baseline", 0)],
        "remove": summary_by_key[("remove", 5)],
        "refine": summary_by_key[("refine", 5)],
    }
    metric_rows = [
        {
            "evidence": "Primary weighted AUROC",
            "baseline": f"{float(primary_l5['baseline']['mean_weighted_auroc']):.4f}",
            "remove": f"{float(primary_l5['remove']['mean_weighted_auroc']):.4f}",
            "refine": f"{float(primary_l5['refine']['mean_weighted_auroc']):.4f}",
            "interpretation": "Higher is better; primary endpoint",
        },
        {
            "evidence": "Macro average precision",
            "baseline": f"{float(heldout[('baseline', 0)]['macro_average_precision_mean']):.4f}",
            "remove": f"{float(heldout[('remove', 5)]['macro_average_precision_mean']):.4f}",
            "refine": f"{float(heldout[('refine', 5)]['macro_average_precision_mean']):.4f}",
            "interpretation": "Higher is better",
        },
        {
            "evidence": "Micro Brier",
            "baseline": f"{float(heldout[('baseline', 0)]['micro_brier_mean']):.4f}",
            "remove": f"{float(heldout[('remove', 5)]['micro_brier_mean']):.4f}",
            "refine": f"{float(heldout[('refine', 5)]['micro_brier_mean']):.4f}",
            "interpretation": "Lower is better",
        },
        {
            "evidence": "Micro NLL",
            "baseline": f"{float(heldout[('baseline', 0)]['micro_nll_mean']):.4f}",
            "remove": f"{float(heldout[('remove', 5)]['micro_nll_mean']):.4f}",
            "refine": f"{float(heldout[('refine', 5)]['micro_nll_mean']):.4f}",
            "interpretation": "Lower; late removal degrades sharply",
        },
        {
            "evidence": "OOF ranking Spearman rho",
            "baseline": "-",
            "remove": "0.887-0.903",
            "refine": "Fixed tables",
            "interpretation": "Rankings are stable across seeds",
        },
        {
            "evidence": "Loop 1 removed-set Jaccard",
            "baseline": "-",
            "remove": "0.470",
            "refine": "-",
            "interpretation": "Exact selected set only moderately stable",
        },
        {
            "evidence": "Selected by >= 3/5 seeds",
            "baseline": "-",
            "remove": "47.29% of union",
            "refine": "-",
            "interpretation": "Threshold membership varies",
        },
        {
            "evidence": "Expert label adjudication",
            "baseline": "Not run",
            "remove": "Not run",
            "refine": "Not run",
            "interpretation": "Label correctness remains unverified",
        },
    ]

    dqs_definition_rows = [
        {
            "term": "E_t",
            "definition": "Valid disease-label entries after the loop-t action",
            "meaning": "Current denominator",
        },
        {
            "term": "M_t",
            "definition": "Cleanlab num_label_issues(y, p_OOF) over E_t",
            "meaning": "Estimated model-consistency issues",
        },
        {
            "term": "Raw DQS",
            "definition": "1 - M_t / |E_t|",
            "meaning": "Health of currently covered entries",
        },
        {
            "term": "Coverage",
            "definition": "|E_t| / |E_0|",
            "meaning": "Fraction of original valid entries retained",
        },
        {
            "term": "Adjusted DQS",
            "definition": "DQS x Coverage = (|E_t| - M_t) / |E_0|",
            "meaning": "Healthy-entry mass vs original denominator",
        },
        {
            "term": "XRV L1 correction",
            "definition": "50,536 / 524,538 = 9.63%; 1.77% used N x 12",
            "meaning": "The old rates used different denominators",
        },
    ]

    protocol_rows = [
        {"component": "Data", "setting": "237,717 MIMIC-CXR-JPG AP/PA train images; 12 labels"},
        {"component": "Test", "setting": "Same 605 expert-labelled studies for all 55 models"},
        {"component": "Seeds", "setting": "7, 13, 42, 97, 123; method and loop paired within seed"},
        {"component": "Remove detection", "setting": "4-fold shuffled OOF; seed-specific MobileNetV3-small; max 100 epochs; patience 10"},
        {"component": "Remove action", "setting": "Top 20% of currently CL-flagged samples each loop; cumulative; five loops"},
        {"component": "Refine action", "setting": "Fixed cumulative XRV-seed13 tables; L5: 4,007 relabels + 2,215 masks; no new LLM calls"},
        {"component": "Final model", "setting": "MobileNetV3-small scratch; 224; Adam 0.001; batch 32; 50 fixed epochs; no early stopping"},
        {"component": "Study score", "setting": "Maximum image aggregation; weighted AUROC from full-precision per-label summaries"},
        {"component": "Uncertainty", "setting": "Paired seeds, shared-study and hierarchical seed/study bootstrap; 1,000 iterations"},
    ]

    rare_rows = [
        {"evidence": "Atelectasis", "support": "212 pos / 2 neg", "effect": "+0.1965", "interpretation": "Too little minority support for stable inference"},
        {"evidence": "Lung Lesion", "support": "51 pos / 2 neg", "effect": "-0.1510", "interpretation": "Too little minority support for stable inference"},
        {"evidence": "Pleural Other", "support": "22 pos / 1 neg", "effect": "+0.2000", "interpretation": "Too little minority support for stable inference"},
        {"evidence": "Consolidation", "support": "70 pos / 21 neg", "effect": "+0.0940", "interpretation": "Clearest stable positive label; 5/5 seeds"},
        {"evidence": "Edema", "support": "165 pos / 87 neg", "effect": "-0.0286", "interpretation": "Stable negative label; 0/5 positive seeds"},
        {"evidence": "Remove L5 shift", "support": "83.09% valid coverage", "effect": "Pneumothorax + retention 51.67%", "interpretation": "Prevalence shifts by -7.64 percentage points"},
        {"evidence": "Refine L5 shift", "support": "99.58% valid coverage", "effect": "Max absolute shift 1.36 pp", "interpretation": "Distribution is largely preserved"},
    ]

    evidence_mode_rows = [
        {"mode": "Frozen initial OOF post-action", "available": "Five seeds: remove and fixed refine", "answers": "Fair action comparison under unchanged initial evidence", "wording": "Model-consistency diagnostic"},
        {"mode": "Dynamic iterative OOF pre-action", "available": "Five seeds: remove only", "answers": "Remaining-data diagnostic before the named action", "wording": "Do not compare directly with frozen post-action"},
        {"mode": "Archived dynamic XRV", "available": "One detection seed", "answers": "Cross-backbone supporting evidence", "wording": "Single-seed support only"},
        {"mode": "Held-out expert test", "available": "Five paired training seeds", "answers": "Downstream training utility", "wording": "Not proof of label correctness"},
        {"mode": "Expert adjudication", "available": "Not yet available", "answers": "Whether proposed changes are clinically correct", "wording": "Required for independent label validation"},
    ]

    dqs_anchor_rows = []
    for method, loop, label in (
        ("baseline", 0, "Baseline"),
        ("remove", 1, "Remove L1"),
        ("remove", 5, "Remove L5"),
        ("refine", 1, "Refine L1"),
        ("refine", 5, "Refine L5"),
    ):
        row = select_one(frozen_rows, method=method, loop=str(loop))
        dqs_anchor_rows.append(
            {
                "condition": label,
                "raw_dqs": f"{float(row['dqs_flattened_mean']):.4f}",
                "coverage": f"{100 * float(row['valid_entry_coverage_mean']):.2f}%",
                "adjusted_dqs": f"{float(row['coverage_adjusted_dqs_mean']):.4f}",
                "healthy_entries": f"{float(row['estimated_healthy_entries_dqs_mean']):,.0f}",
            }
        )

    return {
        "performance_rows": performance_rows,
        "uncertainty_rows": uncertainty_rows,
        "dqs_chart_rows": dqs_chart_rows,
        "dqs_anchor_rows": dqs_anchor_rows,
        "dqs_definition_rows": dqs_definition_rows,
        "metric_rows": metric_rows,
        "protocol_rows": protocol_rows,
        "full_loop_rows": full_loop_rows,
        "rare_rows": rare_rows,
        "evidence_mode_rows": evidence_mode_rows,
        "coverage_by_loop": coverage_by_loop,
    }


def make_forest_plot(rows: list[dict[str, Any]], output: Path) -> None:
    width, height = 1600, 900
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    regular_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    bold_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

    def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
        path = bold_path if bold else regular_path
        return ImageFont.truetype(path, size) if Path(path).exists() else ImageFont.load_default()

    ink = "#1F2430"
    muted = "#657084"
    grid = "#E4E8EF"
    blue = "#2F6FB3"
    red = "#C7332F"

    draw.text((70, 42), "Hierarchical uncertainty across seeds and studies", fill=ink, font=font(42, True))
    draw.text(
        (70, 101),
        "Mean paired AUROC difference; 95% hierarchical bootstrap CI (1,000 resamples)",
        fill=muted,
        font=font(24),
    )

    plot_left, plot_right = 650, 1490
    plot_top, row_gap = 210, 128
    x_min, x_max = -0.07, 0.09

    def x_pos(value: float) -> float:
        return plot_left + (value - x_min) / (x_max - x_min) * (plot_right - plot_left)

    ticks = [-0.06, -0.03, 0.0, 0.03, 0.06, 0.09]
    for tick in ticks:
        x = x_pos(tick)
        color = "#4A4F58" if tick == 0 else grid
        line_width = 4 if tick == 0 else 2
        draw.line((x, plot_top - 28, x, plot_top + row_gap * 3 + 82), fill=color, width=line_width)
        label = "0" if tick == 0 else f"{tick:+.2f}"
        box = draw.textbbox((0, 0), label, font=font(21))
        draw.text((x - (box[2] - box[0]) / 2, 730), label, fill=muted, font=font(21))

    for idx, row in enumerate(rows):
        y = plot_top + idx * row_gap
        draw.line((70, y + 62, 1510, y + 62), fill="#F0F2F6", width=2)
        draw.text((80, y - 24), row["contrast"], fill=ink, font=font(27, True))
        ci_text = f"{row['mean_delta']:+.4f}  [{row['ci_low']:+.4f}, {row['ci_high']:+.4f}]"
        draw.text((80, y + 15), ci_text, fill=muted, font=font(21))
        color = blue if idx == 0 else red
        y_mark = y + 5
        draw.line((x_pos(row["ci_low"]), y_mark, x_pos(row["ci_high"]), y_mark), fill=color, width=8)
        draw.line((x_pos(row["ci_low"]), y_mark - 13, x_pos(row["ci_low"]), y_mark + 13), fill=color, width=5)
        draw.line((x_pos(row["ci_high"]), y_mark - 13, x_pos(row["ci_high"]), y_mark + 13), fill=color, width=5)
        radius = 12
        center = (x_pos(row["mean_delta"]), y_mark)
        if idx == 3:
            draw.ellipse(
                (center[0] - radius, center[1] - radius, center[0] + radius, center[1] + radius),
                fill="#FFFFFF",
                outline=color,
                width=6,
            )
        else:
            draw.ellipse(
                (center[0] - radius, center[1] - radius, center[0] + radius, center[1] + radius),
                fill=color,
                outline=color,
            )
        draw.text((1320, y + 20), row["seed_direction"], fill=muted, font=font(19))

    draw.text((650, 775), "Paired AUROC difference", fill=ink, font=font(24, True))
    draw.text(
        (70, 830),
        "n = 5 paired seeds. Minimum two-sided exact sign-flip p = 0.0625; all five-loop Holm-adjusted tests > 0.05.",
        fill=muted,
        font=font(20),
    )
    image.save(output, format="PNG", optimize=True)


def make_artifact(evidence: dict[str, Any], generated_at: str) -> dict[str, Any]:
    dataset_rows: dict[str, list[dict[str, Any]]] = {
        "headline": [{"paired_seeds": 5, "models": 55, "test_studies": 605, "refine_positive": 5}],
        "performance_trajectory": evidence["performance_rows"],
        "hierarchical_uncertainty": evidence["uncertainty_rows"],
        "dqs_quality_coverage": evidence["dqs_chart_rows"],
        "dqs_definition": evidence["dqs_definition_rows"],
        "heldout_and_stability": evidence["metric_rows"],
        "protocol": evidence["protocol_rows"],
        "full_loop": evidence["full_loop_rows"],
        "rare_and_shift": evidence["rare_rows"],
        "evidence_modes": evidence["evidence_mode_rows"],
    }
    sources = [
        source(
            "performance_ars",
            "Five-seed performance validation",
            "cxr_real_experiment/mobilenet_remove_vs_llm_refine_5seed_ars_validation_20260711_\u4e2d\u6587.md",
            "Full-precision weighted AUROC summaries, coverage, paired tests, and hierarchical bootstrap results.",
        ),
        source(
            "quality_ars",
            "Five-seed data-quality validation",
            "cxr_real_experiment/mobilenet_remove_vs_llm_refine_5seed_data_quality_ars_20260711_\u4e2d\u6587.md",
            "Frozen and dynamic OOF quality, denominator, supplementary scoring, stability, and distribution-shift evidence.",
        ),
        source(
            "protocol",
            "Evaluation protocol and audited slide structure",
            "cxr_real_experiment/evaluation_update_slides_20260712_\u4e2d\u6587.md",
            "Audited protocol, metric definitions, evidence modes, conclusions, and wording controls.",
        ),
    ]

    structured_source_specs = {
        "headline": (
            "headline_snapshot",
            "Evaluation headline snapshot",
            "results_archive/evaluation_20260711_remove_vs_refine_5seed_ars/score_summary.csv",
            "Presentation snapshot of experiment counts and paired-seed coverage.",
            ["score_summary.csv", "evaluation metadata"],
        ),
        "performance_trajectory": (
            "performance_snapshot",
            "Five-seed AUROC trajectory snapshot",
            "results_archive/evaluation_20260711_remove_vs_refine_5seed_ars/score_summary.csv",
            "Verified loop-level AUROC means and standard deviations used by the performance chart.",
            ["score_summary.csv"],
        ),
        "hierarchical_uncertainty": (
            "uncertainty_snapshot",
            "Hierarchical bootstrap contrast snapshot",
            "results_archive/evaluation_20260711_remove_vs_refine_5seed_ars/paired_bootstrap_statistics.csv",
            "Selected observed paired effects and hierarchical seed-and-study bootstrap intervals.",
            ["paired_bootstrap_statistics.csv"],
        ),
        "dqs_quality_coverage": (
            "dqs_snapshot",
            "Frozen-OOF DQS snapshot",
            "results_archive/evaluation_20260711_data_quality_5seed_ars/oof_quality_summary.csv",
            "Raw and coverage-adjusted DQS after each action under unchanged initial OOF evidence.",
            ["oof_quality_summary.csv"],
        ),
        "dqs_definition": (
            "dqs_definition_snapshot",
            "DQS definition and denominator audit snapshot",
            "cxr_real_experiment/mobilenet_remove_vs_llm_refine_5seed_data_quality_ars_20260711_\u4e2d\u6587.md",
            "Audited formula definitions and the corrected XRV Loop 1 denominator calculation.",
            ["Cleanlab dataset.py", "data-quality ARS report"],
        ),
        "heldout_and_stability": (
            "heldout_snapshot",
            "Held-out scoring and stability snapshot",
            "results_archive/evaluation_20260711_data_quality_5seed_ars/heldout_supplementary_metrics_summary.csv",
            "Primary AUROC, supplementary scoring rules, OOF ranking stability, and the expert-audit gap.",
            [
                "score_summary.csv",
                "heldout_supplementary_metrics_summary.csv",
                "oof_ranking_stability.csv",
                "remove_set_stability_summary.csv",
            ],
        ),
        "protocol": (
            "protocol_snapshot",
            "Exact experiment protocol snapshot",
            "cxr_real_experiment/evaluation_update_slides_20260712_\u4e2d\u6587.md",
            "Exact data, model, action, seed, scoring, and uncertainty settings.",
            ["evaluation scripts", "evaluation metadata", "audited slide structure"],
        ),
        "full_loop": (
            "full_loop_snapshot",
            "Full five-seed loop snapshot",
            "results_archive/evaluation_20260711_remove_vs_refine_5seed_ars/score_summary.csv",
            "All remove and refine loop means, paired baseline deltas, and seed directions.",
            ["score_summary.csv"],
        ),
        "rare_and_shift": (
            "rare_shift_snapshot",
            "Rare-label and distribution sensitivity snapshot",
            "results_archive/evaluation_20260711_remove_vs_refine_5seed_ars/per_label_deltas.csv",
            "Rare-label support, direct Loop 5 effects, and training-distribution sensitivity.",
            ["per_label_deltas.csv", "test_label_support.csv", "label_distribution_shift.csv"],
        ),
        "evidence_modes": (
            "evidence_modes_snapshot",
            "Evidence mode and wording-control snapshot",
            "cxr_real_experiment/evaluation_update_slides_20260712_\u4e2d\u6587.md",
            "Distinguishes frozen, dynamic, held-out, and expert evidence and the claims each supports.",
            ["five-seed ARS reports", "audited slide structure"],
        ),
    }
    structured_source_ids: dict[str, str] = {}
    for dataset, (source_id, label, path, description, tables_used) in structured_source_specs.items():
        sql = rows_to_snapshot_sql(dataset_rows[dataset])
        dataset_rows[dataset] = execute_snapshot_sql(sql)
        sources.append(
            source(
                source_id,
                label,
                path,
                description,
                sql=sql,
                executed_at=generated_at,
                tables_used=tables_used,
            )
        )
        structured_source_ids[dataset] = source_id

    cards = [
        {
            "id": "paired_seeds",
            "description": "Same five seeds paired across baseline and cleaning conditions.",
            "dataset": "headline",
            "sourceId": structured_source_ids["headline"],
            "metrics": [{"label": "Paired seeds", "field": "paired_seeds", "format": "number"}],
        },
        {
            "id": "models",
            "description": "5 baseline + 25 remove + 25 fixed-refine models.",
            "dataset": "headline",
            "sourceId": structured_source_ids["headline"],
            "metrics": [{"label": "Models evaluated", "field": "models", "format": "number"}],
        },
        {
            "id": "test_studies",
            "description": "The same expert-labelled study set for every model.",
            "dataset": "headline",
            "sourceId": structured_source_ids["headline"],
            "metrics": [{"label": "Expert test studies", "field": "test_studies", "format": "number"}],
        },
        {
            "id": "refine_positive",
            "description": "Refine Loop 5 improves over baseline in every paired seed.",
            "dataset": "headline",
            "sourceId": structured_source_ids["headline"],
            "metrics": [{"label": "Refine L5 positive seeds", "field": "refine_positive", "format": "number"}],
        },
    ]

    charts = [
        {
            "id": "performance_trajectory",
            "title": "Study-weighted AUROC across five cleaning loops",
            "subtitle": "Mean over five paired seeds; source data also retain seed SD.",
            "type": "line",
            "dataset": "performance_trajectory",
            "sourceId": structured_source_ids["performance_trajectory"],
            "valueFormat": "number",
            "encodings": {
                "x": {"field": "loop", "type": "ordinal", "label": "Cleaning loop"},
                "y": {"field": "mean_auroc", "type": "quantitative", "label": "Study-weighted AUROC", "format": "number"},
                "color": {"field": "method", "type": "nominal", "label": "Method"},
            },
        },
        {
            "id": "hierarchical_uncertainty",
            "title": "Selected paired contrasts and hierarchical uncertainty",
            "subtitle": "Mean AUROC difference; 95% seed-and-study bootstrap intervals are retained in the data table.",
            "type": "bar",
            "dataset": "hierarchical_uncertainty",
            "sourceId": structured_source_ids["hierarchical_uncertainty"],
            "valueFormat": "number",
            "encodings": {
                "x": {"field": "contrast", "type": "nominal", "label": "Contrast"},
                "y": {"field": "mean_delta", "type": "quantitative", "label": "Mean AUROC difference", "format": "number"},
            },
        },
        {
            "id": "dqs_quality_coverage",
            "title": "Raw and coverage-adjusted frozen-OOF DQS",
            "subtitle": "Five-seed mean after each action; baseline raw/adjusted DQS = 0.9277.",
            "type": "line",
            "dataset": "dqs_quality_coverage",
            "sourceId": structured_source_ids["dqs_quality_coverage"],
            "valueFormat": "number",
            "encodings": {
                "x": {"field": "loop", "type": "ordinal", "label": "Post-action loop"},
                "y": {"field": "score", "type": "quantitative", "label": "DQS", "format": "number"},
                "color": {"field": "series", "type": "nominal", "label": "Series"},
            },
        },
    ]

    tables = [
        {
            "id": "dqs_definition",
            "title": "DQS measures estimated consistency on the currently covered entries",
            "subtitle": "Adjusted DQS is a transparent project guardrail, not an official Cleanlab metric.",
            "dataset": "dqs_definition",
            "sourceId": structured_source_ids["dqs_definition"],
            "columns": table_columns([("term", "Term"), ("definition", "Exact definition"), ("meaning", "Interpretation")]),
        },
        {
            "id": "heldout_and_stability",
            "title": "Proper scoring rules and stability expose the remaining evidence gap",
            "subtitle": "Primary AUROC is full precision; AP/Brier/NLL are recomputed supplementary scoring rules.",
            "dataset": "heldout_and_stability",
            "sourceId": structured_source_ids["heldout_and_stability"],
            "columns": table_columns(
                [
                    ("evidence", "Evidence"),
                    ("baseline", "Baseline"),
                    ("remove", "Remove L5"),
                    ("refine", "Refine L5"),
                    ("interpretation", "Interpretation"),
                ]
            ),
        },
        {
            "id": "protocol",
            "title": "Exact reproducibility protocol",
            "subtitle": "Training-seed robustness is tested; full detection-plus-LLM robustness is not.",
            "dataset": "protocol",
            "sourceId": structured_source_ids["protocol"],
            "columns": table_columns([("component", "Component"), ("setting", "Exact setting")]),
        },
        {
            "id": "full_loop",
            "title": "Full five-seed loop table",
            "subtitle": "All values are five-seed means; deltas use the paired no-clean baseline.",
            "dataset": "full_loop",
            "sourceId": structured_source_ids["full_loop"],
            "columns": table_columns(
                [
                    ("loop", "Loop"),
                    ("remove_auroc", "Remove AUROC"),
                    ("remove_delta", "Remove delta"),
                    ("remove_positive", "Remove + seeds"),
                    ("refine_auroc", "Refine AUROC"),
                    ("refine_delta", "Refine delta"),
                    ("refine_positive", "Refine + seeds"),
                    ("direct", "Refine - Remove"),
                ]
            ),
        },
        {
            "id": "rare_and_shift",
            "title": "Rare-label and distribution sensitivity",
            "subtitle": "Large rare-label effects are unstable; repeated removal materially shifts coverage and prevalence.",
            "dataset": "rare_and_shift",
            "sourceId": structured_source_ids["rare_and_shift"],
            "columns": table_columns(
                [
                    ("evidence", "Evidence"),
                    ("support", "Support / coverage"),
                    ("effect", "Loop 5 effect"),
                    ("interpretation", "Interpretation"),
                ]
            ),
        },
        {
            "id": "evidence_modes",
            "title": "Evidence modes and wording controls",
            "subtitle": "Each mode answers a different question; none should be relabelled as expert correctness.",
            "dataset": "evidence_modes",
            "sourceId": structured_source_ids["evidence_modes"],
            "columns": table_columns(
                [
                    ("mode", "Evidence mode"),
                    ("available", "Available evidence"),
                    ("answers", "What it answers"),
                    ("wording", "Correct wording"),
                ]
            ),
        },
    ]

    blocks = [
        {"id": "title", "type": "markdown", "body": f"# {DECK_TITLE}"},
        {
            "id": "summary",
            "type": "markdown",
            "body": (
                "## Executive Summary\n\n"
                "- Seeds 7, 13, 42, 97, and 123 pair baseline, remove, and fixed-refine conditions.\n"
                "- All 55 models use the same 237,717-image training cohort and 605-study expert test set.\n"
                "- The comparison unit is the within-seed AUROC difference under a fixed final-training protocol.\n"
                "- Remove reruns OOF detection per seed; LLM refinement reuses fixed XRV-seed13 cumulative tables."
            ),
        },
        {"id": "headline_metrics", "type": "metric-strip", "cardIds": ["paired_seeds", "models", "test_studies", "refine_positive"]},
        {
            "id": "performance_heading",
            "type": "markdown",
            "body": "## Cleaning trajectories separate early removal from late refinement",
        },
        {"id": "performance_chart", "type": "chart", "chartId": "performance_trajectory"},
        {
            "id": "performance_interpretation",
            "type": "markdown",
            "sourceId": "performance_ars",
            "body": (
                "The two methods show different trajectories across five seeds. Simple remove peaks early, "
                "then loses most of its gain as cumulative deletion reaches 14.11%; fixed refinement rises "
                "later and retains all samples. Loop 5 is exploratory after viewing the full trajectory."
            ),
        },
        {
            "id": "uncertainty_heading",
            "type": "markdown",
            "body": "## Direction is consistent, but uncertainty remains material",
        },
        {"id": "uncertainty_chart", "type": "chart", "chartId": "hierarchical_uncertainty"},
        {
            "id": "uncertainty_interpretation",
            "type": "markdown",
            "sourceId": "performance_ars",
            "body": (
                "The trend is promising, not confirmatory. All four displayed directions are positive, but "
                "their hierarchical 95% intervals include zero; the direct Loop 5 advantage shrinks from "
                "+0.0260 to +0.0068 after excluding labels with minority support below five."
            ),
        },
        {
            "id": "dqs_definition_heading",
            "type": "markdown",
            "body": "## The denominator determines what raw DQS can claim",
        },
        {"id": "dqs_definition_table", "type": "table", "tableId": "dqs_definition"},
        {
            "id": "dqs_evidence_heading",
            "type": "markdown",
            "body": "## Denominator adjustment reverses the simple-remove story",
        },
        {"id": "dqs_chart", "type": "chart", "chartId": "dqs_quality_coverage"},
        {
            "id": "dqs_interpretation",
            "type": "markdown",
            "sourceId": "quality_ars",
            "body": (
                "Refinement improves estimated label health while preserving coverage. Remove Loop 5 has raw "
                "DQS 0.9745 but only 83.09% valid-entry coverage, so adjusted DQS falls to 0.8097; fixed refine "
                "reaches raw/adjusted DQS 0.9374/0.9334 with 99.58% coverage."
            ),
        },
        {
            "id": "scoring_heading",
            "type": "markdown",
            "body": "## Multiple evidence views support the quality-coverage explanation",
        },
        {"id": "scoring_table", "type": "table", "tableId": "heldout_and_stability"},
        {"id": "protocol_table", "type": "table", "tableId": "protocol"},
        {"id": "full_loop_table", "type": "table", "tableId": "full_loop"},
        {"id": "rare_table", "type": "table", "tableId": "rare_and_shift"},
        {"id": "evidence_modes_table", "type": "table", "tableId": "evidence_modes"},
        {
            "id": "conclusions",
            "type": "markdown",
            "sourceId": "performance_ars",
            "body": (
                "## Conclusions\n\n"
                "- Simple remove gives an early performance gain, but repeated removal loses coverage and shifts label distribution.\n"
                "- Fixed LLM refinement has the stronger late quality-coverage point estimate while retaining all training samples.\n"
                "- Five paired seeds make the trajectory more credible than the previous one-off result.\n"
                "- The strongest current claim is promising coverage-preserving refinement, not universal superiority."
            ),
        },
        {
            "id": "recommendations",
            "type": "markdown",
            "body": (
                "## Recommended Actions\n\n"
                "- Do not claim that raw DQS alone proves dataset repair or that every LLM relabel is clinically correct.\n"
                "- Lock one primary endpoint before confirmation and regenerate refinement tables from independent detection seeds.\n"
                "- Expert-audit both flagged entries and randomly sampled unflagged entries.\n"
                "- Report the exact protocol, all loops, and the fixed-table caveat in the thesis and viva."
            ),
        },
    ]

    manifest = {
        "version": 1,
        "surface": "report",
        "title": DECK_TITLE,
        "description": DECK_SUBTITLE,
        "generatedAt": generated_at,
        "cards": cards,
        "charts": charts,
        "tables": tables,
        "sources": sources,
        "blocks": blocks,
    }
    snapshot = {
        "version": 1,
        "generatedAt": generated_at,
        "status": "ready",
        "datasets": dataset_rows,
    }
    return {
        "surface": "report",
        "manifest": manifest,
        "snapshot": snapshot,
        "sources": sources,
        "package_info": {},
    }


def write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def run_portable_builder(artifact_path: Path, report_path: Path) -> None:
    command = [
        "npm",
        "run",
        "report:deliver",
        "--",
        "--input",
        str(artifact_path),
        "--output",
        str(report_path),
    ]
    subprocess.run(command, cwd=PLUGIN_ROOT, check=True)


def png_data_uri(path: Path) -> str:
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def adapt_report_for_slides(report_path: Path, output_path: Path, chart_images: list[Path]) -> None:
    soup = BeautifulSoup(report_path.read_text(encoding="utf-8"), "html.parser")
    figures = soup.select("figure.portable-chart-summary")
    if len(figures) != len(chart_images):
        raise ValueError(f"Expected {len(chart_images)} portable charts, found {len(figures)}")

    alt_texts = [
        "Five-seed study-weighted AUROC trajectories with seed standard deviations",
        "Four selected hierarchical bootstrap contrasts with 95 percent confidence intervals",
        "Raw and coverage-adjusted frozen-OOF DQS across five loops",
    ]
    for figure, image_path, alt_text in zip(figures, chart_images, alt_texts):
        replacement = soup.new_tag("figure")
        replacement["class"] = ["figure", "deck-static-chart"]
        image_tag = soup.new_tag("img")
        image_tag["src"] = png_data_uri(image_path)
        image_tag["alt"] = alt_text
        replacement.append(image_tag)
        caption = soup.new_tag("figcaption")
        caption.string = alt_text
        replacement.append(caption)
        figure.replace_with(replacement)

    for svg in soup.find_all("svg"):
        svg.decompose()
    for source_inventory in soup.select(".portable-sources"):
        source_inventory.decompose()

    h1 = soup.find("h1")
    if h1 is None:
        raise ValueError("Portable report has no h1 title")
    eyebrow = soup.new_tag("p")
    eyebrow["class"] = ["eyebrow"]
    eyebrow.string = "Five paired seeds | MobileNetV3-small | 12 July 2026"
    h1.insert_before(eyebrow)
    lede = soup.new_tag("p")
    lede["class"] = ["lede"]
    lede.string = DECK_SUBTITLE
    h1.insert_after(lede)

    style = soup.new_tag("style")
    style.string = (
        ".deck-static-chart img { display:block; width:100%; height:auto; } "
        ".deck-static-chart figcaption { color:#657084; font-size:0.85rem; }"
    )
    if soup.head:
        soup.head.append(style)
    output_path.write_text(str(soup), encoding="utf-8")


def run_slides_helper(report_path: Path) -> None:
    subprocess.run(
        [sys.executable, str(SLIDES_HELPER), str(report_path), "--out-dir", str(OUTPUT_DIR)],
        check=True,
    )


def find_shape(slide: Any, name: str) -> Any | None:
    return next((shape for shape in slide.shapes if shape.name == name), None)


def style_text_shape(shape: Any, text: str, size: float, color: str, bold: bool = False) -> None:
    shape.text_frame.clear()
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def add_small_label(slide: Any, text: str, x: float, y: float, width: float, color: str) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(0.22))
    box.name = f"label-{text.lower().replace(' ', '-')}"
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(7.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def add_speaker_note(slide: Any, text: str) -> bool:
    try:
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = text
        return True
    except (AttributeError, NotImplementedError):
        return False


def add_text_block(
    slide: Any,
    name: str,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float,
    color: str = "#1F2430",
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> Any:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    box.name = name
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    return box


def add_panel(slide: Any, name: str, x: float, y: float, width: float, height: float) -> Any:
    panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    panel.name = name
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string("F6F8FB")
    panel.line.color.rgb = RGBColor.from_string("DCE2EA")
    panel.line.width = Pt(0.8)
    return panel


def style_table_cell(
    cell: Any,
    text: str,
    *,
    fill: str,
    color: str,
    size: float,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.CENTER,
) -> None:
    cell.text = text
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
    cell.margin_left = Inches(0.04)
    cell.margin_right = Inches(0.04)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def add_mini_example_slide(presentation: Presentation) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_text_block(
        slide,
        "slide-title",
        FINAL_TITLES[-1],
        0.55,
        0.27,
        12.25,
        0.52,
        size=18,
        bold=True,
    )
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.55),
        Inches(0.91),
        Inches(12.23),
        Inches(0.035),
    )
    accent.name = "title-rule"
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor.from_string("2E4780")
    accent.line.fill.background()
    add_text_block(
        slide,
        "mini-example-subtitle",
        "Eight valid binary entries; p is the OOF probability P(y=1).",
        0.55,
        1.02,
        8.6,
        0.28,
        size=10.5,
        color="#657084",
    )

    toy_rows = [
        ("A", "0", "0.10", "0", "No"),
        ("B", "0", "0.20", "0", "No"),
        ("C", "0", "0.30", "0", "No"),
        ("D", "0", "0.90", "1", "YES"),
        ("E", "1", "0.95", "1", "No"),
        ("F", "1", "0.85", "1", "No"),
        ("G", "1", "0.70", "1", "No"),
        ("H", "1", "0.20", "0", "YES"),
    ]
    table_shape = slide.shapes.add_table(
        9,
        5,
        Inches(0.55),
        Inches(1.43),
        Inches(7.05),
        Inches(4.42),
    )
    table_shape.name = "mini-example-entry-table"
    table = table_shape.table
    for column, width in zip(table.columns, [0.52, 1.10, 1.20, 2.28, 1.95]):
        column.width = Inches(width)
    table.rows[0].height = Inches(0.46)
    for row_index in range(1, len(table.rows)):
        table.rows[row_index].height = Inches(0.495)
    headers = ["Entry", "Observed y", "OOF p", "Confident guess", "Counted in M_t?"]
    for column_index, header in enumerate(headers):
        style_table_cell(
            table.cell(0, column_index),
            header,
            fill="#2E4780",
            color="#FFFFFF",
            size=9.4,
            bold=True,
        )
    for row_index, values in enumerate(toy_rows, start=1):
        is_issue = values[-1] == "YES"
        row_fill = "#FDEBE9" if is_issue else ("#FFFFFF" if row_index % 2 else "#F5F7FA")
        for column_index, value in enumerate(values):
            style_table_cell(
                table.cell(row_index, column_index),
                value,
                fill=row_fill,
                color="#C7332F" if is_issue and column_index == 4 else "#1F2430",
                size=9.8,
                bold=is_issue and column_index in (0, 4),
            )

    add_panel(slide, "threshold-panel", 7.86, 1.43, 4.92, 1.38)
    add_text_block(
        slide,
        "threshold-heading",
        "1  Estimate class-specific thresholds",
        8.05,
        1.54,
        4.54,
        0.25,
        size=10.3,
        color="#2E4780",
        bold=True,
    )
    add_text_block(
        slide,
        "threshold-formulas",
        "t0 = mean(1-p | observed y=0) = 0.625\nt1 = mean(p | observed y=1) = 0.675",
        8.05,
        1.84,
        4.54,
        0.61,
        size=11.8,
        bold=True,
    )
    add_text_block(
        slide,
        "threshold-note",
        "A class is confident when its probability reaches its own threshold.",
        8.05,
        2.47,
        4.54,
        0.22,
        size=8.7,
        color="#657084",
    )

    add_panel(slide, "joint-panel", 7.86, 2.96, 4.92, 1.72)
    add_text_block(
        slide,
        "joint-heading",
        "2  Count off-diagonal confident entries",
        8.05,
        3.06,
        4.54,
        0.25,
        size=10.3,
        color="#2E4780",
        bold=True,
    )
    joint_shape = slide.shapes.add_table(
        3,
        3,
        Inches(8.10),
        Inches(3.39),
        Inches(3.30),
        Inches(0.86),
    )
    joint_shape.name = "mini-example-confident-joint"
    joint = joint_shape.table
    for column, width in zip(joint.columns, [1.34, 0.98, 0.98]):
        column.width = Inches(width)
    for row in joint.rows:
        row.height = Inches(0.285)
    joint_values = [
        ("", "Guess 0", "Guess 1"),
        ("Observed 0", "3", "1"),
        ("Observed 1", "1", "3"),
    ]
    for row_index, values in enumerate(joint_values):
        for column_index, value in enumerate(values):
            off_diagonal = (row_index, column_index) in {(1, 2), (2, 1)}
            style_table_cell(
                joint.cell(row_index, column_index),
                value,
                fill="#FDEBE9" if off_diagonal else ("#E8EDF5" if row_index == 0 or column_index == 0 else "#FFFFFF"),
                color="#C7332F" if off_diagonal else "#1F2430",
                size=8.7,
                bold=off_diagonal or row_index == 0 or column_index == 0,
            )
    add_text_block(
        slide,
        "joint-result",
        "M_t = 1 + 1 = 2",
        11.52,
        3.53,
        1.05,
        0.49,
        size=11.0,
        color="#C7332F",
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text_block(
        slide,
        "joint-note",
        "Rows = observed label; columns = Cleanlab's confident class guess.",
        8.05,
        4.35,
        4.54,
        0.21,
        size=8.5,
        color="#657084",
    )

    add_panel(slide, "raw-dqs-panel", 7.86, 4.83, 4.92, 1.02)
    add_text_block(
        slide,
        "raw-dqs-heading",
        "3  Compute raw DQS on the current entries",
        8.05,
        4.94,
        4.54,
        0.25,
        size=10.3,
        color="#2E4780",
        bold=True,
    )
    add_text_block(
        slide,
        "raw-dqs-formula",
        "DQS = 1 - M_t / |E_t| = 1 - 2/8 = 0.75",
        8.05,
        5.28,
        4.54,
        0.32,
        size=13.0,
        bold=True,
    )

    denominator = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.55),
        Inches(6.02),
        Inches(12.23),
        Inches(0.87),
    )
    denominator.name = "denominator-check"
    denominator.fill.solid()
    denominator.fill.fore_color.rgb = RGBColor.from_string("EEF3F8")
    denominator.line.color.rgb = RGBColor.from_string("B7C5D8")
    denominator.line.width = Pt(0.9)
    add_text_block(
        slide,
        "denominator-heading",
        "Denominator check",
        0.75,
        6.10,
        1.45,
        0.22,
        size=9.5,
        color="#2E4780",
        bold=True,
    )
    add_text_block(
        slide,
        "denominator-formula",
        "Delete D and H  ->  |E_1|=6, M_1=0  ->  raw DQS=1.00; coverage=6/8=0.75; adjusted DQS=0.75",
        2.28,
        6.09,
        10.17,
        0.28,
        size=11.4,
        bold=True,
    )
    add_text_block(
        slide,
        "denominator-message",
        "The survivor set looks perfectly clean, but estimated healthy-entry mass remains 6 of the original 8 entries.",
        2.28,
        6.46,
        10.17,
        0.21,
        size=9.0,
        color="#657084",
    )
    add_text_block(
        slide,
        "source-note",
        SOURCE_FOOTERS[-1],
        0.55,
        7.12,
        11.55,
        0.18,
        size=7.0,
        color="#657084",
    )


def postprocess_deck() -> dict[str, Any]:
    deck_path = OUTPUT_DIR / "deck.pptx"
    raw_deck_path = OUTPUT_DIR / "deck_helper_raw.pptx"
    shutil.copy2(deck_path, raw_deck_path)

    raw_plan_path = OUTPUT_DIR / "deck_plan.json"
    raw_plan = json.loads(raw_plan_path.read_text(encoding="utf-8"))
    (OUTPUT_DIR / "deck_plan_helper_raw.json").write_text(
        json.dumps(raw_plan, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )

    presentation = Presentation(str(deck_path))
    if len(presentation.slides) != 12:
        raise ValueError(f"Expected 12 helper slides, found {len(presentation.slides)}")

    # Helper order: cover, summary, 3 charts, 6 tables, conclusion.
    order = [0, 1, 2, 3, 5, 4, 6, 11, 7, 8, 9, 10]
    slide_ids = list(presentation.slides._sldIdLst)
    for slide_id in slide_ids:
        presentation.slides._sldIdLst.remove(slide_id)
    for index in order:
        presentation.slides._sldIdLst.append(slide_ids[index])

    add_mini_example_slide(presentation)
    if len(presentation.slides) != 13:
        raise ValueError(f"Expected 13 final slides, found {len(presentation.slides)}")

    notes_added = 0
    for index, slide in enumerate(presentation.slides):
        title_shape = find_shape(slide, "slide-title")
        if title_shape is None:
            raise ValueError(f"Slide {index + 1} has no named title shape")
        title_size = 30 if index == 0 else 20 if len(FINAL_TITLES[index]) <= 72 else 17.5
        style_text_shape(title_shape, FINAL_TITLES[index], title_size, "#1F2430", bold=True)

        source_shape = find_shape(slide, "source-note")
        if source_shape is not None:
            style_text_shape(source_shape, SOURCE_FOOTERS[index], 7.0, "#657084")

        callout_heading = find_shape(slide, "chart-callout-heading")
        if callout_heading is not None:
            style_text_shape(callout_heading, "Interpretation", 10, "#657084", bold=True)

        if index == 6:
            table_shape = next(
                (shape for shape in slide.shapes if getattr(shape, "has_table", False)),
                None,
            )
            if table_shape is not None:
                for row_index, row in enumerate(table_shape.table.rows):
                    if row_index == 0:
                        continue
                    for column_index, cell in enumerate(row.cells):
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(9.9 if column_index == 0 else 9.6)

        add_small_label(slide, f"{index + 1:02d}", 12.82, 0.08, 0.28, "#657084")
        if index >= 8:
            add_small_label(slide, "BACKUP", 11.72, 0.08, 0.85, "#2E4780")
        notes_added += int(add_speaker_note(slide, SPEAKER_NOTES[index]))

    presentation.save(str(deck_path))
    descriptive_path = OUTPUT_DIR / "evaluation_update_5seed_20260712.pptx"
    shutil.copy2(deck_path, descriptive_path)

    final_plan = []
    for final_index, raw_index in enumerate(order):
        item = dict(raw_plan[raw_index])
        item["slide_number"] = final_index + 1
        item["title"] = FINAL_TITLES[final_index]
        item["section"] = "main" if final_index < 8 else "backup"
        final_plan.append(item)
    final_plan.append(
        {
            "kind": "native_mini_example",
            "title": FINAL_TITLES[-1],
            "elements": [
                "slide-title",
                "mini-example-entry-table",
                "mini-example-confident-joint",
                "denominator-check",
                "source-note",
            ],
            "slide_number": 13,
            "section": "backup",
        }
    )
    write_json(raw_plan_path, final_plan)

    manifest_path = OUTPUT_DIR / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["outputs"]["deck_pptx"] = str(deck_path)
    manifest["outputs"]["descriptive_deck_pptx"] = str(descriptive_path)
    manifest["outputs"]["raw_helper_deck_pptx"] = str(raw_deck_path)
    manifest["outputs"]["final_postflight"] = str(OUTPUT_DIR / "final_postflight_checks.json")
    write_json(manifest_path, manifest)
    return {"notes_added": notes_added, "order": order, "descriptive_path": descriptive_path}


def final_postflight(notes_added: int) -> dict[str, Any]:
    deck_path = OUTPUT_DIR / "deck.pptx"
    presentation = Presentation(str(deck_path))
    checks: list[dict[str, Any]] = []

    def add(name: str, passed: bool, detail: Any) -> None:
        checks.append({"name": name, "status": "passed" if passed else "failed", "detail": detail})

    add("slide_count", len(presentation.slides) == 13, len(presentation.slides))
    ratio = presentation.slide_width / presentation.slide_height
    add("aspect_ratio", abs(ratio - (16 / 9)) < 0.01, ratio)

    actual_titles = []
    all_in_bounds = True
    shape_counts = []
    picture_count = 0
    table_count = 0
    source_count = 0
    for slide in presentation.slides:
        title_shape = find_shape(slide, "slide-title")
        actual_titles.append(title_shape.text.strip() if title_shape is not None else "")
        shape_counts.append(len(slide.shapes))
        source_count += int(find_shape(slide, "source-note") is not None)
        for shape in slide.shapes:
            all_in_bounds = all_in_bounds and shape.left >= 0 and shape.top >= 0
            all_in_bounds = all_in_bounds and shape.left + shape.width <= presentation.slide_width + 10
            all_in_bounds = all_in_bounds and shape.top + shape.height <= presentation.slide_height + 10
            picture_count += int(shape.shape_type == 13)
            table_count += int(getattr(shape, "has_table", False))

    add("titles", actual_titles == FINAL_TITLES, actual_titles)
    add("nonblank_slides", min(shape_counts) >= 3, shape_counts)
    add("shape_bounds", all_in_bounds, "all shapes remain within the 16:9 canvas")
    add("chart_images", picture_count == 3, picture_count)
    add("native_tables", table_count == 8, table_count)
    add("source_notes", source_count == 13, source_count)
    add("speaker_notes", notes_added == 13, notes_added)

    preflight = json.loads((OUTPUT_DIR / "preflight_checks.json").read_text(encoding="utf-8"))
    add("official_helper_preflight", preflight.get("status") == "passed", preflight.get("summary"))

    try:
        with zipfile.ZipFile(deck_path) as archive:
            corrupt = archive.testzip()
        add("pptx_zip_integrity", corrupt is None, corrupt or "all package entries readable")
    except zipfile.BadZipFile as exc:
        add("pptx_zip_integrity", False, str(exc))

    failed = [check for check in checks if check["status"] != "passed"]
    result = {
        "status": "passed" if not failed else "failed",
        "summary": {"checks": len(checks), "passed": len(checks) - len(failed), "failed": len(failed)},
        "checks": checks,
    }
    write_json(OUTPUT_DIR / "final_postflight_checks.json", result)
    if failed:
        raise RuntimeError(f"Final postflight failed: {[check['name'] for check in failed]}")
    return result


def write_build_notes() -> None:
    notes = {
        "audience": "technical",
        "question": "Can the five-seed cleaning results be trusted, and is DQS inflated by deletion?",
        "decision_useful_answer": (
            "The trajectory is more credible but not confirmatory. Raw DQS overstates repeated removal "
            "because coverage shrinks; fixed refinement has the stronger late quality-coverage point estimate."
        ),
        "chart_map": [
            {
                "slide": 3,
                "question": "How do methods change across loops?",
                "family": "highlighted multi-series line with error bars",
                "claim": "Removal peaks early; refinement rises later.",
                "source_image": str(PERFORMANCE_FIGURE),
            },
            {
                "slide": 4,
                "question": "Does uncertainty exclude no difference?",
                "family": "dot and interval forest plot",
                "claim": "Directions are positive but hierarchical intervals include zero.",
                "source_image": str(OUTPUT_DIR / "hierarchical_uncertainty_4contrast.png"),
            },
            {
                "slide": 6,
                "question": "Does DQS improve after holding the original denominator fixed?",
                "family": "paired line panels",
                "claim": "Remove raw DQS rises while adjusted DQS falls; refinement improves both.",
                "source_image": str(DQS_FIGURE),
            },
        ],
        "omissions": [
            "No causal claim: the design compares trained models and diagnostics, not clinical outcomes.",
            "No expert label-accuracy estimate is available.",
            "No full-pipeline five-seed LLM replication: refinement tables are fixed from XRV seed13.",
        ],
    }
    write_json(OUTPUT_DIR / "build_notes.json", notes)


def main() -> int:
    for required in (PERFORMANCE_FIGURE, DQS_FIGURE, PLUGIN_ROOT, SLIDES_HELPER):
        if not required.exists():
            raise FileNotFoundError(required)

    if OUTPUT_DIR.exists():
        shutil.rmtree(OUTPUT_DIR)
    OUTPUT_DIR.mkdir(parents=True)

    evidence = build_evidence()
    forest_path = OUTPUT_DIR / "hierarchical_uncertainty_4contrast.png"
    make_forest_plot(evidence["uncertainty_rows"], forest_path)

    generated_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    artifact = make_artifact(evidence, generated_at)
    artifact_path = OUTPUT_DIR / "artifact.json"
    report_path = OUTPUT_DIR / "report.html"
    conversion_report_path = OUTPUT_DIR / "report_for_slides.html"
    write_json(artifact_path, artifact)
    write_build_notes()

    run_portable_builder(artifact_path, report_path)
    adapt_report_for_slides(
        report_path,
        conversion_report_path,
        [PERFORMANCE_FIGURE, forest_path, DQS_FIGURE],
    )
    run_slides_helper(conversion_report_path)

    preflight = json.loads((OUTPUT_DIR / "preflight_checks.json").read_text(encoding="utf-8"))
    if preflight.get("status") != "passed":
        raise RuntimeError(f"Official slide preflight failed: {preflight.get('summary')}")

    postprocess = postprocess_deck()
    postflight = final_postflight(postprocess["notes_added"])
    print(
        json.dumps(
            {
                "status": postflight["status"],
                "deck": str(OUTPUT_DIR / "deck.pptx"),
                "named_deck": str(postprocess["descriptive_path"]),
                "report": str(report_path),
                "slides": 13,
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
