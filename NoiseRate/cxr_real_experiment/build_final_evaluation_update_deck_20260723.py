#!/usr/bin/env python3
"""Build the final four-seed, eight-loop evaluation update deck."""

from __future__ import annotations

import json
import math
import textwrap
import zipfile
from pathlib import Path
from typing import Any, Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
OUTPUT_DIR = HERE / "final_evaluation_update_20260723"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_DECK = OUTPUT_DIR / "final_evaluation_update_20260723.pptx"

RESULT_ROOT = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "results_mobilenetv3_own_top20_binary_v5_llm_refine_5seed/20260714_065539"
)
EVALUATION_DIR = (
    RESULT_ROOT / "evaluation_posthoc_4seed_8loop_excluding_seed7_20260723"
)
MECHANISM_DIR = (
    RESULT_ROOT / "evaluation_loop_mechanism_audit_4seed_excluding_seed7_20260723"
)
SYNTHETIC_FIGURE = (
    HERE
    / "evaluation_followup_20260713"
    / "synthetic_dqs_validation"
    / "synthetic_dqs_known_truth_trajectories.png"
)

SLIDE_W = 13.333
SLIDE_H = 7.5
MAIN_SLIDE_COUNT = 5


def color(hex_value: str) -> RGBColor:
    value = hex_value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


INK = color("#1F2430")
MUTED = color("#6F768A")
LIGHT_MUTED = color("#8C94A5")
GRID = color("#E3E7EF")
PANEL = color("#F7F8FB")
WHITE = color("#FFFFFF")
BLUE = color("#2E77BB")
BLUE_DARK = color("#2E4780")
BLUE_LIGHT = color("#EAF2FA")
RED = color("#C9362B")
RED_LIGHT = color("#FBECEA")
GREEN = color("#2A8068")
GREEN_LIGHT = color("#E8F4EF")
AMBER = color("#D78B27")
AMBER_LIGHT = color("#F9F0E2")
GREY = color("#525A6B")
GREY_LIGHT = color("#EEF0F4")

HEX_INK = "#1F2430"
HEX_MUTED = "#6F768A"
HEX_GRID = "#DDE1E9"
HEX_BLUE = "#2E77BB"
HEX_RED = "#C9362B"
HEX_GREEN = "#2A8068"
HEX_AMBER = "#D78B27"
HEX_GREY = "#525A6B"

ANALYSIS_FOOTER = (
    "Analysis seeds: 13, 42, 97, 123 (outcome-informed post-hoc set)  |  "
    "Error bars: seed SD unless otherwise stated"
)


def ensure_inputs() -> None:
    required = [
        EVALUATION_DIR / "performance_per_seed_stage.csv",
        EVALUATION_DIR / "performance_four_seed_summary.csv",
        EVALUATION_DIR / "performance_hierarchical_bootstrap.csv",
        EVALUATION_DIR / "performance_paired_statistics.csv",
        EVALUATION_DIR / "quality_four_seed_summary.csv",
        EVALUATION_DIR / "quality_paired_statistics.csv",
        EVALUATION_DIR / "quality_trajectories.png",
        EVALUATION_DIR / "hierarchical_auroc_forest.png",
        MECHANISM_DIR / "loop_mechanism_summary.csv",
        MECHANISM_DIR / "cross_loop_action_transitions.csv",
        MECHANISM_DIR / "review_reuse_diagnostics.png",
        MECHANISM_DIR / "quality_training_diagnostics.png",
        SYNTHETIC_FIGURE,
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing required inputs:\n" + "\n".join(missing))


def load_data() -> dict[str, pd.DataFrame]:
    return {
        "performance_seed": pd.read_csv(EVALUATION_DIR / "performance_per_seed_stage.csv"),
        "performance_summary": pd.read_csv(
            EVALUATION_DIR / "performance_four_seed_summary.csv"
        ),
        "performance_bootstrap": pd.read_csv(
            EVALUATION_DIR / "performance_hierarchical_bootstrap.csv"
        ),
        "performance_stats": pd.read_csv(
            EVALUATION_DIR / "performance_paired_statistics.csv"
        ),
        "quality_summary": pd.read_csv(
            EVALUATION_DIR / "quality_four_seed_summary.csv"
        ),
        "quality_stats": pd.read_csv(
            EVALUATION_DIR / "quality_paired_statistics.csv"
        ),
        "mechanism": pd.read_csv(MECHANISM_DIR / "loop_mechanism_summary.csv"),
        "transitions": pd.read_csv(
            MECHANISM_DIR / "cross_loop_action_transitions.csv"
        ),
    }


def value_at(
    frame: pd.DataFrame,
    column: str,
    **conditions: Any,
) -> float:
    mask = pd.Series(True, index=frame.index)
    for key, expected in conditions.items():
        mask &= frame[key].eq(expected)
    values = frame.loc[mask, column]
    if len(values) != 1:
        raise ValueError(
            f"Expected one row for {conditions} in {column}, found {len(values)}"
        )
    return float(values.iloc[0])


def add_text(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    *,
    size: float = 12,
    text_color: RGBColor = INK,
    bold: bool = False,
    font: str = "Arial",
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.0,
    name: str | None = None,
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = text_color
    return shape


def add_rich_text(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    paragraphs: list[list[dict[str, Any]]],
    *,
    margin: float = 0.0,
    bullet: bool = False,
    name: str | None = None,
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    for paragraph_index, runs in enumerate(paragraphs):
        paragraph = (
            frame.paragraphs[0]
            if paragraph_index == 0
            else frame.add_paragraph()
        )
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(5)
        if bullet:
            paragraph.text = ""
            paragraph.level = 0
            paragraph._p.get_or_add_pPr().insert(
                0,
                paragraph._p._new_buChar(),
            )
        for item in runs:
            run = paragraph.add_run()
            run.text = str(item["text"])
            run.font.name = item.get("font", "Arial")
            run.font.size = Pt(float(item.get("size", 12)))
            run.font.bold = bool(item.get("bold", False))
            run.font.color.rgb = item.get("color", INK)
    return shape


def add_box(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    fill: RGBColor = WHITE,
    line: RGBColor = GRID,
    line_width: float = 0.8,
    radius: bool = False,
    name: str | None = None,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_rule(
    slide,
    x: float,
    y: float,
    width: float,
    *,
    line_color: RGBColor = GRID,
    height: float = 0.01,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = line_color
    shape.line.color.rgb = line_color
    return shape


def add_circle(
    slide,
    x: float,
    y: float,
    diameter: float,
    *,
    fill: RGBColor,
    line: RGBColor | None = None,
    text: str | None = None,
    text_color: RGBColor = WHITE,
    text_size: float = 10,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x),
        Inches(y),
        Inches(diameter),
        Inches(diameter),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if text is not None:
        frame = shape.text_frame
        frame.clear()
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(text_size)
        run.font.bold = True
        run.font.color.rgb = text_color
    return shape


def add_connector(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    line_color: RGBColor = MUTED,
    width: float = 1.2,
):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    connector.line.color.rgb = line_color
    connector.line.width = Pt(width)
    connector.line.end_arrowhead = True
    return connector


def add_picture_contain(
    slide,
    path: Path,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    name: str | None = None,
):
    with Image.open(path) as image:
        image_ratio = image.width / image.height
    box_ratio = width / height
    if image_ratio >= box_ratio:
        rendered_width = width
        rendered_height = width / image_ratio
        rendered_x = x
        rendered_y = y + (height - rendered_height) / 2
    else:
        rendered_height = height
        rendered_width = height * image_ratio
        rendered_x = x + (width - rendered_width) / 2
        rendered_y = y
    picture = slide.shapes.add_picture(
        str(path),
        Inches(rendered_x),
        Inches(rendered_y),
        width=Inches(rendered_width),
        height=Inches(rendered_height),
    )
    if name:
        picture.name = name
    return picture


def add_header(slide, title: str, subtitle: str | None = None) -> None:
    add_text(
        slide,
        0.62,
        0.34,
        12.05,
        0.55,
        title,
        size=23,
        bold=True,
        name="slide-title",
    )
    if subtitle:
        add_text(
            slide,
            0.64,
            0.91,
            11.95,
            0.36,
            subtitle,
            size=10.6,
            text_color=MUTED,
            name="slide-subtitle",
        )


def add_footer(slide, source: str, *, analysis_footer: bool = True) -> None:
    footer = source
    if analysis_footer:
        footer = f"{source}  |  {ANALYSIS_FOOTER}"
    add_text(
        slide,
        0.62,
        7.16,
        12.05,
        0.18,
        footer,
        size=7.4,
        text_color=LIGHT_MUTED,
        name="slide-footer",
    )


def add_notes(slide, note: str) -> None:
    slide.notes_slide.notes_text_frame.text = note


def style_cell(
    cell,
    *,
    fill: RGBColor = WHITE,
    text_color: RGBColor = INK,
    size: float = 10.5,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    frame = cell.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in frame.paragraphs:
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = text_color


def add_table(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    rows: list[list[str]],
    *,
    column_widths: list[float] | None = None,
    header_fill: RGBColor = BLUE_LIGHT,
    header_color: RGBColor = BLUE_DARK,
    body_size: float = 10.3,
    name: str | None = None,
):
    shape = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(x), Inches(y), Inches(width), Inches(height)
    )
    if name:
        shape.name = name
    table = shape.table
    if column_widths:
        if not math.isclose(sum(column_widths), width, rel_tol=1e-6):
            raise ValueError("Column widths must sum to table width")
        for column, column_width in zip(table.columns, column_widths, strict=True):
            column.width = Inches(column_width)
    row_height = height / len(rows)
    for row in table.rows:
        row.height = Inches(row_height)
    for row_index, row_values in enumerate(rows):
        for column_index, value in enumerate(row_values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            style_cell(
                cell,
                fill=header_fill if row_index == 0 else WHITE,
                text_color=header_color if row_index == 0 else INK,
                size=10.1 if row_index == 0 else body_size,
                bold=row_index == 0,
                align=PP_ALIGN.LEFT if column_index == 0 else PP_ALIGN.CENTER,
            )
    return shape


def plot_quality_endpoint(data: dict[str, pd.DataFrame], output: Path) -> None:
    frame = data["quality_summary"]
    methods = ["baseline", "remove", "refine"]
    labels = ["Baseline", "Removal", "Refinement"]
    colors = [HEX_GREY, HEX_BLUE, HEX_RED]
    metrics = [
        ("Entry level", "dqs_flattened_mean", "coverage_adjusted_dqs_mean"),
        (
            "Sample level",
            "sample_issue_free_rate_mean",
            "coverage_adjusted_sample_health_mean",
        ),
    ]

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.labelcolor": HEX_INK,
            "xtick.color": HEX_MUTED,
            "ytick.color": HEX_INK,
        }
    )
    fig, axes = plt.subplots(1, 2, figsize=(10.7, 4.6), dpi=190)
    fig.patch.set_facecolor("white")
    for axis, (title, raw_column, adjusted_column) in zip(
        axes, metrics, strict=True
    ):
        raw_values: list[float] = []
        adjusted_values: list[float] = []
        for method in methods:
            loop = 0 if method == "baseline" else 8
            raw_values.append(value_at(frame, raw_column, method=method, loop=loop))
            adjusted_values.append(
                value_at(frame, adjusted_column, method=method, loop=loop)
            )
        positions = np.arange(len(methods))[::-1]
        axis.set_facecolor("white")
        for position, raw, adjusted, method_color in zip(
            positions, raw_values, adjusted_values, colors, strict=True
        ):
            axis.plot(
                [adjusted, raw],
                [position, position],
                color="#C9CED8",
                linewidth=2,
                zorder=1,
            )
            axis.scatter(
                [raw],
                [position],
                s=70,
                facecolors="white",
                edgecolors=method_color,
                linewidths=2,
                marker="s",
                zorder=3,
            )
            axis.scatter(
                [adjusted],
                [position],
                s=80,
                color=method_color,
                marker="o",
                zorder=4,
            )
            axis.text(
                raw + 0.006,
                position + 0.10,
                f"{raw:.3f}",
                color=method_color,
                fontsize=9.5,
                ha="left",
                va="bottom",
            )
            axis.text(
                adjusted + 0.006,
                position - 0.11,
                f"{adjusted:.3f}",
                color=method_color,
                fontsize=9.5,
                ha="left",
                va="top",
            )
        axis.set_title(title, fontsize=13, weight="bold", color=HEX_INK, pad=12)
        axis.set_yticks(positions)
        axis.set_yticklabels(labels)
        axis.set_xlim(0.64, 1.01)
        axis.set_xticks([0.65, 0.75, 0.85, 0.95, 1.00])
        axis.set_xlabel("DQS")
        axis.grid(axis="x", color=HEX_GRID, linewidth=0.8)
        axis.grid(axis="y", visible=False)
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
        axis.spines["left"].set_visible(False)
        axis.spines["bottom"].set_color("#AEB4C1")
        axis.tick_params(axis="y", length=0)
    handles = [
        Line2D(
            [0],
            [0],
            color=HEX_GREY,
            marker="s",
            markerfacecolor="white",
            markeredgewidth=1.8,
            linewidth=0,
            label="Raw DQS",
        ),
        Line2D(
            [0],
            [0],
            color=HEX_GREY,
            marker="o",
            linewidth=0,
            label="Coverage-adjusted DQS",
        ),
    ]
    fig.legend(
        handles=handles,
        loc="lower center",
        ncol=2,
        frameon=False,
        bbox_to_anchor=(0.5, -0.01),
        fontsize=10,
    )
    fig.subplots_adjust(left=0.13, right=0.99, top=0.87, bottom=0.21, wspace=0.30)
    fig.savefig(output, facecolor="white", bbox_inches="tight", pad_inches=0.06)
    plt.close(fig)


def plot_quality_trajectories(data: dict[str, pd.DataFrame], output: Path) -> None:
    frame = data["quality_summary"].copy()
    specifications = [
        (
            "Raw entry DQS",
            "dqs_flattened_mean",
            "dqs_flattened_sd",
            (0.72, 1.00),
        ),
        (
            "Coverage-adjusted entry DQS",
            "coverage_adjusted_dqs_mean",
            "coverage_adjusted_dqs_sd",
            (0.72, 1.00),
        ),
        (
            "Raw sample DQS",
            "sample_issue_free_rate_mean",
            "sample_issue_free_rate_sd",
            (0.65, 0.92),
        ),
        (
            "Coverage-adjusted sample DQS",
            "coverage_adjusted_sample_health_mean",
            "coverage_adjusted_sample_health_sd",
            (0.65, 0.92),
        ),
    ]
    styles = {
        "remove": {
            "color": HEX_BLUE,
            "error_color": "#205D96",
            "marker": "s",
            "linestyle": "--",
            "label": "Simple removal",
        },
        "refine": {
            "color": HEX_RED,
            "error_color": "#98281F",
            "marker": "o",
            "linestyle": "-",
            "label": "LLM refinement",
        },
    }

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.labelcolor": HEX_INK,
            "xtick.color": HEX_MUTED,
            "ytick.color": HEX_MUTED,
        }
    )
    fig, axes = plt.subplots(2, 2, figsize=(12.0, 5.35), dpi=190, sharex=True)
    fig.patch.set_facecolor("white")
    for index, (axis, specification) in enumerate(
        zip(axes.flat, specifications, strict=True)
    ):
        title, mean_column, sd_column, limits = specification
        baseline = value_at(
            frame,
            mean_column,
            method="baseline",
            loop=0,
        )
        baseline_sd = value_at(
            frame,
            sd_column,
            method="baseline",
            loop=0,
        )
        axis.axhspan(
            baseline - baseline_sd,
            baseline + baseline_sd,
            color=HEX_GREY,
            alpha=0.10,
            linewidth=0,
            zorder=0,
        )
        axis.axhline(
            baseline,
            color=HEX_GREY,
            linewidth=1.8,
            linestyle=(0, (5, 3)),
            zorder=1,
        )
        for method, style in styles.items():
            method_frame = frame.loc[frame["method"].eq(method)].sort_values("loop")
            axis.plot(
                method_frame["loop"],
                method_frame[mean_column],
                color=style["color"],
                linewidth=2.2,
                linestyle=style["linestyle"],
                marker=style["marker"],
                markersize=5.0,
                markerfacecolor="white" if method == "remove" else style["color"],
                markeredgewidth=1.4,
                zorder=3,
            )
            axis.errorbar(
                method_frame["loop"],
                method_frame[mean_column],
                yerr=method_frame[sd_column],
                fmt="none",
                ecolor=style["error_color"],
                elinewidth=1.55,
                capsize=4.6,
                capthick=1.55,
                zorder=4,
            )
        axis.set_title(title, fontsize=11.2, weight="bold", color=HEX_INK, pad=7)
        axis.set_xlim(0.7, 8.3)
        axis.set_ylim(*limits)
        axis.set_xticks(range(1, 9))
        axis.grid(axis="y", color=HEX_GRID, linewidth=0.7)
        axis.grid(axis="x", visible=False)
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
        axis.spines["left"].set_color("#AEB4C1")
        axis.spines["bottom"].set_color("#AEB4C1")
        if index in (0, 2):
            axis.set_ylabel("DQS")
        if index in (2, 3):
            axis.set_xlabel("Cleaning loop")

    handles = [
        Line2D(
            [0],
            [0],
            color=HEX_GREY,
            lw=1.8,
            linestyle=(0, (5, 3)),
            label="Baseline",
        ),
        Line2D(
            [0],
            [0],
            color=HEX_BLUE,
            lw=2.2,
            linestyle="--",
            marker="s",
            markerfacecolor="white",
            label="Simple removal",
        ),
        Line2D(
            [0],
            [0],
            color=HEX_RED,
            lw=2.2,
            marker="o",
            label="LLM refinement",
        ),
    ]
    fig.legend(
        handles=handles,
        loc="upper center",
        ncol=3,
        frameon=False,
        bbox_to_anchor=(0.76, 1.01),
        fontsize=9.7,
        handlelength=2.5,
        columnspacing=1.5,
    )
    fig.text(
        0.075,
        0.965,
        "Mean +/- 1 SD error bars and baseline bands (4 paired seeds)",
        fontsize=9.5,
        weight="bold",
        color=HEX_MUTED,
        ha="left",
    )
    fig.subplots_adjust(
        left=0.075,
        right=0.995,
        top=0.86,
        bottom=0.11,
        wspace=0.15,
        hspace=0.32,
    )
    fig.savefig(output, facecolor="white", bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)


def plot_auroc_trajectory(data: dict[str, pd.DataFrame], output: Path) -> None:
    frame = data["performance_seed"].copy()
    baseline = frame.loc[frame["method"].eq("baseline")]
    baseline_mean = float(baseline["study_weighted_auroc"].mean())
    baseline_sd = float(baseline["study_weighted_auroc"].std(ddof=1))

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10.5,
            "axes.labelcolor": HEX_INK,
            "xtick.color": HEX_MUTED,
            "ytick.color": HEX_MUTED,
        }
    )
    fig, axis = plt.subplots(figsize=(10.7, 5.2), dpi=190)
    fig.patch.set_facecolor("white")
    axis.set_facecolor("white")
    axis.axhspan(
        baseline_mean - baseline_sd,
        baseline_mean + baseline_sd,
        color=HEX_GREY,
        alpha=0.07,
        linewidth=0,
        zorder=0,
    )
    axis.axhline(
        baseline_mean,
        color=HEX_GREY,
        linewidth=2,
        linestyle=(0, (5, 3)),
        zorder=2,
    )
    styles = {
        "remove": {"color": HEX_BLUE, "marker": "s", "linestyle": "--"},
        "refine": {"color": HEX_RED, "marker": "o", "linestyle": "-"},
    }
    for method, style in styles.items():
        method_frame = frame.loc[frame["method"].eq(method)]
        for _, seed_frame in method_frame.groupby("seed"):
            seed_frame = seed_frame.sort_values("loop")
            axis.plot(
                seed_frame["loop"],
                seed_frame["study_weighted_auroc"],
                color=style["color"],
                linewidth=1,
                linestyle=style["linestyle"],
                alpha=0.20,
                zorder=1,
            )
        summary = (
            method_frame.groupby("loop", as_index=False)["study_weighted_auroc"]
            .agg(["mean", "std"])
            .reset_index()
            .sort_values("loop")
        )
        axis.errorbar(
            summary["loop"],
            summary["mean"],
            yerr=summary["std"],
            color=style["color"],
            linewidth=2.7,
            linestyle=style["linestyle"],
            marker=style["marker"],
            markersize=6.5,
            markerfacecolor="white" if method == "remove" else style["color"],
            markeredgewidth=1.7,
            capsize=3.5,
            capthick=1.3,
            zorder=4,
        )
    axis.axvline(8, color="#B3B9C5", linewidth=1.1, linestyle=(0, (2, 3)))
    axis.set_xlim(0.72, 8.28)
    axis.set_ylim(0.68, 0.795)
    axis.set_xticks(range(1, 9))
    axis.set_xlabel("Cleaning loop", labelpad=8)
    axis.set_ylabel("Study-weighted AUROC", labelpad=8)
    axis.grid(axis="y", color=HEX_GRID, linewidth=0.8)
    axis.grid(axis="x", visible=False)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.spines["left"].set_color("#AEB4C1")
    axis.spines["bottom"].set_color("#AEB4C1")
    handles = [
        Line2D(
            [0],
            [0],
            color=HEX_GREY,
            lw=2,
            linestyle=(0, (5, 3)),
            label=f"Baseline ({baseline_mean:.3f})",
        ),
        Line2D(
            [0],
            [0],
            color=HEX_BLUE,
            lw=2.7,
            linestyle="--",
            marker="s",
            markerfacecolor="white",
            label="Simple removal",
        ),
        Line2D(
            [0],
            [0],
            color=HEX_RED,
            lw=2.7,
            marker="o",
            label="LLM refinement",
        ),
    ]
    axis.legend(
        handles=handles,
        loc="upper left",
        frameon=False,
        ncol=3,
        fontsize=9.5,
        handlelength=2.6,
        columnspacing=1.2,
        borderaxespad=0,
    )
    axis.text(
        0.0,
        1.02,
        "Mean +/- 1 SD; faint traces are individual seeds",
        transform=axis.transAxes,
        fontsize=9.3,
        color=HEX_MUTED,
        ha="left",
    )
    fig.subplots_adjust(left=0.11, right=0.99, top=0.90, bottom=0.16)
    fig.savefig(output, facecolor="white", bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)


def plot_review_mechanism(data: dict[str, pd.DataFrame], output: Path) -> None:
    frame = data["mechanism"].sort_values("loop")
    loops = frame["loop"].to_numpy(dtype=int)
    keep = frame["keep_rate_mean"].to_numpy(dtype=float)
    relabel = frame["relabel_rate_mean"].to_numpy(dtype=float)
    mask = frame["mask_rate_mean"].to_numpy(dtype=float)
    first = frame["first_time_review_count_mean"].fillna(0).to_numpy(dtype=float)
    rereview = frame["rereview_count_mean"].fillna(0).to_numpy(dtype=float)
    total = first + rereview
    first_share = np.divide(first, total, out=np.ones_like(first), where=total > 0)
    rereview_share = np.divide(
        rereview, total, out=np.zeros_like(rereview), where=total > 0
    )

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.labelcolor": HEX_INK,
            "xtick.color": HEX_MUTED,
            "ytick.color": HEX_MUTED,
        }
    )
    fig, axes = plt.subplots(2, 1, figsize=(10.6, 5.8), dpi=190, sharex=True)
    fig.patch.set_facecolor("white")

    axis = axes[0]
    axis.bar(loops, keep * 100, color="#9AA2AC", label="Keep")
    axis.bar(
        loops,
        relabel * 100,
        bottom=keep * 100,
        color=HEX_RED,
        label="Relabel",
    )
    axis.bar(
        loops,
        mask * 100,
        bottom=(keep + relabel) * 100,
        color=HEX_AMBER,
        label="Mask",
    )
    axis.set_ylim(0, 100)
    axis.set_ylabel("Successful outcomes (%)")
    axis.set_title(
        "Among successful reviews, most entries are kept",
        loc="left",
        fontsize=12.2,
        weight="bold",
        color=HEX_INK,
    )
    axis.legend(loc="upper right", ncol=3, frameon=False, fontsize=9.2)
    axis.text(
        1,
        keep[0] * 100 / 2,
        f"{keep[0] * 100:.1f}% keep",
        ha="center",
        va="center",
        color="white",
        fontsize=8.8,
        weight="bold",
    )
    axis.text(
        8,
        keep[-1] * 100 / 2,
        f"{keep[-1] * 100:.1f}% keep",
        ha="center",
        va="center",
        color="white",
        fontsize=8.8,
        weight="bold",
    )

    axis = axes[1]
    axis.bar(loops, first_share * 100, color=HEX_RED, label="First-time review")
    axis.bar(
        loops,
        rereview_share * 100,
        bottom=first_share * 100,
        color="#9AA2AC",
        label="Previously reviewed",
    )
    axis.set_ylim(0, 100)
    axis.set_ylabel("Attempted reviews (%)")
    axis.set_xlabel("Refinement loop")
    axis.set_title(
        "Previously reviewed entries dominate the workload",
        loc="left",
        fontsize=12.2,
        weight="bold",
        color=HEX_INK,
    )
    axis.legend(loc="upper right", ncol=2, frameon=False, fontsize=9.2)
    axis.text(
        8,
        first_share[-1] * 100 + rereview_share[-1] * 50,
        f"{rereview_share[-1] * 100:.1f}%\nre-review",
        ha="center",
        va="center",
        color="white",
        fontsize=8.2,
        weight="bold",
    )
    axis.set_xticks(loops)

    for current_axis in axes:
        current_axis.grid(axis="y", color=HEX_GRID, linewidth=0.7)
        current_axis.set_axisbelow(True)
        current_axis.spines["top"].set_visible(False)
        current_axis.spines["right"].set_visible(False)
        current_axis.spines["left"].set_color("#AEB4C1")
        current_axis.spines["bottom"].set_color("#AEB4C1")
    fig.subplots_adjust(left=0.10, right=0.99, top=0.94, bottom=0.12, hspace=0.42)
    fig.savefig(output, facecolor="white", bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)


def plot_transition_bars(data: dict[str, pd.DataFrame], output: Path) -> None:
    frame = data["mechanism"].sort_values("loop")
    late = frame.loc[frame["loop"].isin([5, 6, 7, 8])]
    labels = ["L4→5", "L5→6", "L6→7", "L7→8"]
    all_labels = late["delta_study_weighted_auroc_mean"].to_numpy(dtype=float)
    support = late["delta_support_ge5_weighted_auroc_mean"].to_numpy(dtype=float)
    x = np.arange(len(labels))
    width = 0.34

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10.5,
            "axes.labelcolor": HEX_INK,
            "xtick.color": HEX_INK,
            "ytick.color": HEX_MUTED,
        }
    )
    fig, axis = plt.subplots(figsize=(10.4, 5.1), dpi=190)
    fig.patch.set_facecolor("white")
    axis.set_facecolor("white")
    bars_all = axis.bar(
        x - width / 2,
        all_labels,
        width,
        color=HEX_RED,
        label="All labels",
        zorder=3,
    )
    bars_support = axis.bar(
        x + width / 2,
        support,
        width,
        color=HEX_BLUE,
        label="Minority support >= 5",
        zorder=3,
    )
    axis.axhline(0, color=HEX_GREY, linewidth=1.2)
    axis.set_xticks(x)
    axis.set_xticklabels(labels)
    axis.set_ylabel("Change in study-weighted AUROC")
    axis.set_ylim(-0.052, 0.052)
    axis.grid(axis="y", color=HEX_GRID, linewidth=0.8)
    axis.set_axisbelow(True)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.spines["left"].set_color("#AEB4C1")
    axis.spines["bottom"].set_color("#AEB4C1")
    axis.legend(loc="upper left", frameon=False, fontsize=9.8)
    for bar in [*bars_all, *bars_support]:
        height = bar.get_height()
        axis.text(
            bar.get_x() + bar.get_width() / 2,
            height + (0.0022 if height >= 0 else -0.0022),
            f"{height:+.3f}",
            ha="center",
            va="bottom" if height >= 0 else "top",
            fontsize=8.9,
            color=HEX_INK,
        )
    fig.subplots_adjust(left=0.13, right=0.99, top=0.96, bottom=0.15)
    fig.savefig(output, facecolor="white", bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)


def add_method_badge(
    slide,
    x: float,
    y: float,
    text: str,
    *,
    fill: RGBColor,
    text_color: RGBColor = WHITE,
    width: float = 1.24,
) -> None:
    add_box(
        slide,
        x,
        y,
        width,
        0.34,
        fill=fill,
        line=fill,
        radius=True,
    )
    add_text(
        slide,
        x,
        y + 0.01,
        width,
        0.27,
        text,
        size=9.2,
        text_color=text_color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def slide_1(prs: Presentation, data: dict[str, pd.DataFrame]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Backup: evaluation setup and one-loop pipeline",
        "237,717 chest X-rays; four paired seeds; eight loops; sample removal versus entry-level LLM refinement.",
    )

    card_specs = [
        ("4", "paired training seeds", BLUE_DARK, BLUE_LIGHT),
        ("8", "cleaning loops", BLUE_DARK, BLUE_LIGHT),
        ("+0.048", "AUROC: refine vs remove", RED, RED_LIGHT),
        ("4 / 4", "seeds favor refinement at L8", GREEN, GREEN_LIGHT),
    ]
    x_positions = [0.68, 3.80, 6.92, 10.04]
    for x, (value, label, accent, fill) in zip(
        x_positions, card_specs, strict=True
    ):
        add_box(slide, x, 1.42, 2.62, 1.17, fill=fill, line=fill, radius=True)
        add_text(
            slide,
            x + 0.18,
            1.60,
            2.25,
            0.45,
            value,
            size=22,
            text_color=accent,
            bold=True,
        )
        add_text(
            slide,
            x + 0.18,
            2.10,
            2.25,
            0.28,
            label,
            size=9.8,
            text_color=MUTED,
        )

    add_text(
        slide,
        0.72,
        2.98,
        4.0,
        0.28,
        "ONE PAIRED LOOP",
        size=9.2,
        text_color=MUTED,
        bold=True,
    )
    flow = [
        ("Current-loop 4-fold\nout-of-fold (OOF)\npredictions", BLUE_LIGHT, BLUE_DARK, 9.2),
        ("Confident learning:\ntop 20% of flagged\nsamples", GREY_LIGHT, GREY, 9.4),
        (
            "Remove: delete sample\nRefine: review flagged\nimage-label entries",
            RED_LIGHT,
            RED,
            8.9,
        ),
        ("Apply cumulative actions;\nretrain and test", GREEN_LIGHT, GREEN, 9.8),
    ]
    flow_x = [0.72, 3.48, 6.24, 9.00]
    for index, (x, (text, fill, accent, text_size)) in enumerate(
        zip(flow_x, flow, strict=True)
    ):
        add_box(slide, x, 3.36, 2.26, 1.08, fill=fill, line=fill, radius=True)
        add_circle(
            slide,
            x + 0.16,
            3.67,
            0.43,
            fill=accent,
            text=str(index + 1),
            text_size=10,
        )
        add_text(
            slide,
            x + 0.72,
            3.55,
            1.35,
            0.70,
            text,
            size=text_size,
            bold=True,
            valign=MSO_ANCHOR.MIDDLE,
        )
        if index < len(flow) - 1:
            add_connector(slide, x + 2.32, 3.90, x + 2.70, 3.90)
    add_text(
        slide,
        11.58,
        3.55,
        0.78,
        0.72,
        "repeat\n× 8",
        size=12,
        text_color=RED,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    add_box(slide, 0.72, 4.55, 11.90, 0.32, fill=GREY_LIGHT, line=GREY_LIGHT)
    add_text(
        slide,
        0.94,
        4.62,
        11.45,
        0.18,
        "Selection OOF evidence is regenerated every loop. DQS is evaluated separately with each seed's initial OOF predictions frozen across all loops.",
        size=8.6,
        text_color=GREY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    add_box(slide, 0.72, 5.02, 3.78, 1.44, fill=WHITE, line=GRID)
    add_box(slide, 4.78, 5.02, 3.78, 1.44, fill=WHITE, line=GRID)
    add_box(slide, 8.84, 5.02, 3.78, 1.44, fill=WHITE, line=GRID)
    summary_cards = [
        (
            "QUALITY",
            "Raw Dataset Quality Score (DQS) can reward deletion; adjusted DQS keeps the original denominator.",
            BLUE_DARK,
        ),
        (
            "PERFORMANCE",
            "At Loop 8, refinement has higher held-out AUROC in this four-seed analysis.",
            RED,
        ),
        (
            "MECHANISM",
            "Late loops mostly re-check prior keep decisions; AUROC still oscillates.",
            GREEN,
        ),
    ]
    for x, (heading, body, accent) in zip(
        [0.72, 4.78, 8.84], summary_cards, strict=True
    ):
        add_text(
            slide,
            x + 0.22,
            5.20,
            3.32,
            0.25,
            heading,
            size=9.3,
            text_color=accent,
            bold=True,
        )
        add_text(
            slide,
            x + 0.22,
            5.55,
            3.30,
            0.68,
            body,
            size=10.7,
            text_color=INK,
        )

    add_footer(
        slide,
        "Source: paired MobileNetV3 real-CXR runs; frozen initial OOF evidence for DQS",
    )
    add_notes(
        slide,
        "这里补充一下完整的实验设置。训练集共有二十三万七千七百一十七张胸片，每张图最多对应十二个疾病标签。"
        "这次汇总使用四个 paired training seeds，每个方法都运行八个 cleaning loops。一个 sample 是一张胸片，"
        "一个 entry 是这张胸片和某一个疾病标签组成的一对。每一轮先在当前训练集上重新生成四折 out-of-fold predictions，"
        "再用 confident learning 找到可疑 sample pool，并选择其中最可疑的百分之二十。Simple removal 会删除整张 sample；"
        "LLM refinement 只审阅被 flag 的 image-label entries，并决定 keep、relabel 或 mask。所有 action 都会累积，"
        "之后重新训练 MobileNetV3 五十个 epochs，再在 held-out studies 上测试。需要区分两种 OOF 用法："
        "用于 selection 的 OOF evidence 每轮都会更新，而用于比较 DQS 的初始 OOF predictions 在八轮中保持固定。"
        "本次四个 seed 的汇总是 outcome-informed post-hoc analysis，所以我把它作为当前证据，而不是预先注册的 confirmatory result。",
    )


def slide_2(prs: Presentation, data: dict[str, pd.DataFrame]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "A higher raw DQS can come from deleting the denominator",
        "Raw DQS describes current survivors; adjusted DQS asks what fraction of the original supervision remains and is estimated issue-free.",
    )

    add_box(slide, 0.68, 1.38, 8.35, 4.95, fill=WHITE, line=GRID)
    add_text(
        slide,
        0.92,
        1.58,
        7.82,
        0.32,
        "MINI EXAMPLE: 10 labels, 8 correct and 2 incorrect",
        size=11.2,
        text_color=BLUE_DARK,
        bold=True,
    )
    add_text(
        slide,
        0.92,
        1.96,
        7.82,
        0.34,
        "Each round flags two suspected labels: one is incorrect; the other is a false positive and already correct.",
        size=10.5,
        text_color=MUTED,
    )
    rows = [
        [
            "Dataset state",
            "Correct / remaining",
            "Raw DQS",
            "Coverage",
            "Adjusted DQS",
        ],
        ["Baseline", "8 / 10", "0.800", "1.00", "0.800"],
        ["Removal after L1", "7 / 8", "0.875", "0.80", "0.700"],
        ["Removal after L2", "6 / 6", "1.000", "0.60", "0.600"],
        ["Refinement after L1", "9 / 10", "0.900", "1.00", "0.900"],
        ["Refinement after L2", "10 / 10", "1.000", "1.00", "1.000"],
    ]
    table_shape = add_table(
        slide,
        0.91,
        2.44,
        7.90,
        2.62,
        rows,
        column_widths=[2.16, 1.62, 1.26, 1.16, 1.70],
        name="mini-example-table",
    )
    table = table_shape.table
    for column_index in range(len(rows[0])):
        style_cell(
            table.cell(2, column_index),
            fill=BLUE_LIGHT,
            text_color=BLUE_DARK if column_index == 0 else INK,
            size=10.2,
            bold=column_index == 0,
            align=PP_ALIGN.LEFT if column_index == 0 else PP_ALIGN.CENTER,
        )
        style_cell(
            table.cell(3, column_index),
            fill=BLUE_LIGHT,
            text_color=BLUE_DARK if column_index == 0 else INK,
            size=10.2,
            bold=column_index == 0,
            align=PP_ALIGN.LEFT if column_index == 0 else PP_ALIGN.CENTER,
        )
        style_cell(
            table.cell(4, column_index),
            fill=RED_LIGHT,
            text_color=RED if column_index == 0 else INK,
            size=10.2,
            bold=column_index == 0,
            align=PP_ALIGN.LEFT if column_index == 0 else PP_ALIGN.CENTER,
        )
        style_cell(
            table.cell(5, column_index),
            fill=RED_LIGHT,
            text_color=RED if column_index == 0 else INK,
            size=10.2,
            bold=column_index == 0,
            align=PP_ALIGN.LEFT if column_index == 0 else PP_ALIGN.CENTER,
        )

    add_box(slide, 0.92, 5.30, 7.90, 0.72, fill=PANEL, line=PANEL, radius=True)
    add_text(
        slide,
        1.12,
        5.48,
        7.48,
        0.36,
        "Removal reaches raw DQS = 1.0 while correct labels retained fall from 8/10 to 6/10.",
        size=11.6,
        bold=True,
        text_color=BLUE_DARK,
    )

    add_box(slide, 9.30, 1.38, 3.35, 2.06, fill=WHITE, line=GRID)
    add_text(
        slide,
        9.55,
        1.64,
        2.85,
        0.30,
        "IN THIS TOY EXAMPLE",
        size=9.4,
        text_color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        9.55,
        2.05,
        2.85,
        0.42,
        "Raw = correct remaining / remaining",
        size=11.5,
        text_color=BLUE_DARK,
        bold=True,
    )
    add_text(
        slide,
        9.55,
        2.52,
        2.85,
        0.62,
        "Adjusted = correct remaining / original\n= Raw × Coverage",
        size=11.5,
        text_color=RED,
        bold=True,
    )

    add_box(slide, 9.30, 3.66, 3.35, 1.18, fill=GREEN_LIGHT, line=GREEN_LIGHT)
    add_text(
        slide,
        9.55,
        3.89,
        2.84,
        0.72,
        "Known-truth toy: deleting only incorrect labels keeps adjusted DQS flat. A decline means known-correct labels were also lost.",
        size=10.2,
        text_color=GREEN,
        bold=True,
    )

    add_box(slide, 9.30, 5.08, 3.35, 1.25, fill=GREY_LIGHT, line=GREY_LIGHT)
    add_text(
        slide,
        9.55,
        5.28,
        2.84,
        0.78,
        "Toy: known truth. Real DQS: estimated issue-free under frozen initial OOF evidence and confident learning, not expert-verified correctness.",
        size=9.9,
        text_color=GREY,
        bold=True,
    )
    add_footer(
        slide,
        "The toy uses known truth; the real DQS is a fixed-OOF confident-learning consistency diagnostic",
        analysis_footer=False,
    )
    add_notes(
        slide,
        "First, I want to clarify the difference between raw DQS and adjusted DQS, because this was a little confusing in last week's meeting. "
        "In this toy example, the original dataset has ten labels: eight are correct and two are wrong. So the raw DQS is eight out of ten, or 0.8. "
        "At the start, coverage is one, so the adjusted DQS is also 0.8. Now suppose simple removal deletes two suspicious labels: one wrong label and one correct label. "
        "Eight labels remain, and seven of them are correct. The raw DQS therefore increases to seven over eight, or 0.875. "
        "However, coverage is now eight over ten, so the adjusted DQS is 0.875 times 0.8, which is 0.7. "
        "From the perspective of the original dataset, we have actually reduced the number of correct labels from eight to seven. "
        "This is why I also report adjusted DQS: it prevents removal from looking better only because the denominator becomes smaller. "
        "With LLM refinement, the aim is to correct the wrong label while keeping the data, so raw and adjusted DQS can improve together. "
        "This is a known-truth toy example. In the real experiment, DQS is estimated from frozen initial out-of-fold evidence using confident learning; "
        "it is a model-consistency measure, not expert-verified label accuracy.",
    )


def slide_3(prs: Presentation, data: dict[str, pd.DataFrame], chart: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Removal appears cleaner under raw DQS because it retains less data",
        "Loop 8 endpoints: raw scores describe survivors; adjusted scores retain the original denominator.",
    )

    add_box(slide, 0.68, 1.34, 7.42, 4.62, fill=WHITE, line=GRID)
    add_picture_contain(slide, chart, 0.88, 1.53, 7.02, 4.15, name="quality-chart")
    add_text(
        slide,
        0.94,
        5.66,
        6.95,
        0.22,
        "Coverage at L8 — entries: remove 76.0%, refine 99.1%; samples: remove 74.2%, refine 99.3%.\n"
        "Refinement is below 100% because unresolved entries can be masked.",
        size=8.2,
        text_color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    add_box(slide, 8.34, 1.34, 4.31, 1.48, fill=RED_LIGHT, line=RED_LIGHT)
    add_text(
        slide,
        8.61,
        1.59,
        3.75,
        0.26,
        "LOOP 8 COVERAGE-ADJUSTED DQS",
        size=9.3,
        text_color=RED,
        bold=True,
    )
    add_text(
        slide,
        8.61,
        1.96,
        3.75,
        0.47,
        "Entry:  0.936 refine  vs  0.747 remove",
        size=12.6,
        text_color=INK,
        bold=True,
    )
    add_text(
        slide,
        8.61,
        2.38,
        3.75,
        0.26,
        "Sample: 0.768 refine  vs  0.670 remove",
        size=11.6,
        text_color=INK,
        bold=True,
    )

    add_box(slide, 8.34, 3.03, 4.31, 1.39, fill=WHITE, line=GRID)
    add_text(
        slide,
        8.61,
        3.24,
        3.75,
        0.24,
        "MINI EXAMPLE: ENTRY VS SAMPLE",
        size=9.3,
        text_color=BLUE_DARK,
        bold=True,
    )
    add_text(
        slide,
        8.61,
        3.60,
        3.73,
        0.64,
        "Entry = one image–disease-label pair.\nA sample is issue-free only if all its valid entries are issue-free.",
        size=10.5,
        text_color=INK,
    )

    add_box(slide, 8.34, 4.63, 4.31, 1.33, fill=PANEL, line=PANEL)
    add_text(
        slide,
        8.61,
        4.84,
        3.75,
        0.22,
        "PAIRED FOUR-SEED DIFFERENCES",
        size=9.2,
        text_color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        8.61,
        5.16,
        3.75,
        0.26,
        "Entry  +0.189  (4/4 seeds)",
        size=10.9,
        text_color=RED,
        bold=True,
    )
    add_text(
        slide,
        8.61,
        5.47,
        3.75,
        0.26,
        "Sample +0.098  (4/4 seeds)",
        size=10.9,
        text_color=RED,
        bold=True,
    )

    add_box(slide, 0.68, 6.17, 11.97, 0.70, fill=WHITE, line=GRID)
    add_text(
        slide,
        0.91,
        6.36,
        11.53,
        0.28,
        "Exact sign-flip p = .125 for both (minimum possible with 4 paired seeds); "
        "paired-t p = 2.9×10⁻⁷ entry and 2.5×10⁻⁵ sample, but the normality assumption cannot be checked at n = 4.",
        size=9.6,
        text_color=INK,
        align=PP_ALIGN.CENTER,
    )
    add_footer(
        slide,
        "Source: frozen-initial-OOF DQS after Loop 8 actions; exact paired and paired-t statistics",
    )
    add_notes(
        slide,
        "在完整八轮轨迹的基础上，我再把 Loop 8 单独拿出来量化两种方法的差异。图中的空心方块是 raw DQS，"
        "实心圆是 coverage-adjusted DQS。Simple removal 的 raw entry 和 raw sample DQS 都最高，"
        "但它只保留了大约百分之七十六的 entries 和百分之七十四的 samples。把分母固定回原始数据后，"
        "refinement 的 adjusted entry DQS 是零点九三六，removal 是零点七四七；sample level 分别是零点七六八和零点六七零。"
        "Refinement 的 coverage 略低于百分之百，是因为无法确定的 entries 可以被 mask。Sample DQS 的定义更严格："
        "只要一张图仍有一个有效 entry 被估计为有问题，这张 sample 就不会被算作完全 issue-free。"
        "四个 paired seeds 上，adjusted entry 和 sample DQS 的差值方向都一致。Exact sign-flip p-value 都是零点一二五，"
        "这是四个 seeds 能达到的最小双侧值；paired t-test 的 p-value 分别是二点九乘十的负七次方和二点五乘十的负五次方。"
        "这些 p-values 正是在量化 refinement 和 removal 的 DQS 差异，但由于只有四个 seeds，我仍然把 effect size、方向一致性和分母解释放在首位。",
    )


def slide_4(prs: Presentation, data: dict[str, pd.DataFrame], chart: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "At Loop 8, refinement has higher held-out AUROC across four seeds",
        "Mean study-weighted AUROC across four paired seeds; error bars show +/- 1 SD and faint lines show individual seeds.",
    )
    add_box(slide, 0.68, 1.34, 8.02, 4.95, fill=WHITE, line=GRID)
    add_picture_contain(slide, chart, 0.88, 1.55, 7.62, 4.47, name="auroc-chart")

    add_box(slide, 8.95, 1.34, 3.70, 2.70, fill=WHITE, line=GRID)
    add_text(
        slide,
        9.22,
        1.58,
        3.12,
        0.22,
        "LOOP 8 ENDPOINT",
        size=9.4,
        text_color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        9.22,
        1.93,
        3.15,
        0.38,
        "0.769 refinement",
        size=18,
        text_color=RED,
        bold=True,
    )
    add_text(
        slide,
        9.22,
        2.33,
        3.15,
        0.27,
        "0.721 removal  |  0.728 baseline",
        size=10.4,
        text_color=INK,
    )
    add_rule(slide, 9.22, 2.75, 3.10)
    add_text(
        slide,
        9.22,
        2.95,
        3.15,
        0.27,
        "+0.048 vs removal",
        size=11.4,
        text_color=RED,
        bold=True,
    )
    add_text(
        slide,
        9.22,
        3.28,
        3.15,
        0.26,
        "+0.041 vs baseline",
        size=11.4,
        text_color=RED,
        bold=True,
    )

    add_box(slide, 8.95, 4.30, 3.70, 2.06, fill=GREEN_LIGHT, line=GREEN_LIGHT)
    add_text(
        slide,
        9.22,
        4.55,
        3.15,
        0.24,
        "P-VALUES WITH FOUR SEEDS",
        size=9.2,
        text_color=GREEN,
        bold=True,
    )
    add_text(
        slide,
        9.22,
        4.93,
        3.15,
        0.32,
        "Exact paired p = .125",
        size=13.0,
        text_color=INK,
        bold=True,
    )
    add_text(
        slide,
        9.22,
        5.38,
        3.15,
        0.26,
        "Paired-t p = .0012",
        size=10.8,
        text_color=INK,
        bold=True,
    )
    add_text(
        slide,
        9.22,
        5.72,
        3.15,
        0.26,
        "Holm-adjusted paired-t p = .0093",
        size=10.3,
        text_color=INK,
        bold=True,
    )
    add_text(
        slide,
        9.22,
        6.03,
        3.15,
        0.22,
        "All 4 seeds favour refinement",
        size=9.5,
        text_color=GREEN,
        bold=True,
    )
    add_footer(
        slide,
        "Source: held-out study predictions; means and paired statistics across four training seeds",
    )
    add_notes(
        slide,
        "Next, I want to check whether the difference in DQS transfers to held-out model performance. The red line is LLM refinement, the blue line is simple removal, "
        "and the grey dashed line is the baseline without cleaning. At Loop 8, the mean study-weighted AUROC is 0.769 for refinement, 0.721 for removal, and 0.728 for the baseline. "
        "So refinement is 0.048 higher than removal and 0.041 higher than the baseline. "
        "To check whether this Loop 8 difference is consistent across random seeds, I calculated paired p-values using the four seeds. "
        "The exact p-value is 0.125, the paired t-test p-value is 0.0012, and the Holm-adjusted paired t-test p-value is 0.0093. "
        "All four seeds favour refinement. As on the previous slide, the exact test cannot reach the conventional 0.05 threshold with only four seeds, "
        "so the current result is consistent across these four seeds, but more paired random seeds are still needed for stronger evidence.",
    )


def slide_5(prs: Presentation, data: dict[str, pd.DataFrame], chart: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Later loops mostly re-review entries already judged keep",
        "The mechanism audit explains the falling action yield: review effort increasingly confirms previous decisions instead of making new edits.",
    )

    add_box(slide, 0.68, 1.34, 8.15, 5.28, fill=WHITE, line=GRID)
    add_picture_contain(slide, chart, 0.88, 1.53, 7.75, 4.86, name="review-chart")

    add_box(slide, 9.08, 1.34, 3.57, 2.30, fill=RED_LIGHT, line=RED_LIGHT)
    add_text(
        slide,
        9.34,
        1.58,
        3.05,
        0.26,
        "MINI EXAMPLE + UNIT KEY",
        size=9.3,
        text_color=RED,
        bold=True,
    )
    add_text(
        slide,
        9.34,
        1.84,
        3.03,
        0.60,
        "Attempted review = one flagged\n"
        "image-label entry sent to the LLM.\n"
        "Repeat = same entry seen earlier.",
        size=9.5,
        text_color=INK,
    )
    for index in range(20):
        row = index // 10
        col = index % 10
        is_repeat = index < 16
        add_circle(
            slide,
            9.36 + col * 0.27,
            2.48 + row * 0.28,
            0.17,
            fill=GREY if is_repeat else RED,
        )
    add_text(
        slide,
        9.34,
        3.05,
        3.02,
        0.35,
        "≈ 16–17 were already reviewed.\n≈ 3–4 were first-time reviews.",
        size=9.7,
        text_color=INK,
        bold=True,
    )

    add_box(slide, 9.08, 3.80, 3.57, 1.34, fill=WHITE, line=GRID)
    add_text(
        slide,
        9.34,
        3.96,
        3.03,
        0.38,
        "ACTION RATE = RELABEL OR MASK\nDenominator: successful reviews",
        size=9.4,
        text_color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        9.34,
        4.38,
        3.03,
        0.28,
        "31.5% at L1  →  6.5% at L8",
        size=12.2,
        text_color=RED,
        bold=True,
    )
    add_text(
        slide,
        9.34,
        4.70,
        3.03,
        0.36,
        "keep: unchanged  |  relabel: flip 0↔1\nmask: exclude uncertain",
        size=9.5,
        text_color=MUTED,
    )

    add_box(slide, 9.08, 5.28, 3.57, 1.34, fill=PANEL, line=PANEL)
    add_text(
        slide,
        9.34,
        5.46,
        3.03,
        0.22,
        "POOLED: 4 SEEDS × 8 LOOPS",
        size=9.4,
        text_color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        9.34,
        5.78,
        3.03,
        0.58,
        "58.8% of attempted reviews were repeats.\n"
        "95.5% of repeats were keep → keep.",
        size=9.7,
        text_color=INK,
        bold=True,
    )
    add_footer(
        slide,
        "Source: per-loop LLM action logs and cross-loop review identity audit",
    )
    add_notes(
        slide,
        "To understand why refinement becomes less active in the later loops, I checked what happens inside the LLM review process. "
        "One attempted review means that one flagged image-label entry is sent to the LLM. If the same entry was reviewed in an earlier loop, I count it as a repeat. "
        "Keep means no label change, relabel means switching zero and one, and mask means excluding an entry that remains uncertain. "
        "The top chart uses successful reviews. In Loop 1, 31.5 percent lead to a relabel or mask action. By Loop 8, this falls to 6.5 percent, so each loop is making fewer actual changes. "
        "The bottom chart looks at the attempted review workload. By Loop 8, around 82.5 percent of reviews are repeated entries. In a simple group of twenty reviews, "
        "this means around sixteen or seventeen have already been seen, and only three or four are first-time reviews. "
        "Across all four seeds and eight loops, 58.8 percent of attempted reviews are repeats, and 95.5 percent of repeat transitions are keep to keep. "
        "So the later review budget is mostly being used to confirm earlier keep decisions, and these entries can occupy places that might otherwise go to new suspicious entries. "
        "A practical next experiment is therefore to skip or deprioritise entries already reviewed as keep, and then check whether this increases the action yield and improves held-out performance. "
        "This result explains why the action yield falls, but by itself it does not explain the up-and-down AUROC pattern.",
    )


def slide_6(prs: Presentation, data: dict[str, pd.DataFrame], chart: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "The late-loop oscillation is consistent across these four seeds",
        "The audit supports a multi-factor interpretation rather than a forced monotonic-improvement story.",
    )

    add_box(slide, 0.68, 1.34, 7.55, 4.55, fill=WHITE, line=GRID)
    add_picture_contain(slide, chart, 0.91, 1.58, 7.10, 4.08, name="transition-chart")

    add_box(slide, 8.50, 1.34, 4.15, 1.43, fill=BLUE_LIGHT, line=BLUE_LIGHT)
    add_text(
        slide,
        8.78,
        1.57,
        3.58,
        0.24,
        "MINI EXAMPLE: MINORITY SUPPORT",
        size=9.2,
        text_color=BLUE_DARK,
        bold=True,
    )
    add_text(
        slide,
        8.78,
        1.91,
        3.56,
        0.62,
        "For each label: min(# positive, # negative held-out studies).\n"
        "2 minority studies: 1 changed study = 50%; 80 studies: 1.25%.",
        size=9.6,
        text_color=INK,
        bold=True,
    )

    diagnostic_cards = [
        (
            3.02,
            "LOW-SUPPORT LABELS",
            "Filtering labels with minority support < 5 reduces the dips, but does not remove them.",
            BLUE_DARK,
        ),
        (
            4.10,
            "SELECTION / OOF EVIDENCE",
            "Per-loop top-20% cutoffs and adjacent-loop selected-set overlap change smoothly; this is separate from frozen DQS evidence.",
            GREEN,
        ),
        (
            5.18,
            "TRAINING VARIATION",
            "Validation loss sometimes improves when AUROC falls, and worsens when AUROC recovers.",
            AMBER,
        ),
    ]
    for y, heading, body, accent in diagnostic_cards:
        add_box(slide, 8.50, y, 4.15, 0.86, fill=WHITE, line=GRID)
        add_text(
            slide,
            8.76,
            y + 0.16,
            3.60,
            0.20,
            heading,
            size=8.8,
            text_color=accent,
            bold=True,
        )
        add_text(
            slide,
            8.76,
            y + 0.41,
            3.60,
            0.34,
            body,
            size=9.2,
            text_color=INK,
        )

    add_box(slide, 0.68, 6.12, 11.97, 0.78, fill=RED_LIGHT, line=RED_LIGHT)
    add_text(
        slide,
        0.95,
        6.32,
        11.43,
        0.40,
        "Conclusion: refinement has higher Loop-8 AUROC in this post-hoc four-seed analysis, but later-loop performance is not monotonic. "
        "Current seeds jointly vary selection, LLM actions, and model training, so the mechanism evidence remains observational.",
        size=10.6,
        text_color=RED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(
        slide,
        "Source: four-seed Loop 1–8 mechanism audit; support sensitivity, selection overlap, and training logs",
    )
    add_notes(
        slide,
        "Finally, I went back to the fluctuations from Loop 4 to Loop 8. The AUROC drops from Loop 4 to 5, recovers from 5 to 6, "
        "drops again from 6 to 7, and recovers again from 7 to 8. All four seeds move in the same direction for each of these four transitions, "
        "so this is not caused by only one unusual seed. The first check is low-support labels. For some labels, the minority side has only one or two held-out studies, "
        "so a ranking change in one study can have a large effect on AUROC. After excluding labels with minority support below five, the dips become smaller but do not disappear. "
        "The second check is the selection evidence. The top-20-percent cutoff and the overlap between selected sets in adjacent loops both change smoothly, "
        "with no clear break at Loop 5 or Loop 7. The third check is the training process. Sometimes validation loss improves while AUROC falls, "
        "and sometimes the opposite happens, so training loss alone cannot explain the pattern either. "
        "Overall, the oscillation is consistent across these four seeds. Low-support labels amplify it, and repeated keep decisions reduce the later action yield, "
        "but no single factor explains all of the changes. These mechanism results are still observational. "
        "The clearest next test is to reduce repeated keep reviews and then see whether the later-loop action yield and AUROC trajectory change.",
    )


def backup_1(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Backup: exact DQS definitions",
        "The same construction is used at entry and sample level; only the unit of analysis changes.",
    )
    entry_rows = [
        ["Term", "Exact definition", "Interpretation"],
        [
            "Eₜ",
            "Valid disease-label entries after the Loop-t action",
            "Current entry denominator",
        ],
        [
            "Mₜ",
            "Estimated issues among Eₜ from confident-learning analysis of OOF predictions",
            "Estimated model-consistency issues",
        ],
        ["Raw entry DQS", "1 − Mₜ / |Eₜ|", "Health of currently covered entries"],
        ["Entry coverage", "|Eₜ| / |E₀|", "Fraction of original entries retained"],
        [
            "Adjusted entry DQS",
            "(|Eₜ| − Mₜ) / |E₀| = Raw × Coverage",
            "Healthy entry mass vs original denominator",
        ],
    ]
    sample_rows = [
        ["Term", "Exact definition", "Interpretation"],
        [
            "Sₜ",
            "Samples with at least one valid disease-label entry after the action",
            "Current sample denominator",
        ],
        [
            "Qₜ",
            "Samples in Sₜ with at least one valid entry flagged as an issue",
            "Strict problematic-sample count",
        ],
        ["Raw sample DQS", "1 − Qₜ / |Sₜ|", "Fraction of covered samples issue-free"],
        ["Sample coverage", "|Sₜ| / |S₀|", "Fraction of original samples retained"],
        [
            "Adjusted sample DQS",
            "(|Sₜ| − Qₜ) / |S₀| = Raw × Coverage",
            "Issue-free sample mass vs original denominator",
        ],
    ]
    add_text(
        slide,
        0.72,
        1.37,
        5.9,
        0.26,
        "ENTRY LEVEL",
        size=9.4,
        text_color=BLUE_DARK,
        bold=True,
    )
    add_table(
        slide,
        0.72,
        1.72,
        12.0,
        2.16,
        entry_rows,
        column_widths=[1.65, 6.75, 3.60],
        body_size=9.5,
    )
    add_text(
        slide,
        0.72,
        4.10,
        5.9,
        0.26,
        "SAMPLE LEVEL",
        size=9.4,
        text_color=RED,
        bold=True,
    )
    add_table(
        slide,
        0.72,
        4.45,
        12.0,
        2.16,
        sample_rows,
        column_widths=[1.65, 6.75, 3.60],
        header_fill=RED_LIGHT,
        header_color=RED,
        body_size=9.5,
    )
    add_footer(
        slide,
        "Definitions used in the four-seed frozen-initial-OOF quality evaluation",
        analysis_footer=False,
    )
    add_notes(
        slide,
        "这里给出 DQS 的精确定义。Entry level 中，E_t 表示 Loop t 的 action 之后仍然有效的 disease-label entries，"
        "M_t 表示其中根据固定 OOF predictions 和 confident learning 估计仍然存在 issue 的 entries。"
        "Raw entry DQS 用当前有效 entries 作为分母；entry coverage 表示原始 entries 中还有多少被保留；"
        "adjusted entry DQS 把分母固定为最初的 E_0。Sample level 使用完全对应的结构。S_t 是仍有至少一个有效 entry 的 samples，"
        "Q_t 是其中至少有一个有效 entry 被估计为 issue 的 samples。因此 sample DQS 是一个更严格的指标："
        "一张图只有在所有有效 entries 都没有被 flag 时，才会被算作 issue-free。",
    )


def backup_2(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Backup: known-truth controls validate the intended DQS behavior",
        "5,000 binary samples, 12% injected label noise, 10 synthetic seeds, 5 cleaning loops, and frozen informative OOF evidence.",
    )
    add_box(slide, 0.68, 1.34, 11.97, 5.60, fill=WHITE, line=GRID)
    add_picture_contain(
        slide,
        SYNTHETIC_FIGURE,
        0.88,
        1.50,
        11.57,
        5.23,
        name="synthetic-controls",
    )
    add_footer(
        slide,
        "Source: synthetic known-ground-truth validation; mean +/- SD across 10 seeds",
        analysis_footer=False,
    )
    add_notes(
        slide,
        "这里用已知真值的 synthetic experiment 检查 DQS 的方向是否符合预期。实验包含五千个二分类 samples，"
        "初始标签噪声率是百分之十二，共使用十个 random seeds，每个运行五个 cleaning loops。"
        "左上角把错误标签改正时，raw 和 adjusted DQS 都上升。右上角只删除错误标签时，raw DQS 接近一点零，"
        "而 adjusted DQS 仍然记录 coverage loss。左下角随机删除标签时，raw DQS 基本稳定，但 adjusted DQS 会下降，"
        "因为一部分正确信息也被删除。右下角故意把正确标签改错时，两种 DQS 都下降。"
        "因此，在已知真值的设置中，adjusted DQS 能够同时反映标签质量和数据保留量；真实实验中的数值仍然取决于 OOF evidence 的可靠性。",
    )


def backup_3(prs: Presentation, chart: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Across eight loops, adjusted DQS separates refinement from deletion",
        "Raw DQS rises as removal shrinks the denominator; coverage-adjusted DQS falls for removal while refinement improves.",
    )
    add_box(slide, 0.68, 1.34, 11.97, 5.04, fill=WHITE, line=GRID)
    add_picture_contain(
        slide,
        chart,
        0.86,
        1.49,
        11.61,
        4.70,
        name="quality-trajectories",
    )
    add_box(slide, 0.68, 6.49, 11.97, 0.43, fill=GREEN_LIGHT, line=GREEN_LIGHT)
    add_text(
        slide,
        0.91,
        6.60,
        11.50,
        0.22,
        "4/4 paired seeds agree: exact sign-flip p = .125; the conventional p < .05 threshold "
        "is only possible with at least 6 consistently directed seeds.",
        size=9.3,
        text_color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_notes(
        slide,
        "This slide shows the real DQS trajectories across all eight loops. The top row is entry level and the bottom row is sample level. "
        "The left column shows raw DQS, and the right column shows coverage-adjusted DQS. Blue is simple removal, red is LLM refinement, "
        "the grey dashed line is the baseline, and the vertical error bars show plus or minus one standard deviation across four training seeds. "
        "The baseline bands show the same uncertainty. The entry-level standard deviations are very small, so some error bars are still shorter than the point markers. "
        "If we only look at the left column, removal seems to keep improving: both raw entry DQS and raw sample DQS increase faster than refinement. "
        "But the right column fixes the denominator to the original dataset, and the interpretation changes. The adjusted DQS for removal keeps falling because more training information is deleted, "
        "while refinement keeps almost all of the coverage and gradually improves. Entry-level and sample-level results point in the same direction, "
        "so this is not just a Loop 8 endpoint; it is a consistent pattern across the eight loops. "
        "To check whether this DQS difference is consistent across random seeds, I calculated an exact paired p-value using the four seeds. The result is 0.125. "
        "All four seeds favour refinement over removal, so the direction is fully consistent in the current results. "
        "However, the conventional significance threshold is below 0.05. With only four seeds, this exact test cannot reach that threshold even when all four agree. "
        "At least six consistently directed seeds are needed before a p-value below 0.05 becomes possible. "
        "So we may need to add more paired random seeds in a planned follow-up. For now, the result shows four-out-of-four directional agreement, but not significance at the conventional 0.05 level.",
    )


def backup_4(prs: Presentation, data: dict[str, pd.DataFrame]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Backup: complete hierarchical AUROC contrasts and p-values",
        "Effect sizes and seed + study intervals are primary; p-values are supporting evidence with limited four-seed resolution.",
    )
    add_box(slide, 0.68, 1.34, 8.06, 5.55, fill=WHITE, line=GRID)
    add_picture_contain(
        slide,
        EVALUATION_DIR / "hierarchical_auroc_forest.png",
        0.88,
        1.53,
        7.66,
        5.18,
        name="auroc-forest",
    )
    rows = [
        ["Contrast", "Δ", "HCI 95%", "Exact p", "t p", "Holm t p"],
        ["Refine L8 − Remove L8", "+.048", "[+.014, +.075]", ".125", ".0012", ".0093"],
        ["Refine L8 − Baseline", "+.041", "[+.009, +.076]", ".125", ".0196", ".1373"],
        ["Remove L8 − Baseline", "−.007", "[−.051, +.039]", ".625", ".5812", "1.000"],
        ["Refine L8 − Refine L5", "+.051", "[+.008, +.086]", ".125", ".0566", ".3398"],
    ]
    add_table(
        slide,
        8.98,
        1.52,
        3.67,
        2.70,
        rows,
        column_widths=[1.25, 0.43, 0.88, 0.40, 0.34, 0.37],
        body_size=7.8,
    )
    add_box(slide, 8.98, 4.46, 3.67, 1.19, fill=AMBER_LIGHT, line=AMBER_LIGHT)
    add_text(
        slide,
        9.22,
        4.68,
        3.20,
        0.22,
        "WHY THE TESTS DISAGREE",
        size=9.0,
        text_color=AMBER,
        bold=True,
    )
    add_text(
        slide,
        9.22,
        5.01,
        3.18,
        0.48,
        "Exact sign-flip is robust but coarse at n=4. The paired t-test uses magnitude but assumes a normal difference distribution.",
        size=9.3,
        text_color=INK,
    )
    add_box(slide, 8.98, 5.87, 3.67, 1.02, fill=PANEL, line=PANEL)
    add_text(
        slide,
        9.22,
        6.08,
        3.20,
        0.56,
        "The four-seed analysis set was selected after inspecting outcomes; treat it as post-hoc evidence, not a preregistered confirmatory test.",
        size=9.1,
        text_color=MUTED,
        bold=True,
    )
    add_footer(
        slide,
        "HCI = hierarchical seed + study bootstrap interval; eight planned contrasts used for Holm adjustment",
    )
    add_notes(
        slide,
        "这里把 held-out AUROC 的统计结果完整展开。左侧森林图同时展示全部十二个标签和 minority support 至少为五的敏感性分析，"
        "每一行都给出 mean difference 和同时重采样 training seed 与 held-out study 的 hierarchical interval。"
        "右侧表格列出主要 contrasts 的 effect、hierarchical interval、exact sign-flip p-value、paired t-test p-value，"
        "以及 paired t-test 的 Holm-adjusted p-value。Exact sign-flip test 对分布假设要求较少，但四个 seeds 只有十六种符号组合，"
        "所以最小双侧 p-value 是零点一二五。Paired t-test 会利用差值大小，因此 p-value 更小，"
        "但差值的正态分布假设无法在 n 等于四时检查。再加上当前四-seed集合是在查看 outcome 后确定的，"
        "这些结果应当被解释为 post-hoc evidence，而不是 preregistered confirmatory evidence。",
    )


def backup_5(prs: Presentation, data: dict[str, pd.DataFrame]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Backup: proper scoring rules expose overconfident errors",
        "AUROC measures ranking; Brier and NLL also assess the probabilities assigned to held-out labels.",
    )
    rows = [
        ["Loop 8 metric", "Baseline", "Removal", "Refinement", "Better direction"],
        ["Study-weighted AUROC", "0.7280", "0.7209", "0.7692", "Higher"],
        ["Macro average precision", "0.8553", "0.8556", "0.8675", "Higher"],
        ["Micro Brier", "0.1631", "0.1602", "0.1527", "Lower"],
        ["Micro NLL", "0.7071", "1.0615", "0.6984", "Lower"],
    ]
    add_table(
        slide,
        0.72,
        1.54,
        7.45,
        2.60,
        rows,
        column_widths=[2.47, 1.16, 1.16, 1.34, 1.32],
        body_size=10.2,
    )
    add_box(slide, 0.72, 4.42, 7.45, 1.52, fill=PANEL, line=PANEL)
    add_text(
        slide,
        0.98,
        4.66,
        6.92,
        0.23,
        "INTERPRETATION",
        size=9.2,
        text_color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        0.98,
        5.00,
        6.92,
        0.69,
        "Refinement beats removal on Brier and NLL in 4/4 seeds. "
        "NLL versus baseline improves in only 1/4 seeds, so calibration improvement over baseline is not consistent.",
        size=10.8,
        text_color=INK,
        bold=True,
    )

    add_box(slide, 8.48, 1.54, 4.17, 4.40, fill=WHITE, line=GRID)
    add_text(
        slide,
        8.76,
        1.80,
        3.60,
        0.26,
        "MINI EXAMPLE: BOTH PREDICTIONS ARE WRONG",
        size=9.3,
        text_color=BLUE_DARK,
        bold=True,
    )
    add_text(
        slide,
        8.76,
        2.25,
        3.60,
        0.52,
        "True label = 0",
        size=15,
        text_color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_box(slide, 8.86, 2.95, 1.55, 1.40, fill=BLUE_LIGHT, line=BLUE_LIGHT)
    add_box(slide, 10.73, 2.95, 1.55, 1.40, fill=RED_LIGHT, line=RED_LIGHT)
    add_text(
        slide,
        9.02,
        3.17,
        1.23,
        0.30,
        "P(y=1)=.60",
        size=11.0,
        text_color=BLUE_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        9.02,
        3.62,
        1.23,
        0.34,
        "NLL = 0.92",
        size=11.0,
        text_color=BLUE_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        10.89,
        3.17,
        1.23,
        0.30,
        "P(y=1)=.99",
        size=11.0,
        text_color=RED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        10.89,
        3.62,
        1.23,
        0.34,
        "NLL = 4.61",
        size=11.0,
        text_color=RED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        8.82,
        4.70,
        3.48,
        0.82,
        "NLL penalizes the .99 mistake much more because the model was confidently wrong.",
        size=10.6,
        text_color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(
        slide,
        "Source: held-out predictions at Loop 8; averages across four training seeds",
    )
    add_notes(
        slide,
        "除了 AUROC，我还检查了 proper scoring rules。AUROC 只评价正负样本的排序，Brier score 和 negative log-likelihood"
        " 还会评价预测概率是否校准，以及模型在出错时是否过度自信。右边的小例子中，真实标签是零，"
        "两个模型都错误地预测为正类。把正类概率预测成零点六时，NLL 是零点九二；把它预测成零点九九时，"
        "NLL 上升到四点六一，因为第二个模型是非常自信地犯错。Loop 8 时，refinement 的 Brier 和 NLL 都优于 removal，"
        "而且四个 seeds 方向一致。不过 refinement 相对 baseline 的 NLL 只在一个 seed 上改善，"
        "所以目前可以说 refinement 比 removal 的概率质量更好，但不能说它相对 baseline 带来了稳定的 calibration improvement。",
    )


def backup_6(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Backup: detailed review-reuse and training diagnostics",
        "These checks support the mechanism summary but do not identify a single causal driver of AUROC.",
    )
    add_box(slide, 0.68, 1.34, 6.02, 5.56, fill=WHITE, line=GRID)
    add_box(slide, 6.92, 1.34, 5.73, 5.56, fill=WHITE, line=GRID)
    add_picture_contain(
        slide,
        MECHANISM_DIR / "review_reuse_diagnostics.png",
        0.88,
        1.60,
        5.62,
        5.04,
        name="review-reuse-detail",
    )
    add_picture_contain(
        slide,
        MECHANISM_DIR / "quality_training_diagnostics.png",
        7.12,
        1.60,
        5.33,
        5.04,
        name="training-detail",
    )
    add_footer(
        slide,
        "Source: four-seed mechanism audit; exploratory associations are repeated-measure and uncorrected",
    )
    add_notes(
        slide,
        "这里给出机制审计的完整诊断图。左侧把 first-time reviews 和 repeated reviews 分开，可以看到随着 loop 增加，"
        "第一次出现的 entries 持续减少，重复 entries 占据越来越多的 review workload，而且重复 review 的 action rate 很低。"
        "右侧把 DQS、selection cutoff、相邻 loop overlap、training loss 和 AUROC change 放在同一张图中。"
        "这些指标都没有在 Loop 5 或 Loop 7 出现共同的结构性断点。图中的 association analysis 还包含同一 seed 的重复测量，"
        "并且没有对多重比较进行 correction，因此它们只能用于排查可能机制，不能被解释为某个因素导致 AUROC 波动的因果证据。",
    )


def validate_data(data: dict[str, pd.DataFrame]) -> dict[str, Any]:
    quality = data["quality_summary"]
    performance = data["performance_summary"]
    stats = data["performance_stats"]
    quality_stats = data["quality_stats"]
    bootstrap = data["performance_bootstrap"]
    mechanism = data["mechanism"]

    checks: dict[str, bool] = {}
    checks["four_seeds"] = set(data["performance_seed"]["seed"].unique()) == {
        13,
        42,
        97,
        123,
    }
    checks["eight_loops"] = set(
        data["performance_seed"].loc[
            data["performance_seed"]["method"].eq("refine"), "loop"
        ]
    ) == set(range(1, 9))
    checks["refine_l8_auroc"] = math.isclose(
        value_at(
            performance,
            "study_weighted_auroc_mean",
            method="refine",
            loop=8,
        ),
        0.769208,
        abs_tol=5e-7,
    )
    checks["primary_auroc_delta"] = math.isclose(
        value_at(
            stats,
            "mean_improvement",
            metric="study_weighted_auroc",
            comparison="refine_L8_vs_remove_L8",
        ),
        0.048338,
        abs_tol=5e-7,
    )
    checks["primary_hierarchical_interval"] = (
        math.isclose(
            value_at(
                bootstrap,
                "ci_low_2p5",
                scope="all_labels",
                comparison="refine_L8_vs_remove_L8",
                uncertainty="hierarchical_seed_and_study",
            ),
            0.014225,
            abs_tol=5e-7,
        )
        and math.isclose(
            value_at(
                bootstrap,
                "ci_high_97p5",
                scope="all_labels",
                comparison="refine_L8_vs_remove_L8",
                uncertainty="hierarchical_seed_and_study",
            ),
            0.074729,
            abs_tol=5e-7,
        )
    )
    checks["adjusted_entry_endpoint"] = math.isclose(
        value_at(
            quality,
            "coverage_adjusted_dqs_mean",
            method="refine",
            loop=8,
        ),
        0.936014,
        abs_tol=5e-7,
    )
    checks["adjusted_sample_endpoint"] = math.isclose(
        value_at(
            quality,
            "coverage_adjusted_sample_health_mean",
            method="refine",
            loop=8,
        ),
        0.768126,
        abs_tol=5e-7,
    )
    checks["quality_exact_p"] = math.isclose(
        value_at(
            quality_stats,
            "exact_signflip_p_two_sided",
            metric="coverage_adjusted_dqs",
            comparison="refine_L8_vs_remove_L8",
        ),
        0.125,
        abs_tol=1e-12,
    )
    checks["loop8_rereview_share"] = math.isclose(
        value_at(mechanism, "rereview_count_mean", loop=8)
        / (
            value_at(mechanism, "rereview_count_mean", loop=8)
            + value_at(mechanism, "first_time_review_count_mean", loop=8)
        ),
        0.8248,
        abs_tol=0.002,
    )
    failed = [name for name, passed in checks.items() if not passed]
    if failed:
        raise AssertionError("Data validation failed: " + ", ".join(failed))
    return {"status": "passed", "checks": checks}


def validate_deck(path: Path, expected_titles: list[str]) -> dict[str, Any]:
    presentation = Presentation(path)
    if len(presentation.slides) != len(expected_titles):
        raise AssertionError(
            f"Expected {len(expected_titles)} slides, got {len(presentation.slides)}"
        )
    observed_titles: list[str] = []
    notes_present = 0
    out_of_bounds: list[dict[str, Any]] = []
    visible_text: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        title_shape = next(
            (shape for shape in slide.shapes if shape.name == "slide-title"),
            None,
        )
        if title_shape is None:
            raise AssertionError(f"Slide {slide_index} has no named title")
        observed_titles.append(title_shape.text.strip())
        note = slide.notes_slide.notes_text_frame.text.strip()
        if note:
            notes_present += 1
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                visible_text.append(shape.text)
            if (
                shape.left < -2
                or shape.top < -2
                or shape.left + shape.width > presentation.slide_width + 2
                or shape.top + shape.height > presentation.slide_height + 2
            ):
                out_of_bounds.append(
                    {
                        "slide": slide_index,
                        "shape": shape.name,
                        "left": shape.left,
                        "top": shape.top,
                        "width": shape.width,
                        "height": shape.height,
                    }
                )
    if observed_titles != expected_titles:
        raise AssertionError(
            "Title mismatch:\n"
            + json.dumps(
                {"expected": expected_titles, "observed": observed_titles},
                ensure_ascii=False,
                indent=2,
            )
        )
    if notes_present != len(expected_titles):
        raise AssertionError(
            f"Speaker notes missing: {notes_present}/{len(expected_titles)}"
        )
    if out_of_bounds:
        raise AssertionError(
            "Out-of-bounds shapes:\n"
            + json.dumps(out_of_bounds, ensure_ascii=False, indent=2)
        )
    combined = "\n".join(visible_text)
    forbidden = ["cleanlab", "XRV", "five-seed analysis set", "3-seed"]
    found_forbidden = [term for term in forbidden if term.lower() in combined.lower()]
    if found_forbidden:
        raise AssertionError(f"Forbidden visible terms: {found_forbidden}")
    required = [
        "Exact sign-flip p = .125",
        "paired-t p = .0012",
        "coverage-adjusted",
        "post-hoc",
    ]
    missing_required = [
        term for term in required if term.lower() not in combined.lower()
    ]
    if missing_required:
        raise AssertionError(f"Required visible text missing: {missing_required}")
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad is not None:
            raise AssertionError(f"Corrupt PPTX archive member: {bad}")
        note_parts = [
            name for name in archive.namelist() if name.startswith("ppt/notesSlides/")
        ]
        if len([name for name in note_parts if name.endswith(".xml")]) < len(
            expected_titles
        ):
            raise AssertionError("PPTX note parts are missing")
    return {
        "status": "passed",
        "slides": len(expected_titles),
        "main_slides": MAIN_SLIDE_COUNT,
        "backup_slides": len(expected_titles) - MAIN_SLIDE_COUNT,
        "speaker_notes": notes_present,
        "shape_bounds": "passed",
        "pptx_zip_test": "passed",
        "forbidden_terms": "passed",
        "required_statistical_disclosures": "passed",
    }


def write_outline(path: Path, titles: list[str], prs: Presentation) -> None:
    lines = [
        "# Final evaluation update: slide structure and speaker notes",
        "",
        "## Main slides",
        "",
    ]
    for index, (title, slide) in enumerate(zip(titles, prs.slides, strict=True), start=1):
        if index == MAIN_SLIDE_COUNT + 1:
            lines.extend(["## Backup slides", ""])
        notes = slide.notes_slide.notes_text_frame.text.strip()
        lines.extend(
            [
                f"### Slide {index}: {title}",
                "",
                notes,
                "",
            ]
        )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    ensure_inputs()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    data = load_data()
    data_validation = validate_data(data)

    quality_trajectory_chart = ASSET_DIR / "quality_trajectories_main.png"
    auroc_chart = ASSET_DIR / "auroc_trajectory.png"
    review_chart = ASSET_DIR / "review_mechanism.png"
    transition_chart = ASSET_DIR / "late_loop_transitions.png"
    plot_quality_trajectories(data, quality_trajectory_chart)
    plot_auroc_trajectory(data, auroc_chart)
    plot_review_mechanism(data, review_chart)
    plot_transition_bars(data, transition_chart)

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    presentation.core_properties.title = (
        "Four-seed eight-loop reliability and mechanism update"
    )
    presentation.core_properties.subject = (
        "DQS denominator checks, downstream performance, p-values, and loop audit"
    )
    presentation.core_properties.author = "Yihang"
    presentation.core_properties.keywords = (
        "DQS, confident learning, LLM refinement, removal, AUROC, mechanism audit"
    )

    slide_2(presentation, data)
    backup_3(presentation, quality_trajectory_chart)
    slide_4(presentation, data, auroc_chart)
    slide_5(presentation, data, review_chart)
    slide_6(presentation, data, transition_chart)
    slide_1(presentation, data)
    backup_1(presentation)
    backup_2(presentation)
    backup_4(presentation, data)
    backup_5(presentation, data)
    backup_6(presentation)

    expected_titles = [
        "A higher raw DQS can come from deleting the denominator",
        "Across eight loops, adjusted DQS separates refinement from deletion",
        "At Loop 8, refinement has higher held-out AUROC across four seeds",
        "Later loops mostly re-review entries already judged keep",
        "The late-loop oscillation is consistent across these four seeds",
        "Backup: evaluation setup and one-loop pipeline",
        "Backup: exact DQS definitions",
        "Backup: known-truth controls validate the intended DQS behavior",
        "Backup: complete hierarchical AUROC contrasts and p-values",
        "Backup: proper scoring rules expose overconfident errors",
        "Backup: detailed review-reuse and training diagnostics",
    ]
    presentation.save(OUTPUT_DECK)
    deck_validation = validate_deck(OUTPUT_DECK, expected_titles)

    outline_path = OUTPUT_DIR / "final_evaluation_update_20260723_talk_track_zh.md"
    write_outline(outline_path, expected_titles, presentation)
    preflight = {
        "status": "passed",
        "data_validation": data_validation,
        "deck_validation": deck_validation,
        "assets": {
            "quality_trajectories": str(quality_trajectory_chart),
            "auroc_trajectory": str(auroc_chart),
            "review_mechanism": str(review_chart),
            "late_loop_transitions": str(transition_chart),
        },
    }
    (OUTPUT_DIR / "preflight_checks.json").write_text(
        json.dumps(preflight, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    delivery = {
        "status": "passed",
        "final_deck": str(OUTPUT_DECK),
        "talk_track": str(outline_path),
        "slides": len(expected_titles),
        "main_slides": MAIN_SLIDE_COUNT,
        "backup_slides": len(expected_titles) - MAIN_SLIDE_COUNT,
        "speaker_notes_added": len(expected_titles),
        "preflight": str(OUTPUT_DIR / "preflight_checks.json"),
    }
    (OUTPUT_DIR / "delivery_summary.json").write_text(
        json.dumps(delivery, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(delivery, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
