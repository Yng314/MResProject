#!/usr/bin/env python3
"""Build the current reliability-validation report and slide deck."""

from __future__ import annotations

import base64
import json
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from bs4 import BeautifulSoup
from pptx import Presentation


HERE = Path(__file__).resolve().parent
OUTPUT_DIR = HERE / "reliability_validation_report_20260715"
EVIDENCE_DIR = HERE / "evaluation_followup_20260713"
SYNTH_DIR = EVIDENCE_DIR / "synthetic_dqs_validation"
QUALITY_DIR = EVIDENCE_DIR / "real_data_quality_figures"

SYNTH_SUMMARY = SYNTH_DIR / "synthetic_dqs_summary.csv"
SYNTH_ERRORS = SYNTH_DIR / "synthetic_dqs_error_summary.csv"
SYNTH_METADATA = SYNTH_DIR / "synthetic_dqs_validation.json"
REAL_SUMMARY = QUALITY_DIR / "quality_metrics_five_seed_summary.csv"
REAL_METADATA = QUALITY_DIR / "quality_figure_metadata.json"

CHART_IMAGES = [
    SYNTH_DIR / "synthetic_dqs_known_truth_trajectories.png",
    SYNTH_DIR / "synthetic_dqs_oof_quality_sensitivity.png",
    QUALITY_DIR / "entry_dqs_common_scale_five_seed.png",
    QUALITY_DIR / "sample_issue_free_rate_five_seed.png",
]

SLIDES_SKILL_DIR = Path(
    "/homes/yz3522/.codex/plugins/cache/openai-curated-remote/"
    "data-analytics/0.2.8-13ceeea1f599/skills/build-report/report-to-google-slides"
)
PLUGIN_ROOT = SLIDES_SKILL_DIR.parents[2]
SLIDES_HELPER = SLIDES_SKILL_DIR / "scripts" / "report_to_google_slides.py"

REPORT_TITLE = "Reliability checks for iterative label cleaning"
REPORT_SUBTITLE = (
    "Known-truth validation, denominator-aware quality measures, and sample-level evidence; "
    "final downstream Loop 5/8 endpoints remain pending."
)

SCENARIO_LABELS = {
    "oracle_correction": "Correct known errors",
    "oracle_removal": "Remove known errors",
    "random_removal": "Remove random entries",
    "harmful_correction": "Corrupt clean entries",
}


def require_files(paths: list[Path]) -> None:
    missing = [str(path) for path in paths if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing required inputs: {missing}")


def close(actual: float, expected: float, tolerance: float = 5e-4) -> None:
    if abs(actual - expected) > tolerance:
        raise AssertionError(f"Expected {expected:.6f}, got {actual:.6f}")


def validate_inputs(
    synthetic: pd.DataFrame,
    errors: pd.DataFrame,
    real: pd.DataFrame,
    synth_meta: dict[str, Any],
    real_meta: dict[str, Any],
) -> None:
    if len(synthetic) != 48:
        raise AssertionError(f"Expected 48 synthetic summary rows, got {len(synthetic)}")
    if len(errors) != 8:
        raise AssertionError(f"Expected 8 synthetic error rows, got {len(errors)}")
    if int(synth_meta["design"]["n_seeds"]) != 10:
        raise AssertionError("Synthetic seed count is not 10")
    if int(synth_meta["design"]["n_samples"]) != 5000:
        raise AssertionError("Synthetic entry count is not 5000")
    if list(real_meta["seeds"]) != [7, 13, 42, 97, 123]:
        raise AssertionError("Real-data seed set changed")

    baseline = real.loc[(real["method"] == "baseline") & (real["loop"] == 0)].iloc[0]
    removal = real.loc[(real["method"] == "remove") & (real["loop"] == 5)].iloc[0]
    refinement = real.loc[(real["method"] == "refine") & (real["loop"] == 5)].iloc[0]
    close(float(baseline["dqs_flattened_mean"]), 0.9277272571)
    close(float(removal["coverage_adjusted_dqs_mean"]), 0.8096839504)
    close(float(refinement["coverage_adjusted_dqs_mean"]), 0.9334233173)
    close(float(removal["coverage_adjusted_sample_health_mean"]), 0.6978500892)
    close(float(refinement["coverage_adjusted_sample_health_mean"]), 0.7535942469)


def source(
    source_id: str,
    label: str,
    path: str,
    description: str,
    metric_definitions: list[str] | None = None,
) -> dict[str, Any]:
    if path.endswith(".csv"):
        sql = f"SELECT * FROM read_csv_auto('{path}')"
    elif path.endswith(".json"):
        sql = f"SELECT * FROM read_json_auto('{path}')"
    else:
        sql = f"SELECT content FROM read_text('{path}')"
    query: dict[str, Any] = {
        "description": description,
        "engine": "duckdb",
        "sql": sql,
        "language": "sql",
        "executed_at": "2026-07-15T12:00:00Z",
    }
    if metric_definitions:
        query["metric_definitions"] = metric_definitions
    return {
        "id": source_id,
        "label": label,
        "path": path,
        "query": query,
    }


def source_paths() -> dict[str, str]:
    base = "cxr_real_experiment/evaluation_followup_20260713"
    return {
        "synthetic": f"{base}/synthetic_dqs_validation/synthetic_dqs_summary.csv",
        "synthetic_errors": f"{base}/synthetic_dqs_validation/synthetic_dqs_error_summary.csv",
        "synthetic_metadata": f"{base}/synthetic_dqs_validation/synthetic_dqs_validation.json",
        "quality": f"{base}/real_data_quality_figures/quality_metrics_five_seed_summary.csv",
        "quality_metadata": f"{base}/real_data_quality_figures/quality_figure_metadata.json",
        "statistics": f"{base}/own_top20_refinement_protocol.md",
    }


def real_row(real: pd.DataFrame, method: str, loop: int) -> pd.Series:
    rows = real.loc[(real["method"] == method) & (real["loop"] == loop)]
    if len(rows) != 1:
        raise ValueError(f"Expected one row for {method} loop {loop}, got {len(rows)}")
    return rows.iloc[0]


def build_entry_trajectory(real: pd.DataFrame) -> list[dict[str, Any]]:
    baseline = real_row(real, "baseline", 0)
    rows: list[dict[str, Any]] = []
    series = [
        ("Removal: raw", "remove", "dqs_flattened"),
        ("Removal: adjusted", "remove", "coverage_adjusted_dqs"),
        ("Refinement: raw", "refine", "dqs_flattened"),
        ("Refinement: adjusted", "refine", "coverage_adjusted_dqs"),
    ]
    for label, method, metric in series:
        rows.append(
            {
                "loop": 0,
                "series": label,
                "score": float(baseline[f"{metric}_mean"]),
                "sd": float(baseline[f"{metric}_std"]),
                "method": method,
                "view": "raw" if metric == "dqs_flattened" else "adjusted",
            }
        )
        for loop in range(1, 6):
            row = real_row(real, method, loop)
            rows.append(
                {
                    "loop": loop,
                    "series": label,
                    "score": float(row[f"{metric}_mean"]),
                    "sd": float(row[f"{metric}_std"]),
                    "method": method,
                    "view": "raw" if metric == "dqs_flattened" else "adjusted",
                }
            )
    return rows


def build_sample_trajectory(real: pd.DataFrame) -> list[dict[str, Any]]:
    baseline = real_row(real, "baseline", 0)
    rows: list[dict[str, Any]] = []
    series = [
        ("Removal: evaluable samples", "remove", "sample_issue_free_rate"),
        ("Removal: original denominator", "remove", "coverage_adjusted_sample_health"),
        ("Refinement: evaluable samples", "refine", "sample_issue_free_rate"),
        ("Refinement: original denominator", "refine", "coverage_adjusted_sample_health"),
    ]
    for label, method, metric in series:
        rows.append(
            {
                "loop": 0,
                "series": label,
                "score": float(baseline[f"{metric}_mean"]),
                "sd": float(baseline[f"{metric}_std"]),
                "method": method,
                "view": "current" if metric == "sample_issue_free_rate" else "adjusted",
            }
        )
        for loop in range(1, 6):
            row = real_row(real, method, loop)
            rows.append(
                {
                    "loop": loop,
                    "series": label,
                    "score": float(row[f"{metric}_mean"]),
                    "sd": float(row[f"{metric}_std"]),
                    "method": method,
                    "view": "current" if metric == "sample_issue_free_rate" else "adjusted",
                }
            )
    return rows


def build_artifact(
    synthetic: pd.DataFrame,
    errors: pd.DataFrame,
    real: pd.DataFrame,
    synth_meta: dict[str, Any],
) -> dict[str, Any]:
    generated_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    paths = source_paths()

    sources = [
        source(
            "synthetic_summary",
            "Known-truth synthetic validation",
            paths["synthetic"],
            "Ten-seed Loop 0-5 validation under four known-truth cleaning scenarios.",
            [
                "Raw DQS is compared with current-label accuracy among covered entries.",
                "Adjusted DQS is compared with healthy-label mass relative to the original denominator.",
            ],
        ),
        source(
            "synthetic_errors",
            "Synthetic OOF sensitivity",
            paths["synthetic_errors"],
            "Mean absolute error of raw and adjusted DQS under informative and weak OOF evidence.",
        ),
        source(
            "synthetic_design",
            "Synthetic design",
            paths["synthetic_metadata"],
            "Known-truth design and seed settings.",
        ),
        source(
            "quality_summary",
            "Five-seed quality",
            paths["quality"],
            "Loop 0-5 entry/sample means and SD.",
            [
                "Adjusted entry DQS = raw entry DQS multiplied by valid-entry coverage.",
                "Coverage-adjusted sample DQS = raw sample DQS multiplied by evaluable-sample coverage.",
            ],
        ),
        source(
            "quality_definitions",
            "Quality metric definitions",
            paths["quality_metadata"],
            "Definitions, denominator identities, seed set, and error-bar convention.",
        ),
        source(
            "statistics_plan",
            "Locked endpoint statistics plan",
            paths["statistics"],
            "Paired effect sizes, hierarchical confidence intervals, exact sign-flip tests, and Holm adjustment.",
        ),
    ]

    informative_errors = errors.loc[errors["evidence_quality"] == "informative_oof"]
    weak_errors = errors.loc[errors["evidence_quality"] == "weak_oof"]
    informative_mae = float(
        pd.concat([informative_errors["raw_abs_error"], informative_errors["adjusted_abs_error"]]).mean()
    )
    weak_mae = float(
        pd.concat([weak_errors["raw_abs_error"], weak_errors["adjusted_abs_error"]]).mean()
    )

    baseline = real_row(real, "baseline", 0)
    removal = real_row(real, "remove", 5)
    refinement = real_row(real, "refine", 5)

    headline = [
        {
            "known_noise_rate": float(synth_meta["design"]["noise_rate"]),
            "synthetic_seeds": int(synth_meta["design"]["n_seeds"]),
            "real_seeds": 5,
            "refinement_sample_coverage": float(refinement["sample_coverage_mean"]),
        }
    ]

    definitions = [
        {
            "metric": "Raw entry DQS",
            "plain_definition": "Estimated clean fraction among entries that remain valid now",
            "question_answered": "How clean are the surviving entries?",
        },
        {
            "metric": "Valid-entry coverage",
            "plain_definition": "Current valid entries divided by baseline valid entries",
            "question_answered": "How much of the original entry denominator remains?",
        },
        {
            "metric": "Adjusted entry DQS",
            "plain_definition": "Raw entry DQS multiplied by valid-entry coverage",
            "question_answered": "What fraction of the original entries remain and are estimated clean?",
        },
        {
            "metric": "Raw sample DQS",
            "plain_definition": "Fraction of evaluable samples with no valid entry flagged by confident learning",
            "question_answered": "How clean are the currently evaluable samples?",
        },
        {
            "metric": "Coverage-adjusted sample DQS",
            "plain_definition": "Raw sample DQS multiplied by evaluable-sample coverage",
            "question_answered": "What fraction of the original samples remain and are estimated clean?",
        },
    ]

    synth_l5 = synthetic.loc[
        (synthetic["evidence_quality"] == "informative_oof") & (synthetic["loop"] == 5)
    ].copy()
    synthetic_chart: list[dict[str, Any]] = []
    synthetic_table: list[dict[str, Any]] = []
    scenario_order = [
        "oracle_correction",
        "oracle_removal",
        "random_removal",
        "harmful_correction",
    ]
    for scenario in scenario_order:
        row = synth_l5.loc[synth_l5["scenario"] == scenario].iloc[0]
        label = SCENARIO_LABELS[scenario]
        values = [
            ("Accuracy among labels that remain", float(row["true_current_accuracy_mean"])),
            ("Raw DQS", float(row["raw_dqs_mean"])),
            (
                "Fraction of original labels that remain and are correct",
                float(row["true_healthy_mass_mean"]),
            ),
            ("Coverage-adjusted DQS", float(row["adjusted_dqs_mean"])),
        ]
        synthetic_chart.extend(
            {"scenario": label, "series": series, "score": value}
            for series, value in values
        )
        synthetic_table.append(
            {
                "scenario": label,
                "truth_covered": float(row["true_current_accuracy_mean"]),
                "raw_dqs": float(row["raw_dqs_mean"]),
                "truth_original": float(row["true_healthy_mass_mean"]),
                "adjusted_dqs": float(row["adjusted_dqs_mean"]),
                "coverage": float(row["coverage_mean"]),
            }
        )

    sensitivity: list[dict[str, Any]] = []
    for row in errors.itertuples(index=False):
        label = SCENARIO_LABELS[str(row.scenario)]
        evidence = "Informative OOF" if row.evidence_quality == "informative_oof" else "Weak OOF"
        sensitivity.extend(
            [
                {
                    "scenario": label,
                    "series": f"{evidence}: raw",
                    "mae": float(row.raw_abs_error),
                },
                {
                    "scenario": label,
                    "series": f"{evidence}: adjusted",
                    "mae": float(row.adjusted_abs_error),
                },
            ]
        )

    endpoint_quality = [
        {
            "method": "Baseline",
            "entry_raw": float(baseline["dqs_flattened_mean"]),
            "entry_adjusted": float(baseline["coverage_adjusted_dqs_mean"]),
            "entry_coverage": float(baseline["valid_entry_coverage_mean"]),
            "sample_raw": float(baseline["sample_issue_free_rate_mean"]),
            "sample_adjusted": float(baseline["coverage_adjusted_sample_health_mean"]),
            "sample_coverage": float(baseline["sample_coverage_mean"]),
        },
        {
            "method": "Simple removal, Loop 5",
            "entry_raw": float(removal["dqs_flattened_mean"]),
            "entry_adjusted": float(removal["coverage_adjusted_dqs_mean"]),
            "entry_coverage": float(removal["valid_entry_coverage_mean"]),
            "sample_raw": float(removal["sample_issue_free_rate_mean"]),
            "sample_adjusted": float(removal["coverage_adjusted_sample_health_mean"]),
            "sample_coverage": float(removal["sample_coverage_mean"]),
        },
        {
            "method": "LLM refinement, Loop 5",
            "entry_raw": float(refinement["dqs_flattened_mean"]),
            "entry_adjusted": float(refinement["coverage_adjusted_dqs_mean"]),
            "entry_coverage": float(refinement["valid_entry_coverage_mean"]),
            "sample_raw": float(refinement["sample_issue_free_rate_mean"]),
            "sample_adjusted": float(refinement["coverage_adjusted_sample_health_mean"]),
            "sample_coverage": float(refinement["sample_coverage_mean"]),
        },
    ]

    statistics_terms = [
        {
            "term": "Effect size",
            "plain_meaning": "The average paired change, such as refinement AUROC minus baseline AUROC",
            "reporting_role": "Primary: states how large the observed effect is",
        },
        {
            "term": "95% confidence interval",
            "plain_meaning": "A range of effect sizes compatible with seed and study uncertainty",
            "reporting_role": "Primary: shows precision and whether material alternatives remain plausible",
        },
        {
            "term": "Exact paired p-value",
            "plain_meaning": "How surprising the paired seed differences would be under no systematic effect",
            "reporting_role": "Supporting: with five seeds, the smallest two-sided exact value is 0.0625",
        },
        {
            "term": "Holm-adjusted p-value",
            "plain_meaning": "The exact p-value corrected for several pre-specified comparisons",
            "reporting_role": "Supporting: limits false positives from multiple testing",
        },
        {
            "term": "Mean +/- SD error bar",
            "plain_meaning": "The observed spread across seeds around the mean trajectory",
            "reporting_role": "Descriptive only: it is not a confidence interval or p-value",
        },
    ]

    datasets = {
        "headline": headline,
        "definitions": definitions,
        "synthetic_endpoint": synthetic_chart,
        "synthetic_exact": synthetic_table,
        "sensitivity": sensitivity,
        "entry_trajectory": build_entry_trajectory(real),
        "sample_trajectory": build_sample_trajectory(real),
        "endpoint_quality": endpoint_quality,
        "statistics_terms": statistics_terms,
        "validation_metrics": [
            {
                "informative_mae": informative_mae,
                "weak_mae": weak_mae,
                "informative_oof_auroc": float(informative_errors["oof_auroc_vs_truth"].mean()),
                "weak_oof_auroc": float(weak_errors["oof_auroc_vs_truth"].mean()),
            }
        ],
    }

    cards = [
        {
            "id": "known_noise",
            "description": "The synthetic experiment knows exactly which labels are wrong.",
            "dataset": "headline",
            "sourceId": "synthetic_design",
            "metrics": [{"label": "Known synthetic noise", "field": "known_noise_rate", "format": "percent"}],
        },
        {
            "id": "synthetic_seeds",
            "description": "Independent synthetic repetitions used for mean and SD.",
            "dataset": "headline",
            "sourceId": "synthetic_design",
            "metrics": [{"label": "Synthetic seeds", "field": "synthetic_seeds", "format": "number"}],
        },
        {
            "id": "real_seeds",
            "description": "Paired seeds in the current entry- and sample-level diagnostics.",
            "dataset": "headline",
            "sourceId": "quality_summary",
            "metrics": [{"label": "Real-data seeds", "field": "real_seeds", "format": "number"}],
        },
        {
            "id": "refinement_sample_coverage",
            "description": "Evaluable-sample coverage retained at Loop 5.",
            "dataset": "headline",
            "sourceId": "quality_summary",
            "metrics": [
                {
                    "label": "Refinement sample coverage",
                    "field": "refinement_sample_coverage",
                    "format": "percent",
                }
            ],
        },
    ]

    charts = [
        {
            "id": "synthetic_known_truth",
            "title": "Loop 5 synthetic truth targets and corresponding DQS values",
            "subtitle": "Four known-truth controls; the slide version shows complete Loop 0-5 trajectories and mean +/- SD over 10 seeds.",
            "showDescription": True,
            "type": "bar",
            "dataset": "synthetic_endpoint",
            "sourceId": "synthetic_summary",
            "valueFormat": "number",
            "encodings": {
                "x": {"field": "scenario", "type": "nominal", "label": "Known action"},
                "y": {"field": "score", "type": "quantitative", "label": "Score", "format": "number"},
                "color": {"field": "series", "type": "nominal", "label": "Quantity"},
            },
            "settings": {"groupMode": "grouped", "showValues": False},
            "surface": {"surface": "export", "showControls": False},
        },
        {
            "id": "oof_sensitivity",
            "title": "DQS error under informative and weak OOF evidence",
            "subtitle": "Mean absolute error against each score's matching known-truth target.",
            "showDescription": True,
            "type": "bar",
            "dataset": "sensitivity",
            "sourceId": "synthetic_errors",
            "valueFormat": "number",
            "encodings": {
                "x": {"field": "scenario", "type": "nominal", "label": "Known action"},
                "y": {"field": "mae", "type": "quantitative", "label": "Mean absolute error", "format": "number"},
                "color": {"field": "series", "type": "nominal", "label": "Evidence and metric"},
            },
            "settings": {"groupMode": "grouped", "showValues": False},
            "surface": {"surface": "export", "showControls": False},
        },
        {
            "id": "entry_quality",
            "title": "Entry-level raw and denominator-adjusted quality across five loops",
            "subtitle": "Five-seed mean; slide version adds individual seed trajectories and +/- 1 SD error bars on shared axes.",
            "showDescription": True,
            "type": "line",
            "dataset": "entry_trajectory",
            "sourceId": "quality_summary",
            "valueFormat": "number",
            "encodings": {
                "x": {"field": "loop", "type": "ordinal", "label": "Cleaning loop"},
                "y": {"field": "score", "type": "quantitative", "label": "Entry-level quality", "format": "number"},
                "color": {"field": "series", "type": "nominal", "label": "Method and denominator"},
                "tooltip": [
                    {"field": "sd", "type": "quantitative", "label": "Seed SD", "format": "number"},
                ],
            },
            "settings": {"showPoints": "always"},
            "surface": {"surface": "export", "showControls": False},
        },
        {
            "id": "sample_quality",
            "title": "Raw and coverage-adjusted sample DQS across five loops",
            "subtitle": "A sample is problematic if any valid entry is flagged; slide version adds individual seeds and +/- 1 SD.",
            "showDescription": True,
            "type": "line",
            "dataset": "sample_trajectory",
            "sourceId": "quality_summary",
            "valueFormat": "number",
            "encodings": {
                "x": {"field": "loop", "type": "ordinal", "label": "Cleaning loop"},
                "y": {"field": "score", "type": "quantitative", "label": "Sample-level proportion", "format": "number"},
                "color": {"field": "series", "type": "nominal", "label": "Method and denominator"},
                "tooltip": [
                    {"field": "sd", "type": "quantitative", "label": "Seed SD", "format": "number"},
                ],
            },
            "settings": {"showPoints": "always"},
            "surface": {"surface": "export", "showControls": False},
        },
    ]

    def columns(items: list[tuple[str, str, str]]) -> list[dict[str, str]]:
        return [{"field": field, "label": label, "type": kind} for field, label, kind in items]

    tables = [
        {
            "id": "metric_definitions",
            "title": "The paired metrics answer two different denominator questions",
            "subtitle": "Raw metrics describe current survivors; adjusted metrics retain the baseline denominator.",
            "showDescription": True,
            "dataset": "definitions",
            "sourceId": "quality_definitions",
            "density": "spacious",
            "columns": columns(
                [
                    ("metric", "Metric", "text"),
                    ("plain_definition", "Plain definition", "text"),
                    ("question_answered", "Question answered", "text"),
                ]
            ),
        },
        {
            "id": "synthetic_exact",
            "title": "Backup: exact Loop 5 synthetic targets and scores",
            "subtitle": "Informative OOF evidence; means across 10 synthetic seeds.",
            "showDescription": True,
            "dataset": "synthetic_exact",
            "sourceId": "synthetic_summary",
            "density": "spacious",
            "columns": [
                {"field": "scenario", "label": "Known action", "type": "text"},
                {"field": "truth_covered", "label": "Accuracy among retained labels", "format": "number"},
                {"field": "raw_dqs", "label": "Raw DQS", "format": "number"},
                {"field": "truth_original", "label": "Fraction retained and correct", "format": "number"},
                {
                    "field": "adjusted_dqs",
                    "label": "Coverage-adjusted DQS",
                    "format": "number",
                },
                {"field": "coverage", "label": "Coverage", "format": "percent"},
            ],
        },
        {
            "id": "endpoint_quality",
            "title": "Backup: exact five-seed quality anchors",
            "subtitle": "Loop 5 means; adjusted quantities use the corresponding baseline denominator.",
            "showDescription": True,
            "dataset": "endpoint_quality",
            "sourceId": "quality_summary",
            "density": "spacious",
            "columns": [
                {"field": "method", "label": "Condition", "type": "text"},
                {"field": "entry_raw", "label": "Entry raw", "format": "number"},
                {"field": "entry_adjusted", "label": "Entry adjusted", "format": "number"},
                {"field": "entry_coverage", "label": "Entry coverage", "format": "percent"},
                {"field": "sample_raw", "label": "Raw sample DQS", "format": "number"},
                {"field": "sample_adjusted", "label": "Adjusted sample DQS", "format": "number"},
                {"field": "sample_coverage", "label": "Sample coverage", "format": "percent"},
            ],
        },
        {
            "id": "statistics_terms",
            "title": "Backup: how to interpret the final statistical quantities",
            "subtitle": "Effect size and uncertainty lead; p-values are limited-resolution supporting evidence with five seeds.",
            "showDescription": True,
            "dataset": "statistics_terms",
            "sourceId": "statistics_plan",
            "density": "spacious",
            "columns": columns(
                [
                    ("term", "Quantity", "text"),
                    ("plain_meaning", "Plain meaning", "text"),
                    ("reporting_role", "Role in the update", "text"),
                ]
            ),
        },
    ]

    blocks = [
        {"id": "title", "type": "markdown", "body": f"# {REPORT_TITLE}"},
        {
            "id": "summary",
            "type": "markdown",
            "body": (
                "## Technical summary\n\n"
                "- Known-truth controls show that raw DQS tracks survivor cleanliness, while adjusted DQS tracks retained healthy mass.\n"
                "- Five-seed entry- and sample-level diagnostics expose the same denominator effect in the real-data pipeline.\n"
                "- Refinement shows a modest quality gain while retaining almost all evaluable samples; repeated removal looks cleaner among survivors but loses healthy mass against the original denominator.\n"
                "- These checks validate metric interpretation and robustness; they do not establish final downstream model superiority."
            ),
        },
        {
            "id": "headline_metrics",
            "type": "metric-strip",
            "cardIds": ["known_noise", "synthetic_seeds", "real_seeds", "refinement_sample_coverage"],
        },
        {
            "id": "definitions_heading",
            "type": "markdown",
            "body": "## Raw quality and retained healthy mass are complementary, not competing, views",
        },
        {"id": "definitions_table", "type": "table", "tableId": "metric_definitions"},
        {
            "id": "definitions_interpretation",
            "type": "markdown",
            "body": (
                "A simple example fixes the interpretation: if 100 entries contain 80 healthy labels and all 20 erroneous labels are removed, "
                "survivor cleanliness is 100%, but retained healthy mass against the original denominator remains 80%."
            ),
        },
        {
            "id": "synthetic_heading",
            "type": "markdown",
            "body": "## Known-truth controls reproduce the expected behavior across cleaning loops",
        },
        {"id": "synthetic_chart_block", "type": "chart", "chartId": "synthetic_known_truth"},
        {
            "id": "synthetic_interpretation",
            "type": "markdown",
            "sourceId": "synthetic_summary",
            "body": (
                "Correcting known errors raises both truth targets and both scores. Removing known errors raises raw DQS to 0.994 while adjusted DQS stays near the retained healthy mass of 0.880; random removal leaves survivor accuracy near 0.880 but lowers healthy mass to 0.774, which adjusted DQS exposes."
            ),
        },
        {
            "id": "sensitivity_heading",
            "type": "markdown",
            "body": "## DQS is an evidence-dependent diagnostic, not a direct truth measurement",
        },
        {"id": "sensitivity_chart_block", "type": "chart", "chartId": "oof_sensitivity"},
        {
            "id": "sensitivity_interpretation",
            "type": "markdown",
            "sourceId": "synthetic_errors",
            "body": (
                f"Across scenarios and matching targets, mean absolute error is {informative_mae:.3f} with informative OOF evidence and {weak_mae:.3f} with weak OOF evidence. "
                "The synthetic experiment therefore validates direction and denominator interpretation, not perfect calibration."
            ),
        },
        {
            "id": "entry_heading",
            "type": "markdown",
            "body": "## Entry-level adjustment separates cleaner survivors from retained healthy mass",
        },
        {"id": "entry_chart_block", "type": "chart", "chartId": "entry_quality"},
        {
            "id": "entry_interpretation",
            "type": "markdown",
            "sourceId": "quality_summary",
            "body": (
                "At Loop 5, removal reaches raw DQS 0.974 but adjusted DQS falls to 0.810 as valid-entry coverage declines to 83.1%. "
                "Refinement reaches raw/adjusted DQS 0.937/0.933 while retaining 99.6% entry coverage."
            ),
        },
        {
            "id": "sample_heading",
            "type": "markdown",
            "body": "## The sample-level diagnostic reaches the same denominator conclusion",
        },
        {"id": "sample_chart_block", "type": "chart", "chartId": "sample_quality"},
        {
            "id": "sample_interpretation",
            "type": "markdown",
            "sourceId": "quality_summary",
            "body": (
                "Removal raises raw sample DQS from 0.736 to 0.865, but 80.6% sample coverage leaves coverage-adjusted sample DQS at 0.698. "
                "Refinement reaches 0.756 raw sample DQS and 0.754 coverage-adjusted sample DQS while retaining 99.7% sample coverage."
            ),
        },
        {"id": "endpoint_table_block", "type": "table", "tableId": "endpoint_quality"},
        {
            "id": "limitations",
            "type": "markdown",
            "body": (
                "## Conclusions\n\n"
                "- Established: raw and adjusted metrics target different, interpretable quantities.\n"
                "- Established: the denominator effect is visible at both entry and sample levels across five seeds.\n"
                "- Limitation: DQS remains conditional on OOF evidence quality, and the sample metric is a strict model-based diagnostic rather than expert ground truth.\n"
                "- Not yet established: the final downstream Loop 5/8 effect size and uncertainty."
            ),
        },
        {
            "id": "next_steps",
            "type": "markdown",
            "body": (
                "## Recommended Actions\n\n"
                "- Finish the locked Loop 5/8 endpoint jobs.\n"
                "- Re-run the existing endpoint pipeline without changing comparisons after seeing results.\n"
                "- Report paired effect sizes, seed directions, hierarchical 95% intervals, exact paired p-values, and Holm-adjusted p-values.\n"
                "- Keep the synthetic and sample-level analyses as reliability evidence rather than treating them as downstream performance endpoints."
            ),
        },
        {"id": "synthetic_exact_block", "type": "table", "tableId": "synthetic_exact"},
        {"id": "statistics_terms_block", "type": "table", "tableId": "statistics_terms"},
    ]

    manifest = {
        "version": 1,
        "surface": "report",
        "title": REPORT_TITLE,
        "description": REPORT_SUBTITLE,
        "generatedAt": generated_at,
        "cards": cards,
        "charts": charts,
        "tables": tables,
        "sources": sources,
        "blocks": blocks,
    }
    return {
        "surface": "report",
        "manifest": manifest,
        "snapshot": {
            "version": 1,
            "generatedAt": generated_at,
            "status": "ready",
            "datasets": datasets,
        },
        "sources": sources,
        "package_info": {},
    }


def write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def run_portable_builder(artifact_path: Path, report_path: Path) -> None:
    subprocess.run(
        [
            "npm",
            "run",
            "report:deliver",
            "--",
            "--input",
            str(artifact_path),
            "--output",
            str(report_path),
        ],
        cwd=PLUGIN_ROOT,
        check=True,
    )


def png_data_uri(path: Path) -> str:
    return "data:image/png;base64," + base64.b64encode(path.read_bytes()).decode("ascii")


def adapt_report_for_slides(report_path: Path, output_path: Path) -> None:
    soup = BeautifulSoup(report_path.read_text(encoding="utf-8"), "html.parser")
    figures = soup.select("figure.portable-chart-summary")
    if len(figures) != len(CHART_IMAGES):
        raise ValueError(f"Expected {len(CHART_IMAGES)} portable charts, found {len(figures)}")

    alt_texts = [
        "Known-ground-truth validation of raw and coverage-adjusted DQS across four actions",
        "DQS mean absolute error under informative and weak OOF evidence",
        "Five-seed entry-level raw and adjusted DQS with individual seeds and standard deviations",
        "Five-seed raw and coverage-adjusted sample DQS with individual seeds and standard deviations",
    ]
    for figure, image_path, alt_text in zip(figures, CHART_IMAGES, alt_texts, strict=True):
        replacement = soup.new_tag("figure")
        replacement["class"] = ["figure", "deck-static-chart"]
        image_tag = soup.new_tag("img")
        image_tag["src"] = png_data_uri(image_path)
        image_tag["alt"] = alt_text
        replacement.append(image_tag)
        caption = soup.new_tag("figcaption")
        caption.string = alt_text
        replacement.append(caption)
        figure.replace_with(replacement)

    for svg in soup.find_all("svg"):
        svg.decompose()
    for source_inventory in soup.select(".portable-sources"):
        source_inventory.decompose()
    for tooltip in soup.select(".portable-markdown .portable-source-tooltip"):
        value = tooltip.select_one(".portable-source-value-text")
        tooltip.replace_with(value.get_text(" ", strip=True) if value else tooltip.get_text(" ", strip=True))

    h1 = soup.find("h1")
    if h1 is None:
        raise ValueError("Portable report has no h1 title")
    eyebrow = soup.new_tag("p")
    eyebrow["class"] = ["eyebrow"]
    eyebrow.string = "Current evaluation evidence | 15 July 2026"
    h1.insert_before(eyebrow)
    lede = soup.new_tag("p")
    lede["class"] = ["lede"]
    lede.string = REPORT_SUBTITLE
    h1.insert_after(lede)

    style = soup.new_tag("style")
    style.string = (
        ".deck-static-chart img{display:block;width:100%;height:auto;}"
        ".deck-static-chart figcaption{color:#657084;font-size:.85rem;}"
    )
    if soup.head:
        soup.head.append(style)
    output_path.write_text(str(soup), encoding="utf-8")


def run_slides_helper(report_path: Path) -> None:
    subprocess.run(
        [sys.executable, str(SLIDES_HELPER), str(report_path), "--out-dir", str(OUTPUT_DIR)],
        check=True,
    )


def note_for_title(title: str) -> str:
    lowered = title.lower()
    if "reliability checks" in lowered:
        return (
            "这次只汇报已经完成的 reliability checks，不提前解释还在运行的最终 endpoint。"
            "主线是三个问题：指标是否跟随已知真值、提升是否只是分母缩小、sample level 是否得到同样结论。"
        )
    if "technical summary" in lowered or "executive summary" in lowered or "completed checks already" in lowered:
        return (
            "先给结论：synthetic、entry-level 和 sample-level 三种证据互相吻合。"
            "它们说明 raw 与 adjusted 回答不同问题，但还不能替代最终下游模型效果。"
        )
    if "paired metrics" in lowered or "complementary" in lowered or "different denominator questions" in lowered:
        return (
            "用 100 个 entries 的小例子解释：原来 80 个健康、20 个错误；把 20 个错误全删掉后，"
            "幸存者质量是 100%，但原始分母下健康信息仍只有 80%。raw 看前者，adjusted 看后者。"
        )
    if "known-truth" in lowered or "synthetic truth" in lowered:
        return (
            "四个 panel 是正反 controls。修正错误时两种分数都上升；删除错误时 raw 接近 1，"
            "adjusted 保持在剩余健康质量附近；随机删除只让 adjusted 下降；故意破坏干净标签时两者都下降。"
            "补充 reliability：informative OOF 下 matching-target MAE 是 0.019，weak OOF 下是 0.055。"
        )
    if "dqs reliability" in lowered or "evidence-dependent" in lowered or "informative and weak" in lowered:
        return (
            "这里主动说明限制：DQS 不是直接真值。OOF evidence 好时，平均绝对误差约 1.9 个百分点；"
            "evidence 弱时约 5.5 个百分点，所以我们验证的是方向和分母解释，不是完美校准。"
        )
    if "entry-level" in lowered or "healthy entry mass" in lowered:
        return (
            "左图看当前留下来的 entries，右图把分母固定回 baseline。Removal 左边持续上升，"
            "但右边下降到 0.810；refinement raw 和 adjusted 都略高于 baseline，并保留约 99.6% entry coverage。"
        )
    if "sample-level" in lowered or "sample level" in lowered:
        return (
            "sample 定义很严格：任一有效疾病标签被 flagged，整个 sample 就算 problematic。"
            "Removal 的 raw sample DQS 升到 0.865，但 coverage-adjusted sample DQS 降到 0.698；"
            "refinement 的 coverage-adjusted sample DQS 为 0.754，并保留约 99.7% sample coverage。"
        )
    if "exact five-seed quality" in lowered:
        return "这张是精确数字查表页，主讲时可以快速带过；老师追问分母或小数时再使用。"
    if "completed evidence supports" in lowered or "does and does not" in lowered or "next reporting" in lowered:
        return (
            "结论要守住边界：目前已经验证 metric interpretation 和 denominator effect。"
            "最终 performance effect、CI 和 p-values 等作业完成后再按锁定 pipeline 一次性补上。"
        )
    if "exact loop 5 synthetic" in lowered:
        return "这是 synthetic 的精确 Loop 5 backup，用来回答图上每条线对应哪个真实目标以及误差有多大。"
    if "statistical quantities" in lowered or "final statistics" in lowered:
        return (
            "p-value 只回答无系统效果时当前差异有多意外，不是方法有效的概率。"
            "五个 paired seeds 的双侧 exact p 最小是 0.0625，所以主结论应先看 effect size、CI 和 seed direction。"
        )
    return "按页面标题先讲结论，再指出证据和适用边界；不延伸到尚未完成的最终 endpoint。"


def preferred_title(item: dict[str, Any]) -> str:
    kind = str(item.get("kind") or "")
    if kind == "cover":
        return "Reliability checks completed"
    if kind == "executive_summary":
        return "The completed checks already answer the denominator question"
    if kind == "chart_evidence":
        return {
            1: "Known-truth controls reproduce the expected score behavior",
            2: "DQS reliability depends on OOF evidence quality",
            3: "Removal raises raw DQS but lowers coverage-adjusted DQS",
            4: "Sample DQS shows the same denominator effect",
        }.get(int(item.get("chart_index") or 0), str(item.get("title") or ""))
    if kind == "table":
        return {
            1: "Raw and coverage-adjusted DQS use different denominators",
            2: "Backup: exact five-seed quality anchors",
            3: "Backup: exact Loop 5 synthetic targets and scores",
            4: "Backup: how to interpret the final statistics",
        }.get(int(item.get("table_index") or 0), str(item.get("title") or ""))
    if kind == "conclusions_implications":
        return "What the completed evidence supports, and what remains pending"
    return str(item.get("title") or "")


def slide_sort_key(item: dict[str, Any], original_index: int) -> tuple[int, int]:
    kind = str(item.get("kind") or "")
    if kind == "cover":
        return (0, original_index)
    if kind == "executive_summary":
        return (1, original_index)
    if kind == "table" and int(item.get("table_index") or 0) == 1:
        return (2, original_index)
    if kind == "chart_evidence":
        return (2 + int(item.get("chart_index") or 0), original_index)
    if kind == "conclusions_implications":
        return (7, original_index)
    if kind == "table":
        return (7 + int(item.get("table_index") or 0), original_index)
    return (20, original_index)


def replace_slide_title(slide: Any, old_title: str, new_title: str) -> None:
    if not new_title or new_title == old_title:
        return
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if " ".join(shape.text.split()) != " ".join(old_title.split()):
            continue
        paragraphs = shape.text_frame.paragraphs
        if not paragraphs:
            continue
        first = paragraphs[0]
        if first.runs:
            first.runs[0].text = new_title
            for run in first.runs[1:]:
                run.text = ""
        else:
            first.text = new_title
        for paragraph in paragraphs[1:]:
            for run in paragraph.runs:
                run.text = ""
        return
    raise ValueError(f"Could not find slide title: {old_title}")


def replace_shape_text_containing(slide: Any, needle: str, replacement: str) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or needle not in shape.text:
            continue
        paragraphs = shape.text_frame.paragraphs
        first = paragraphs[0]
        if first.runs:
            first.runs[0].text = replacement
            for run in first.runs[1:]:
                run.text = ""
        else:
            first.text = replacement
        for paragraph in paragraphs[1:]:
            for run in paragraph.runs:
                run.text = ""
        return
    raise ValueError(f"Could not find slide text containing: {needle}")


def finalize_deck(raw_deck: Path, output_deck: Path, plan_path: Path) -> int:
    presentation = Presentation(raw_deck)
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    if len(presentation.slides) != len(plan):
        raise ValueError(f"Slide/plan mismatch: {len(presentation.slides)} vs {len(plan)}")

    keep = [
        index
        for index, item in enumerate(plan)
        if str(item.get("kind") or "") != "executive_summary"
        and not (
            str(item.get("kind") or "") == "chart_evidence"
            and int(item.get("chart_index") or 0) == 2
        )
    ]
    order = sorted(keep, key=lambda index: slide_sort_key(plan[index], index))
    slide_ids = list(presentation.slides._sldIdLst)
    for slide_id in slide_ids:
        presentation.slides._sldIdLst.remove(slide_id)
    for index in order:
        presentation.slides._sldIdLst.append(slide_ids[index])
    plan = [plan[index] for index in order]

    notes_added = 0
    for slide_number, (slide, item) in enumerate(zip(presentation.slides, plan, strict=True), start=1):
        old_title = str(item.get("title") or "")
        title = preferred_title(item)
        replace_slide_title(slide, old_title, title)
        if str(item.get("kind") or "") == "chart_evidence" and int(item.get("chart_index") or 0) == 1:
            replace_shape_text_containing(
                slide,
                "Removing known errors raises",
                "Raw DQS follows survivor cleanliness; adjusted DQS follows healthy mass against the original denominator. Matching-target MAE is 0.019 with informative OOF evidence and 0.055 with weak evidence.",
            )
        item["title"] = title
        item["slide_number"] = slide_number
        item["section"] = "backup" if title.lower().startswith("backup:") else "main"
        slide.notes_slide.notes_text_frame.text = note_for_title(title)
        notes_added += 1
    presentation.save(output_deck)
    write_json(plan_path, plan)
    return notes_added


def main() -> int:
    require_files(
        [
            SYNTH_SUMMARY,
            SYNTH_ERRORS,
            SYNTH_METADATA,
            REAL_SUMMARY,
            REAL_METADATA,
            SLIDES_HELPER,
            *CHART_IMAGES,
        ]
    )
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    synthetic = pd.read_csv(SYNTH_SUMMARY)
    errors = pd.read_csv(SYNTH_ERRORS)
    real = pd.read_csv(REAL_SUMMARY)
    synth_meta = json.loads(SYNTH_METADATA.read_text(encoding="utf-8"))
    real_meta = json.loads(REAL_METADATA.read_text(encoding="utf-8"))
    validate_inputs(synthetic, errors, real, synth_meta, real_meta)

    artifact_path = OUTPUT_DIR / "artifact.json"
    report_path = OUTPUT_DIR / "report.html"
    conversion_report_path = OUTPUT_DIR / "report_for_slides.html"
    final_deck_path = OUTPUT_DIR / "reliability_validation_update_20260715.pptx"

    write_json(artifact_path, build_artifact(synthetic, errors, real, synth_meta))
    run_portable_builder(artifact_path, report_path)
    adapt_report_for_slides(report_path, conversion_report_path)
    run_slides_helper(conversion_report_path)

    preflight = json.loads((OUTPUT_DIR / "preflight_checks.json").read_text(encoding="utf-8"))
    if preflight.get("status") != "passed":
        raise RuntimeError(f"Slides preflight failed: {preflight}")
    notes_added = finalize_deck(
        OUTPUT_DIR / "deck.pptx",
        final_deck_path,
        OUTPUT_DIR / "deck_plan.json",
    )
    shutil.copy2(final_deck_path, OUTPUT_DIR / "deck_with_notes.pptx")

    summary = {
        "status": "passed",
        "artifact": str(artifact_path),
        "report": str(report_path),
        "conversion_report": str(conversion_report_path),
        "raw_deck": str(OUTPUT_DIR / "deck.pptx"),
        "final_deck": str(final_deck_path),
        "slides": len(Presentation(final_deck_path).slides),
        "main_slides": 6,
        "backup_slides": 3,
        "speaker_notes_added": notes_added,
        "google_slides_import": "pending connector availability",
    }
    write_json(OUTPUT_DIR / "delivery_summary.json", summary)
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
