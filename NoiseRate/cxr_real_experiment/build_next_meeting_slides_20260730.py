#!/usr/bin/env python3
"""Build the next-meeting PowerPoint from the canonical Chinese slide outline."""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
SOURCE_MD = HERE / "next_meeting_slides_outline_中文.md"
OUTPUT_DIR = HERE / "next_meeting_slides_20260730"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_PPTX = OUTPUT_DIR / "next_meeting_update_20260730.pptx"

PAGE2_TRAJECTORY = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "loop_oscillation_forensics_4seed_20260727/"
    "three_label_subset_comparison_mean_trajectory.csv"
)
PAGE2_BASELINE_ROOT = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "results_mobilenetv3_noes50_clean_3seed/20260707_123320"
)
PAGE2_SEEDS = (13, 42, 97, 123)
PAGE2_DRIVER_TRIAD = {
    "Atelectasis",
    "Lung Opacity",
    "Pneumothorax",
}

STANDARD_LOOP_METRICS = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "results_mobilenetv3_own_top20_binary_v5_llm_refine_5seed/"
    "20260714_065539/seed_13/llm_refine/loop_metrics.csv"
)
UNSEEN_LOOP_METRICS = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "results_mobilenetv3_own_top20_binary_v5_refine_unseen_extension/"
    "20260726_seed42_to_loop15/seed_42/llm_refine/loop_metrics.csv"
)
UNSEEN_DQS = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "results_mobilenetv3_own_top20_binary_v5_refine_unseen_extension/"
    "20260726_seed42_to_loop15/analysis/"
    "seed42_loop01_15_sample_dqs_frozen_vs_dynamic.csv"
)
STANDARD_DQS = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "results_mobilenetv3_own_top20_binary_v5_llm_refine_5seed/"
    "20260714_065539/seed_13/llm_refine/analysis/"
    "seed13_loop01_15_sample_dqs_frozen.csv"
)

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT_FAMILY = "Times New Roman"
PLOT_FONT_FAMILY = "Times New Roman"
PLOT_FONT_FALLBACK = "Liberation Serif"


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


INK = rgb("#18212F")
MUTED = rgb("#667085")
LIGHT_MUTED = rgb("#98A2B3")
BLUE = rgb("#0077BB")
BLUE_DARK = rgb("#24557A")
BLUE_LIGHT = rgb("#E8F3F9")
ORANGE = rgb("#EE7733")
ORANGE_LIGHT = rgb("#FCEEE5")
TEAL = rgb("#009988")
TEAL_LIGHT = rgb("#E5F4F1")
RED = rgb("#CC3311")
RED_LIGHT = rgb("#FAECE8")
AMBER = rgb("#C57A00")
AMBER_LIGHT = rgb("#FCF2DE")
GREY = rgb("#6B7280")
GREY_LIGHT = rgb("#EEF1F5")
GRID = rgb("#D9DEE7")
PANEL = rgb("#F7F8FA")
WHITE = rgb("#FFFFFF")

HEX_INK = "#18212F"
HEX_MUTED = "#667085"
HEX_BLUE = "#0077BB"
HEX_BLUE_DARK = "#24557A"
HEX_ORANGE = "#EE7733"
HEX_TEAL = "#009988"
HEX_RED = "#CC3311"
HEX_GREY = "#6B7280"
HEX_GRID = "#D9DEE7"
HEX_LIGHT_GREY = "#B8C0CC"

TITLES = [
    "From Open Questions to Validation",
    "What Drives the AUROC Oscillation?",
    "Can Additional Seeds Strengthen the Statistical Evidence?",
    "More Loops and Skipping Repeated Entries",
    "Why AUROC and DQS Are Not Direct Module Validation",
    "Med-PaLM: Can the LLM Correct a Flagged Entry?",
    "REFLACX: Can Confident Learning Find Real Disagreements?",
    "Current Interpretation and Next Steps",
]
NOTE_PAGES = set(range(1, 9))


def add_text(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    *,
    size: float = 14,
    text_color: RGBColor = INK,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.0,
    name: str | None = None,
    font: str = FONT_FAMILY,
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = text_color
    return shape


def add_multiline(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    lines: list[dict[str, Any]],
    *,
    fill: RGBColor | None = None,
    line: RGBColor | None = None,
    margin: float = 0.12,
    name: str | None = None,
):
    if fill is not None:
        add_box(
            slide,
            x,
            y,
            width,
            height,
            fill=fill,
            line=line or fill,
            name=f"{name}-box" if name else None,
        )
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    for index, item in enumerate(lines):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.alignment = item.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(item.get("space_before", 0))
        p.space_after = Pt(item.get("space_after", 4))
        p.level = int(item.get("level", 0))
        if item.get("bullet"):
            p.text = f"• {item['text']}"
            if p.runs:
                run = p.runs[0]
            else:
                run = p.add_run()
                run.text = p.text
        else:
            run = p.add_run()
            run.text = item["text"]
        run.font.name = item.get("font", FONT_FAMILY)
        run.font.size = Pt(item.get("size", 13))
        run.font.bold = item.get("bold", False)
        run.font.color.rgb = item.get("color", INK)
    return shape


def add_box(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    fill: RGBColor = WHITE,
    line: RGBColor = GRID,
    line_width: float = 0.8,
    name: str | None = None,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_circle(
    slide,
    x: float,
    y: float,
    diameter: float,
    *,
    fill: RGBColor,
    line: RGBColor | None = None,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x),
        Inches(y),
        Inches(diameter),
        Inches(diameter),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_rule(
    slide,
    x: float,
    y: float,
    width: float,
    *,
    line_color: RGBColor = GRID,
    line_width: float = 1.0,
):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x),
        Inches(y),
        Inches(x + width),
        Inches(y),
    )
    line.line.color.rgb = line_color
    line.line.width = Pt(line_width)
    return line


def add_picture_contain(
    slide,
    path: Path,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    name: str | None = None,
):
    with Image.open(path) as image:
        image_ratio = image.width / image.height
    box_ratio = width / height
    if image_ratio >= box_ratio:
        render_w = width
        render_h = width / image_ratio
        render_x = x
        render_y = y + (height - render_h) / 2
    else:
        render_h = height
        render_w = height * image_ratio
        render_x = x + (width - render_w) / 2
        render_y = y
    shape = slide.shapes.add_picture(
        str(path),
        Inches(render_x),
        Inches(render_y),
        width=Inches(render_w),
        height=Inches(render_h),
    )
    if name:
        shape.name = name
    return shape


def new_slide(prs: Presentation, page: int, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = WHITE
    add_text(
        slide,
        0.55,
        0.28,
        11.9,
        0.55,
        title,
        size=25,
        bold=True,
        name="slide-title",
    )
    add_text(
        slide,
        12.38,
        0.34,
        0.38,
        0.28,
        f"{page:02d}",
        size=10,
        text_color=LIGHT_MUTED,
        align=PP_ALIGN.RIGHT,
        name="page-number",
    )
    add_rule(slide, 0.55, 0.92, 12.2, line_color=GRID, line_width=0.7)
    return slide


def add_section_label(slide, x: float, y: float, text: str, color_value: RGBColor):
    width = min(4.8, SLIDE_W - x - 0.55)
    add_text(
        slide,
        x,
        y,
        width,
        0.28,
        text.upper(),
        size=10,
        text_color=color_value,
        bold=True,
    )


def add_takeaway(
    slide,
    text: str,
    *,
    color_value: RGBColor = BLUE_DARK,
    y: float = 6.70,
):
    add_rule(slide, 0.62, y - 0.10, 12.08, line_color=color_value, line_width=1.8)
    add_text(
        slide,
        0.72,
        y,
        11.9,
        0.42,
        text,
        size=13.2,
        text_color=color_value,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_source(slide, text: str):
    add_text(
        slide,
        0.62,
        7.26,
        12.08,
        0.16,
        text,
        size=7.2,
        text_color=LIGHT_MUTED,
    )


def parse_notes() -> dict[int, str]:
    text = SOURCE_MD.read_text(encoding="utf-8")
    page_sections = re.findall(
        r"^## Page (\d+) - .*?\n(.*?)(?=^---\n\n## Page |\Z)",
        text,
        flags=re.MULTILINE | re.DOTALL,
    )
    notes: dict[int, str] = {}
    for page_text, section in page_sections:
        page = int(page_text)
        if page not in NOTE_PAGES:
            continue
        match = re.search(
            r"^### 讲稿\s*\n\n(.*?)(?=^### |\Z)",
            section,
            flags=re.MULTILINE | re.DOTALL,
        )
        if not match:
            raise ValueError(f"Missing canonical speaker note for Page {page}")
        notes[page] = match.group(1).strip()
    if set(notes) != NOTE_PAGES:
        raise ValueError(f"Unexpected note pages: {sorted(notes)}")
    return notes


def configure_plot_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": [
                PLOT_FONT_FAMILY,
                PLOT_FONT_FALLBACK,
                "Nimbus Roman",
                "DejaVu Serif",
            ],
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.labelcolor": HEX_MUTED,
            "xtick.color": HEX_MUTED,
            "ytick.color": HEX_MUTED,
        }
    )


def page2_baseline_means() -> dict[str, float]:
    records: list[dict[str, Any]] = []
    for seed in PAGE2_SEEDS:
        summary = pd.read_csv(
            PAGE2_BASELINE_ROOT
            / f"seed_{seed}"
            / "baseline_no_clean"
            / "test_study_auroc_summary.csv"
        )
        specifications = {
            "all_12_labels": pd.Series(True, index=summary.index),
            "driver_triad_only": summary["label_name"].isin(PAGE2_DRIVER_TRIAD),
            "remaining_9_labels": ~summary["label_name"].isin(PAGE2_DRIVER_TRIAD),
        }
        for specification, mask in specifications.items():
            frame = summary[
                mask
                & summary["study_auroc_binary"].notna()
                & (summary["study_valid_count"] > 0)
            ]
            records.append(
                {
                    "specification": specification,
                    "seed": seed,
                    "baseline_auroc": float(
                        np.average(
                            frame["study_auroc_binary"],
                            weights=frame["study_valid_count"],
                        )
                    ),
                }
            )
    baseline = pd.DataFrame(records)
    return (
        baseline.groupby("specification")["baseline_auroc"].mean().to_dict()
    )


def page4_baseline_values() -> dict[int, float]:
    values: dict[int, float] = {}
    for seed in (13, 42):
        summary = pd.read_csv(
            PAGE2_BASELINE_ROOT
            / f"seed_{seed}"
            / "baseline_no_clean"
            / "test_study_auroc_summary.csv"
        )
        frame = summary[
            summary["study_auroc_binary"].notna()
            & (summary["study_valid_count"] > 0)
        ]
        values[seed] = float(
            np.average(
                frame["study_auroc_binary"],
                weights=frame["study_valid_count"],
            )
        )
    return values


def plot_page2(path: Path) -> None:
    trajectory = pd.read_csv(PAGE2_TRAJECTORY)
    trajectory = trajectory[trajectory["loop"].between(4, 8)].copy()
    baselines = page2_baseline_means()
    loops = np.array([4, 5, 6, 7, 8])

    def series(specification: str) -> tuple[np.ndarray, np.ndarray]:
        frame = (
            trajectory[trajectory["specification"] == specification]
            .sort_values("loop")
            .set_index("loop")
            .loc[loops]
        )
        return frame["mean"].to_numpy(), frame["std"].to_numpy()

    all_labels, all_labels_sd = series("all_12_labels")
    triad, triad_sd = series("driver_triad_only")
    remaining, remaining_sd = series("remaining_9_labels")

    fig, ax = plt.subplots(figsize=(10.6, 5.7), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.errorbar(
        loops,
        all_labels,
        yerr=all_labels_sd,
        color=HEX_INK,
        marker="o",
        linewidth=2.6,
        markersize=6,
        capsize=3.5,
        elinewidth=1.35,
        label="All 12 labels",
    )
    ax.errorbar(
        loops,
        triad,
        yerr=triad_sd,
        color=HEX_RED,
        marker="s",
        linewidth=2.6,
        markersize=6,
        capsize=3.5,
        elinewidth=1.35,
        label="Atelectasis + Lung Opacity + Pneumothorax",
    )
    ax.errorbar(
        loops,
        remaining,
        yerr=remaining_sd,
        color=HEX_GREY,
        marker="^",
        linewidth=2.3,
        markersize=6,
        capsize=3.5,
        elinewidth=1.35,
        label="Remaining 9 labels",
    )
    baseline_styles = [
        ("all_12_labels", HEX_INK),
        ("driver_triad_only", HEX_RED),
        ("remaining_9_labels", HEX_GREY),
    ]
    for specification, color in baseline_styles:
        value = baselines[specification]
        ax.axhline(
            value,
            color=color,
            linestyle=(0, (5, 4)),
            linewidth=1.55,
            alpha=0.82,
            zorder=0,
        )
        ax.text(
            8.12,
            value,
            f"baseline {value:.3f}",
            color=color,
            fontsize=8.1,
            va="bottom",
            ha="left",
        )
    ax.text(
        0.01,
        1.02,
        "Mean ± SD across seeds 13, 42, 97 and 123; dashed lines show matched baseline means",
        transform=ax.transAxes,
        color=HEX_MUTED,
        fontsize=8.8,
        va="bottom",
    )
    ax.set_xlim(3.8, 8.72)
    ax.set_ylim(0.43, 0.83)
    ax.set_xticks(loops)
    ax.set_xlabel("Loop")
    ax.set_ylabel("Study-weighted AUROC")
    ax.grid(axis="y", color=HEX_GRID, linewidth=0.8)
    legend = ax.legend(
        loc="lower center",
        bbox_to_anchor=(0.5, -0.26),
        ncol=1,
        frameon=False,
        fontsize=8.7,
    )
    for text in legend.get_texts():
        text.set_color(HEX_INK)
    fig.tight_layout(rect=(0.02, 0.08, 0.99, 0.97))
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def plot_page4_auroc(path: Path) -> None:
    standard = pd.read_csv(STANDARD_LOOP_METRICS).sort_values("loop_id")
    unseen = pd.read_csv(UNSEEN_LOOP_METRICS).sort_values("loop_id")
    baselines = page4_baseline_values()
    fig, ax = plt.subplots(figsize=(11.4, 3.15), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.plot(
        standard["loop_id"],
        standard["test_study_weighted_auroc"],
        color=HEX_BLUE,
        marker="o",
        linewidth=2.2,
        markersize=4.5,
        label="Standard refinement (seed 13)",
    )
    ax.plot(
        unseen["loop_id"],
        unseen["test_study_weighted_auroc"],
        color=HEX_ORANGE,
        marker="s",
        linewidth=2.2,
        markersize=4.3,
        label="Unseen-entry refinement (seed 42)",
    )
    for seed, color in ((13, HEX_BLUE), (42, HEX_ORANGE)):
        value = baselines[seed]
        ax.axhline(
            value,
            color=color,
            linestyle=(0, (5, 4)),
            linewidth=1.45,
            alpha=0.78,
            zorder=0,
        )
        ax.text(
            15.18,
            value,
            f"seed {seed} baseline {value:.3f}",
            color=color,
            fontsize=8.0,
            ha="right",
            va="bottom",
        )
    ax.axvline(8.5, color=HEX_LIGHT_GREY, linestyle="--", linewidth=1.4)
    ax.text(
        8.38,
        0.800,
        "Unseen-entry selection begins",
        color=HEX_MUTED,
        fontsize=8.3,
        ha="right",
        va="top",
    )
    unseen_loop10 = unseen.loc[unseen["loop_id"] == 10].iloc[0]
    ax.annotate(
        f"Loop 10: {unseen_loop10['test_study_weighted_auroc']:.3f}",
        xy=(
            unseen_loop10["loop_id"],
            unseen_loop10["test_study_weighted_auroc"],
        ),
        xytext=(16, 7),
        textcoords="offset points",
        ha="left",
        va="bottom",
        color=HEX_ORANGE,
        fontsize=8.8,
        fontweight="bold",
        bbox={
            "boxstyle": "round,pad=0.22",
            "facecolor": "white",
            "edgecolor": HEX_ORANGE,
            "linewidth": 0.8,
            "alpha": 0.94,
        },
        arrowprops={
            "arrowstyle": "-",
            "color": HEX_ORANGE,
            "linewidth": 0.9,
        },
    )
    ax.set_xlim(0.7, 15.3)
    ax.set_ylim(0.69, 0.805)
    ax.set_xticks(range(1, 16))
    ax.set_ylabel("Study-weighted AUROC")
    ax.grid(axis="y", color=HEX_GRID, linewidth=0.8)
    ax.legend(loc="lower left", frameon=False, fontsize=8.5, ncol=2)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout(pad=0.8)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def plot_page4_dqs(path: Path) -> None:
    standard = pd.read_csv(STANDARD_DQS).sort_values("state_after_loop")
    unseen = pd.read_csv(UNSEEN_DQS)
    unseen = unseen[unseen["mode"] == "frozen_initial_oof_post_action"].sort_values(
        "state_after_loop"
    )
    fig, ax = plt.subplots(figsize=(11.4, 3.05), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.plot(
        standard["state_after_loop"],
        standard["adjusted_sample_dqs"],
        color=HEX_GREY,
        marker="o",
        linewidth=2.3,
        markersize=4.5,
        label="Standard repeated-review (seed 13)",
    )
    ax.plot(
        unseen["state_after_loop"],
        unseen["adjusted_sample_dqs"],
        color=HEX_ORANGE,
        marker="s",
        linewidth=2.3,
        markersize=4.5,
        label="Unseen-entry run (seed 42; skip from Loop 9)",
    )
    ax.axvline(8.5, color=HEX_LIGHT_GREY, linestyle="--", linewidth=1.4)
    ax.text(
        8.62,
        0.829,
        "Selection policy diverges",
        color=HEX_MUTED,
        fontsize=8.1,
        va="top",
    )
    final_points = [
        (standard.iloc[-1], HEX_GREY, -12),
        (unseen.iloc[-1], HEX_ORANGE, 9),
    ]
    for row, color, vertical_offset in final_points:
        ax.annotate(
            f"{row['adjusted_sample_dqs']:.3f}",
            xy=(row["state_after_loop"], row["adjusted_sample_dqs"]),
            xytext=(0, vertical_offset),
            textcoords="offset points",
            ha="center",
            va="bottom" if vertical_offset >= 0 else "top",
            color=color,
            fontsize=8.7,
            fontweight="bold",
        )
    ax.set_xlim(-0.3, 15.3)
    ax.set_ylim(0.725, 0.833)
    ax.set_xticks(range(0, 16))
    ax.set_xlabel("Completed loop")
    ax.set_ylabel("Coverage-adjusted sample DQS")
    ax.grid(axis="y", color=HEX_GRID, linewidth=0.8)
    ax.legend(loc="upper left", frameon=False, fontsize=8.5, ncol=2)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout(pad=0.8)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def slide_1(prs: Presentation, note: str) -> None:
    slide = new_slide(prs, 1, TITLES[0])
    add_section_label(slide, 0.68, 1.12, "Questions from the last meeting", BLUE)
    questions = [
        "What drives the late-loop AUROC oscillation?",
        "Are the results repeatable across more random seeds?",
        "Do more loops converge, and are repeated reviews wasting the review budget?",
        "Improved AUROC and DQS do not directly verify detection or correction.",
    ]
    for index, question in enumerate(questions, start=1):
        y = 1.52 + (index - 1) * 1.16
        add_circle(slide, 0.72, y + 0.05, 0.36, fill=BLUE_LIGHT, line=BLUE)
        add_text(
            slide,
            0.72,
            y + 0.075,
            0.36,
            0.25,
            str(index),
            size=10.5,
            text_color=BLUE_DARK,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            1.25,
            y,
            4.55,
            0.72,
            question,
            size=14.4 if index < 4 else 13.6,
            bold=index == 4,
            text_color=INK if index < 4 else BLUE_DARK,
            valign=MSO_ANCHOR.MIDDLE,
        )

    add_rule(slide, 6.15, 1.15, 0.01, line_color=GRID, line_width=1.0)
    divider = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(6.15),
        Inches(1.15),
        Inches(6.15),
        Inches(6.55),
    )
    divider.line.color.rgb = GRID
    divider.line.width = Pt(1)

    add_section_label(slide, 6.55, 1.12, "What we did this week", TEAL)
    work = [
        ("DONE", "Investigated Loop 4-8 AUROC oscillation", TEAL, TEAL_LIGHT),
        (
            "RUNNING",
            "Added two seeds for the fixed Loop 8 comparison",
            AMBER,
            AMBER_LIGHT,
        ),
        ("DONE", "Extended refinement to 15 loops", TEAL, TEAL_LIGHT),
        ("DONE", "Tested skipping previously reviewed entries", TEAL, TEAL_LIGHT),
        ("DONE", "Med-PaLM validation of LLM correction", TEAL, TEAL_LIGHT),
        ("PENDING", "REFLACX validation of CL detection", GREY, GREY_LIGHT),
    ]
    for index, (status, task, status_color, fill_color) in enumerate(work):
        y = 1.48 + index * 0.78
        add_box(
            slide,
            6.55,
            y,
            1.15,
            0.38,
            fill=fill_color,
            line=fill_color,
        )
        add_text(
            slide,
            6.55,
            y + 0.02,
            1.15,
            0.31,
            status,
            size=9.1,
            bold=True,
            text_color=status_color,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            7.95,
            y - 0.01,
            4.35,
            0.44,
            task,
            size=14.0,
            valign=MSO_ANCHOR.MIDDLE,
        )
        if index < len(work) - 1:
            add_rule(slide, 7.95, y + 0.56, 4.35, line_color=GREY_LIGHT, line_width=0.6)
    slide.notes_slide.notes_text_frame.text = note


def slide_2(prs: Presentation, note: str, chart: Path) -> None:
    slide = new_slide(prs, 2, TITLES[1])
    add_text(
        slide,
        0.68,
        1.09,
        7.8,
        0.3,
        "Loop 4-8 study-weighted AUROC",
        size=11,
        text_color=MUTED,
        bold=True,
    )
    add_picture_contain(slide, chart, 0.58, 1.35, 8.15, 4.9, name="page2-chart")
    add_text(
        slide,
        0.78,
        6.28,
        7.75,
        0.24,
        "Three labels driving most of the movement: Atelectasis · Lung Opacity · Pneumothorax",
        size=10.2,
        text_color=RED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_section_label(slide, 9.0, 1.12, "Evidence", RED)
    callouts = [
        (
            "01",
            "The three driver labels oscillate much more strongly.",
            RED_LIGHT,
            RED,
        ),
        (
            "02",
            "Remaining 9 stay stable and +0.022 to +0.037 above baseline.",
            GREY_LIGHT,
            GREY,
        ),
        (
            "03",
            "Only 11 immediate reversals among 12,730 state changes.",
            BLUE_LIGHT,
            BLUE,
        ),
    ]
    for index, (number, text, fill, accent) in enumerate(callouts):
        y = 1.55 + index * 1.20
        add_box(slide, 9.0, y, 3.7, 0.96, fill=fill, line=fill)
        add_text(
            slide,
            9.18,
            y + 0.16,
            0.42,
            0.28,
            number,
            size=10,
            bold=True,
            text_color=accent,
        )
        add_text(
            slide,
            9.72,
            y + 0.12,
            2.72,
            0.64,
            text,
            size=12.5,
            bold=index == 0,
            valign=MSO_ANCHOR.MIDDLE,
        )
    add_multiline(
        slide,
        9.0,
        5.27,
        3.7,
        0.88,
        [
            {
                "text": "SUPPORT SENSITIVITY",
                "size": 9.3,
                "bold": True,
                "color": BLUE,
                "space_after": 3,
            },
            {
                "text": "Excluding labels with <5 minority-class test studies reduces amplitude; all four adjacent directions remain.",
                "size": 10.4,
                "color": MUTED,
            },
        ],
        fill=PANEL,
        line=GRID,
    )
    add_takeaway(
        slide,
        "Most labels retain a stable gain over baseline; aggregate oscillation is concentrated in a small high-variance subset.",
        color_value=RED,
    )
    add_source(
        slide,
        "Mean ± SD across four seeds; dashed lines show matched baseline means. Driver labels were identified post hoc.",
    )
    slide.notes_slide.notes_text_frame.text = note


def placeholder_slide(
    prs: Presentation,
    page: int,
    title: str,
    label: str,
    note: str,
) -> None:
    slide = new_slide(prs, page, title)
    add_box(
        slide,
        1.55,
        2.08,
        10.25,
        3.18,
        fill=PANEL,
        line=GRID,
        line_width=1.2,
        name="result-placeholder",
    )
    add_text(
        slide,
        1.55,
        2.77,
        10.25,
        0.38,
        label.upper(),
        size=11,
        text_color=LIGHT_MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    slide.notes_slide.notes_text_frame.text = note
    add_text(
        slide,
        1.55,
        3.25,
        10.25,
        0.70,
        "Results pending",
        size=24,
        text_color=GREY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def slide_4(prs: Presentation, note: str, auroc_chart: Path, dqs_chart: Path) -> None:
    slide = new_slide(prs, 4, TITLES[3])
    add_text(
        slide,
        0.62,
        1.07,
        8.95,
        0.24,
        "Downstream performance",
        size=10.2,
        text_color=MUTED,
        bold=True,
    )
    add_picture_contain(
        slide,
        auroc_chart,
        0.55,
        1.25,
        9.0,
        2.42,
        name="page4-auroc-chart",
    )
    add_text(
        slide,
        0.62,
        3.64,
        8.95,
        0.24,
        "Sample-level DQS: within-run change under a fixed reference",
        size=10.2,
        text_color=MUTED,
        bold=True,
    )
    add_picture_contain(
        slide,
        dqs_chart,
        0.55,
        3.82,
        9.0,
        2.45,
        name="page4-dqs-chart",
    )

    add_section_label(slide, 9.82, 1.12, "Review efficiency", ORANGE)
    metrics = [
        ("4.3%", "Old method\naction yield", GREY, GREY_LIGHT),
        ("29.5%", "Unseen-entry\naction yield", ORANGE, ORANGE_LIGHT),
        ("0", "Historical-entry overlap\n(by design)", TEAL, TEAL_LIGHT),
    ]
    for index, (value, label, accent, fill) in enumerate(metrics):
        y = 1.55 + index * 1.37
        add_box(slide, 9.82, y, 2.90, 1.05, fill=fill, line=fill)
        add_text(
            slide,
            10.02,
            y + 0.11,
            0.95,
            0.45,
            value,
            size=20,
            bold=True,
            text_color=accent,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            11.04,
            y + 0.10,
            1.43,
            0.70,
            label,
            size=10.8,
            text_color=INK,
            valign=MSO_ANCHOR.MIDDLE,
        )
    add_multiline(
        slide,
        9.82,
        5.70,
        2.90,
        0.63,
        [
            {
                "text": "Exploratory comparison",
                "size": 9.2,
                "bold": True,
                "color": RED,
                "space_after": 2,
            },
            {
                "text": "Different seeds; not a paired causal test.",
                "size": 9.3,
                "color": MUTED,
            },
        ],
        fill=RED_LIGHT,
        line=RED_LIGHT,
    )
    add_takeaway(
        slide,
        "The unseen-entry trajectory shows a steeper DQS rise and reaches AUROC 0.786 at Loop10, but neither shows a stable plateau by Loop15.",
        color_value=ORANGE,
    )
    slide.notes_slide.notes_text_frame.text = note


def add_pipeline_arrow(slide, x: float, y: float, width: float = 0.38) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(0.46),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GREY_LIGHT
    shape.line.color.rgb = GREY_LIGHT


def slide_5(prs: Presentation, note: str) -> None:
    slide = new_slide(prs, 5, TITLES[4])
    add_text(
        slide,
        0.75,
        1.18,
        11.9,
        0.60,
        "AUROC and DQS may improve, but neither directly proves that CL found the correct issues or that the LLM corrected them correctly.",
        size=16.8,
        bold=True,
        text_color=BLUE_DARK,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    pipeline_y = 2.12
    boxes = [
        (0.65, 1.65, "Original\nlabels", PANEL, GREY),
        (2.78, 2.34, "Confident Learning\nselection", BLUE_LIGHT, BLUE),
        (5.60, 2.34, "LLM\nkeep / relabel / mask", ORANGE_LIGHT, ORANGE),
        (8.42, 1.65, "Refined\ndataset", TEAL_LIGHT, TEAL),
    ]
    for x, width, text, fill, accent in boxes:
        add_box(slide, x, pipeline_y, width, 0.90, fill=fill, line=accent, line_width=1.2)
        add_text(
            slide,
            x + 0.10,
            pipeline_y + 0.10,
            width - 0.20,
            0.68,
            text,
            size=13.5,
            bold=True,
            text_color=accent if x != 0.65 else INK,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
    add_pipeline_arrow(slide, 2.40, pipeline_y + 0.22)
    add_pipeline_arrow(slide, 5.22, pipeline_y + 0.22)
    add_pipeline_arrow(slide, 8.04, pipeline_y + 0.22)

    add_text(
        slide,
        2.90,
        3.12,
        2.08,
        0.46,
        "Did it find true\nlabel problems?",
        size=10.8,
        text_color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        5.74,
        3.12,
        2.05,
        0.46,
        "Did it choose the\ncorrect action?",
        size=10.8,
        text_color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    add_multiline(
        slide,
        10.45,
        1.90,
        2.25,
        1.24,
        [
            {
                "text": "AUROC",
                "size": 15,
                "bold": True,
                "color": TEAL,
                "space_after": 5,
            },
            {
                "text": "Held-out ranking utility",
                "size": 11.2,
                "bold": True,
                "color": INK,
            },
            {
                "text": "Does not verify individual selections or actions.",
                "size": 9.7,
                "color": MUTED,
            },
        ],
        fill=TEAL_LIGHT,
        line=TEAL_LIGHT,
    )
    add_multiline(
        slide,
        10.45,
        3.28,
        2.25,
        1.24,
        [
            {
                "text": "DQS",
                "size": 15,
                "bold": True,
                "color": BLUE,
                "space_after": 5,
            },
            {
                "text": "Model-label consistency",
                "size": 11.2,
                "bold": True,
                "color": INK,
            },
            {
                "text": "Uses the same OOF evidence as CL; not independent.",
                "size": 9.7,
                "color": MUTED,
            },
        ],
        fill=BLUE_LIGHT,
        line=BLUE_LIGHT,
    )

    add_section_label(slide, 0.78, 4.22, "External validation", RED)
    add_multiline(
        slide,
        0.78,
        4.60,
        5.72,
        1.12,
        [
            {
                "text": "CORRECTION",
                "size": 9.2,
                "bold": True,
                "color": ORANGE,
                "space_after": 4,
            },
            {
                "text": "Can the LLM correct an externally flagged entry?",
                "size": 14.0,
                "bold": True,
            },
            {"text": "Med-PaLM radiologist reference", "size": 11.3, "color": MUTED},
        ],
        fill=ORANGE_LIGHT,
        line=ORANGE,
    )
    add_multiline(
        slide,
        6.68,
        4.60,
        5.72,
        1.12,
        [
            {
                "text": "DETECTION",
                "size": 9.2,
                "bold": True,
                "color": BLUE,
                "space_after": 4,
            },
            {
                "text": "Can CL enrich independently identified disagreements?",
                "size": 14.0,
                "bold": True,
            },
            {"text": "REFLACX radiologist reference", "size": 11.3, "color": MUTED},
        ],
        fill=BLUE_LIGHT,
        line=BLUE,
    )
    add_takeaway(
        slide,
        "Detection and correction therefore require separate external validation.",
        color_value=RED,
    )
    slide.notes_slide.notes_text_frame.text = note


def slide_6(prs: Presentation, note: str) -> None:
    slide = new_slide(prs, 6, TITLES[5])
    add_section_label(slide, 0.72, 1.16, "Expert reference", BLUE)
    add_multiline(
        slide,
        0.72,
        1.50,
        5.20,
        2.10,
        [
            {
                "text": "Med-PaLM validation set",
                "size": 17.5,
                "bold": True,
                "color": INK,
                "space_after": 12,
            },
            {
                "text": "1,378 report-finding entries",
                "size": 22,
                "bold": True,
                "color": BLUE,
                "space_after": 8,
            },
            {
                "text": "Reviewed by 3 US board-certified radiologists",
                "size": 13.5,
                "color": INK,
                "space_after": 11,
            },
            {
                "text": "Adjudicated radiologist reference",
                "size": 12.2,
                "bold": True,
                "color": BLUE_DARK,
            },
        ],
        fill=BLUE_LIGHT,
        line=BLUE_LIGHT,
        margin=0.22,
    )
    add_multiline(
        slide,
        0.72,
        3.92,
        5.20,
        1.66,
        [
            {
                "text": "CURRENT PROJECT SETTING",
                "size": 9.6,
                "bold": True,
                "color": MUTED,
                "space_after": 7,
            },
            {"text": "12 target labels", "size": 13.2, "bold": True},
            {"text": "U-Ones binary classification", "size": 13.2, "bold": True},
            {
                "text": "498 compatible benchmark entries",
                "size": 18.5,
                "bold": True,
                "color": TEAL,
                "space_before": 7,
            },
        ],
        fill=PANEL,
        line=GRID,
        margin=0.22,
    )

    add_section_label(slide, 6.35, 1.16, "Benchmark results", ORANGE)
    add_box(slide, 6.35, 1.50, 2.85, 1.70, fill=GREY_LIGHT, line=GREY_LIGHT)
    add_text(
        slide,
        6.52,
        1.72,
        2.51,
        0.32,
        "ORIGINAL LABELS",
        size=10,
        text_color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        6.52,
        2.05,
        2.51,
        0.68,
        "74.5%",
        size=34,
        text_color=GREY,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        6.52,
        2.76,
        2.51,
        0.28,
        "Agreement across all 498 entries",
        size=9.5,
        text_color=MUTED,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_box(slide, 9.55, 1.50, 2.85, 1.70, fill=TEAL_LIGHT, line=TEAL_LIGHT)
    add_text(
        slide,
        9.72,
        1.72,
        2.51,
        0.32,
        "LLM-RETAINED DECISIONS",
        size=10,
        text_color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        9.72,
        2.05,
        2.51,
        0.68,
        "96.2%",
        size=34,
        text_color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        9.72,
        2.76,
        2.51,
        0.28,
        "Agreement with expert reference",
        size=9.5,
        text_color=MUTED,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        6.35,
        3.43,
        6.05,
        0.62,
        "Same entry-specific LLM refinement as our pipeline;\nonly the source of flagged entries changes.",
        size=14.5,
        bold=True,
        text_color=INK,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_section_label(slide, 6.35, 4.27, "Action-level evidence", ORANGE)
    action_metrics = [
        ("90.1%", "Relabel\nprecision", ORANGE, ORANGE_LIGHT),
        ("85.8%", "Relabel\nrecall", BLUE, BLUE_LIGHT),
        ("3.2%", "Harmful\nflip rate", TEAL, TEAL_LIGHT),
    ]
    for index, (value, label, accent, fill) in enumerate(action_metrics):
        x = 6.35 + index * 2.08
        add_box(slide, x, 4.58, 1.89, 1.05, fill=fill, line=fill)
        add_text(
            slide,
            x + 0.10,
            4.68,
            0.88,
            0.50,
            value,
            size=20,
            bold=True,
            text_color=accent,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            x + 1.00,
            4.68,
            0.78,
            0.52,
            label,
            size=10.0,
            bold=True,
            text_color=INK,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
    add_text(
        slide,
        6.52,
        5.82,
        5.72,
        0.38,
        "Supports correction on flagged cases; does not test upstream CL detection.",
        size=11.6,
        text_color=RED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_takeaway(
        slide,
        "On externally flagged entries, the existing LLM refinement stage achieves high relabel precision with few harmful flips.",
        color_value=TEAL,
    )
    add_source(
        slide,
        "Source: PhysioNet, Application of Med-PaLM 2 in the refinement of MIMIC-CXR labels, v1.0.0.",
    )
    slide.notes_slide.notes_text_frame.text = note


def slide_7(prs: Presentation, note: str) -> None:
    slide = new_slide(prs, 7, TITLES[6])

    add_box(slide, 11.38, 1.08, 1.30, 0.42, fill=ORANGE_LIGHT, line=ORANGE_LIGHT)
    add_text(
        slide,
        11.52,
        1.17,
        1.02,
        0.20,
        "PENDING",
        size=9.3,
        text_color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    add_section_label(slide, 0.68, 1.16, "Independent radiologist reference", BLUE)
    add_box(slide, 0.68, 1.50, 4.64, 2.08, fill=BLUE_LIGHT, line=BLUE_LIGHT)
    add_text(
        slide,
        0.92,
        1.73,
        4.16,
        0.34,
        "REFLACX",
        size=17.5,
        text_color=BLUE_DARK,
        bold=True,
    )
    add_text(
        slide,
        0.92,
        2.14,
        4.16,
        0.58,
        "2,616 MIMIC-CXR images",
        size=24,
        text_color=BLUE,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        0.92,
        2.86,
        4.16,
        0.42,
        "Independently read and labelled by 5 radiologists",
        size=13.0,
        text_color=INK,
        bold=True,
    )

    add_box(slide, 0.68, 3.82, 4.64, 1.48, fill=TEAL_LIGHT, line=TEAL_LIGHT)
    add_text(
        slide,
        0.92,
        4.03,
        4.16,
        0.26,
        "AFTER MATCHING IT TO OUR SETTING",
        size=9.5,
        text_color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    subset_metrics = [
        (0.86, "1,646", "images"),
        (2.19, "3,172", "valid entries"),
        (3.66, "6", "mapped labels"),
    ]
    for x, value, label in subset_metrics:
        add_text(
            slide,
            x,
            4.37,
            1.28,
            0.42,
            value,
            size=22,
            text_color=TEAL,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            x,
            4.84,
            1.28,
            0.22,
            label,
            size=10.5,
            text_color=INK,
            bold=True,
            align=PP_ALIGN.CENTER,
        )

    add_section_label(slide, 5.74, 1.16, "Mini iterative experiment", ORANGE)
    pipeline_boxes = [
        (5.74, 1.42, "Original MIMIC\nlabels", PANEL, GREY),
        (7.66, 1.38, "4-fold OOF\n+ CL", BLUE_LIGHT, BLUE),
        (9.53, 1.42, "LLM\nrefinement", ORANGE_LIGHT, ORANGE),
        (11.44, 1.24, "Next\nloop", TEAL_LIGHT, TEAL),
    ]
    for index, (x, width, text, fill, accent) in enumerate(pipeline_boxes):
        add_box(slide, x, 1.50, width, 0.94, fill=fill, line=accent, line_width=1.0)
        add_text(
            slide,
            x + 0.08,
            1.67,
            width - 0.16,
            0.54,
            text,
            size=11.0,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        if index < len(pipeline_boxes) - 1:
            next_x = pipeline_boxes[index + 1][0]
            add_pipeline_arrow(slide, x + width + 0.09, 1.74, next_x - (x + width) - 0.18)

    add_box(slide, 5.74, 2.68, 6.94, 0.52, fill=PANEL, line=GRID)
    add_text(
        slide,
        5.92,
        2.82,
        6.58,
        0.22,
        "REFLACX labels are hidden from the pipeline and used only for evaluation.",
        size=10.7,
        text_color=BLUE_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    add_section_label(slide, 5.74, 3.48, "Planned measurements", ORANGE)
    measurement_cards = [
        (
            5.74,
            "01",
            "TRUE LABEL QUALITY",
            "Correct valid entries\n÷ original 3,172 entries",
            BLUE,
            BLUE_LIGHT,
        ),
        (
            8.07,
            "02",
            "DQS VALIDITY",
            "DQS vs true label\nquality across loops",
            TEAL,
            TEAL_LIGHT,
        ),
        (
            10.40,
            "03",
            "CL DETECTION",
            "Precision · Recall\nAUPRC · Enrichment",
            ORANGE,
            ORANGE_LIGHT,
        ),
    ]
    for x, number, heading, detail, accent, fill in measurement_cards:
        add_box(slide, x, 3.82, 2.16, 1.86, fill=fill, line=fill)
        add_text(
            slide,
            x + 0.16,
            4.02,
            0.32,
            0.28,
            number,
            size=10,
            text_color=accent,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            x + 0.55,
            4.02,
            1.44,
            0.26,
            heading,
            size=9.2,
            text_color=accent,
            bold=True,
        )
        add_text(
            slide,
            x + 0.17,
            4.54,
            1.82,
            0.78,
            detail,
            size=12.2,
            text_color=INK,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )

    add_takeaway(
        slide,
        "Tests whether CL finds real disagreements and whether DQS tracks reference-based data quality.",
        color_value=BLUE_DARK,
    )
    add_source(
        slide,
        "Source: REFLACX v1.0.0, PhysioNet; 2,616 unique MIMIC-CXR images and five radiologists.",
    )
    slide.notes_slide.notes_text_frame.text = note


def slide_8(prs: Presentation, note: str) -> None:
    slide = new_slide(prs, 8, TITLES[7])
    columns = [
        (
            "WHAT WE LEARNED\nABOUT ITERATION",
            BLUE,
            BLUE_LIGHT,
            [
                "Oscillation is concentrated in a few high-variance labels.",
                "Influential test studies remain a working hypothesis.",
                "No stable plateau by Loop15; AUROC remains non-monotonic.",
                "Unseen-entry run: higher action yield and steeper DQS (exploratory; different seeds).",
            ],
        ),
        (
            "WHAT EXTERNAL\nEVIDENCE SUPPORTS",
            TEAL,
            TEAL_LIGHT,
            [
                "Med-PaLM supports LLM refinement when a suspicious entry is supplied.",
                "Relabel precision 90.1%; recall 85.8%.",
                "Harmful flip rate 3.2%.",
            ],
        ),
        (
            "WHAT IS STILL\nPENDING",
            ORANGE,
            ORANGE_LIGHT,
            [
                "Two new paired seeds: final exact p-value.",
                "REFLACX: CL detection enrichment.",
                "REFLACX: DQS correspondence with external reference quality.",
            ],
        ),
    ]
    for index, (heading, accent, fill, bullets) in enumerate(columns):
        x = 0.62 + index * 4.22
        add_box(slide, x, 1.28, 3.86, 3.72, fill=fill, line=fill)
        add_text(
            slide,
            x + 0.22,
            1.52,
            3.42,
            0.62,
            heading,
            size=11.5,
            bold=True,
            text_color=accent,
        )
        add_rule(slide, x + 0.22, 2.28, 3.42, line_color=accent, line_width=1.4)
        add_multiline(
            slide,
            x + 0.18,
            2.52,
            3.50,
            2.10,
            [
                {
                    "text": bullet,
                    "size": 11.3,
                    "color": INK,
                    "bullet": True,
                    "space_after": 8,
                }
                for bullet in bullets
            ],
            margin=0.04,
        )

    add_section_label(slide, 0.72, 5.28, "Next steps", RED)
    steps = [
        "Complete added-seed paired analysis",
        "Complete REFLACX detection pilot",
        "Combine detection and correction evidence",
    ]
    for index, step in enumerate(steps, start=1):
        x = 0.72 + (index - 1) * 4.08
        add_circle(slide, x, 5.75, 0.34, fill=RED_LIGHT, line=RED)
        add_text(
            slide,
            x,
            5.775,
            0.34,
            0.22,
            str(index),
            size=9.3,
            text_color=RED,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            x + 0.46,
            5.66,
            2.42,
            0.63,
            step,
            size=10.8,
            bold=True,
            valign=MSO_ANCHOR.MIDDLE,
        )
    add_takeaway(
        slide,
        "The project is moving from showing that metrics change to testing whether detection and correction actually work.",
        color_value=BLUE_DARK,
    )
    slide.notes_slide.notes_text_frame.text = note


def ensure_inputs() -> None:
    required = [
        SOURCE_MD,
        PAGE2_TRAJECTORY,
        STANDARD_LOOP_METRICS,
        UNSEEN_LOOP_METRICS,
        UNSEEN_DQS,
        STANDARD_DQS,
    ]
    required.extend(
        PAGE2_BASELINE_ROOT
        / f"seed_{seed}"
        / "baseline_no_clean"
        / "test_study_auroc_summary.csv"
        for seed in PAGE2_SEEDS
    )
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing required inputs:\n" + "\n".join(missing))


def validate_deck(path: Path, notes: dict[int, str]) -> dict[str, Any]:
    prs = Presentation(path)
    if len(prs.slides) != 8:
        raise AssertionError(f"Expected 8 slides, found {len(prs.slides)}")
    observed_titles: list[str] = []
    out_of_bounds: list[dict[str, Any]] = []
    unexpected_fonts: list[dict[str, Any]] = []
    note_checks: dict[str, str] = {}
    for page, slide in enumerate(prs.slides, start=1):
        title_shape = next(
            (shape for shape in slide.shapes if shape.name == "slide-title"), None
        )
        if title_shape is None:
            raise AssertionError(f"Slide {page} is missing its title")
        observed_titles.append(title_shape.text.strip())
        actual_note = slide.notes_slide.notes_text_frame.text.strip()
        expected_note = notes.get(page, "")
        if actual_note != expected_note:
            raise AssertionError(f"Slide {page} speaker note does not match source")
        note_checks[str(page)] = "present" if actual_note else "intentionally blank"
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if (
                            run.text.strip()
                            and run.font.name != FONT_FAMILY
                        ):
                            unexpected_fonts.append(
                                {
                                    "page": page,
                                    "shape": shape.name,
                                    "text": run.text[:80],
                                    "font": run.font.name,
                                }
                            )
            if (
                shape.left < -2
                or shape.top < -2
                or shape.left + shape.width > prs.slide_width + 2
                or shape.top + shape.height > prs.slide_height + 2
            ):
                out_of_bounds.append(
                    {
                        "page": page,
                        "shape": shape.name,
                        "left": int(shape.left),
                        "top": int(shape.top),
                        "width": int(shape.width),
                        "height": int(shape.height),
                    }
                )
    if observed_titles != TITLES:
        raise AssertionError(
            json.dumps(
                {"expected": TITLES, "observed": observed_titles},
                ensure_ascii=False,
                indent=2,
            )
        )
    if out_of_bounds:
        raise AssertionError(
            "Out-of-bounds shapes:\n"
            + json.dumps(out_of_bounds, ensure_ascii=False, indent=2)
        )
    if unexpected_fonts:
        raise AssertionError(
            "Unexpected visible slide fonts:\n"
            + json.dumps(unexpected_fonts, ensure_ascii=False, indent=2)
        )
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise AssertionError(f"Corrupt PPTX member: {bad_member}")
        note_parts = [
            name
            for name in archive.namelist()
            if name.startswith("ppt/notesSlides/") and name.endswith(".xml")
        ]
        if len(note_parts) < len(NOTE_PAGES):
            raise AssertionError("PPTX package is missing speaker-note parts")
    return {
        "status": "passed",
        "slides": 8,
        "completed_slides": [1, 2, 4, 5, 6, 7, 8],
        "placeholder_slides": [3],
        "results_pending_slides": [3, 7],
        "speaker_notes": note_checks,
        "shape_bounds": "passed",
        "visible_slide_fonts": FONT_FAMILY,
        "pptx_zip_integrity": "passed",
    }


def write_manifest(notes: dict[int, str], preflight: dict[str, Any]) -> None:
    manifest = {
        "status": "passed",
        "source_outline": str(SOURCE_MD),
        "deck": str(OUTPUT_PPTX),
        "presentation_font": FONT_FAMILY,
        "page2_baseline_means": page2_baseline_means(),
        "slides": [
            {
                "page": page,
                "title": title,
                "status": (
                    "placeholder"
                    if page == 3
                    else "design_shown_results_pending"
                    if page == 7
                    else "complete"
                ),
                "speaker_note_chars": len(notes.get(page, "")),
            }
            for page, title in enumerate(TITLES, start=1)
        ],
        "assets": {
            "page2_oscillation": str(ASSET_DIR / "page2_auroc_oscillation.png"),
            "page4_auroc": str(ASSET_DIR / "page4_auroc_1_15.png"),
            "page4_dqs": str(ASSET_DIR / "page4_sample_dqs_0_15.png"),
        },
        "preflight": preflight,
    }
    (OUTPUT_DIR / "delivery_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> int:
    ensure_inputs()
    configure_plot_style()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    notes = parse_notes()

    page2_chart = ASSET_DIR / "page2_auroc_oscillation.png"
    page4_auroc = ASSET_DIR / "page4_auroc_1_15.png"
    page4_dqs = ASSET_DIR / "page4_sample_dqs_0_15.png"
    plot_page2(page2_chart)
    plot_page4_auroc(page4_auroc)
    plot_page4_dqs(page4_dqs)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "From open questions to validation"
    prs.core_properties.subject = (
        "AUROC oscillation, extended refinement, Med-PaLM correction validation, "
        "and REFLACX detection validation"
    )
    prs.core_properties.author = "Yihang"
    prs.core_properties.keywords = (
        "confident learning, LLM refinement, AUROC, DQS, Med-PaLM, REFLACX"
    )

    slide_1(prs, notes[1])
    slide_2(prs, notes[2], page2_chart)
    placeholder_slide(
        prs,
        3,
        TITLES[2],
        "Reserved for added-seed paired analysis",
        notes[3],
    )
    slide_4(prs, notes[4], page4_auroc, page4_dqs)
    slide_5(prs, notes[5])
    slide_6(prs, notes[6])
    slide_7(prs, notes[7])
    slide_8(prs, notes[8])
    prs.save(OUTPUT_PPTX)

    preflight = validate_deck(OUTPUT_PPTX, notes)
    write_manifest(notes, preflight)
    print(
        json.dumps(
            {
                "status": "passed",
                "deck": str(OUTPUT_PPTX),
                "manifest": str(OUTPUT_DIR / "delivery_manifest.json"),
                "preflight": preflight,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
