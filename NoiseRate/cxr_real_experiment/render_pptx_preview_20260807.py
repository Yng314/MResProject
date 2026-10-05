#!/usr/bin/env python3
"""Render lightweight previews from PPTX shape geometry for visual QA."""

from __future__ import annotations

import io
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN


HERE = Path(__file__).resolve().parent
DECK = HERE / "next_meeting_slides_20260807/next_meeting_update_20260807.pptx"
OUT = HERE / "next_meeting_slides_20260807/preview"
WIDTH = 1600
HEIGHT = 900
FONT_ROOT = Path("/homes/yz3522/.local/share/fonts/msttcorefonts")
FONTS = {
    (False, False): FONT_ROOT / "Times_New_Roman.ttf",
    (True, False): FONT_ROOT / "Times_New_Roman_Bold.ttf",
    (False, True): FONT_ROOT / "Times_New_Roman_Italic.ttf",
    (True, True): FONT_ROOT / "Times_New_Roman_Bold_Italic.ttf",
}


def color_tuple(color, fallback=(255, 255, 255)):
    try:
        value = color.rgb
        if value is None:
            return fallback
        return tuple(value)
    except (AttributeError, ValueError, TypeError):
        return fallback


def shape_box(shape, prs):
    sx = WIDTH / prs.slide_width
    sy = HEIGHT / prs.slide_height
    return (
        int(shape.left * sx),
        int(shape.top * sy),
        int((shape.left + shape.width) * sx),
        int((shape.top + shape.height) * sy),
    )


def wrap_line(draw, text, font, max_width):
    if not text:
        return [""]
    words = text.split()
    lines = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if draw.textlength(trial, font=font) <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def render_text(draw, shape, box, prs):
    frame = shape.text_frame
    sx = WIDTH / prs.slide_width
    sy = HEIGHT / prs.slide_height
    left = box[0] + int(frame.margin_left * sx)
    right = box[2] - int(frame.margin_right * sx)
    top = box[1] + int(frame.margin_top * sy)
    bottom = box[3] - int(frame.margin_bottom * sy)
    rendered = []
    for paragraph in frame.paragraphs:
        text = paragraph.text
        run = next((item for item in paragraph.runs if item.text), None)
        bold = bool(run.font.bold) if run is not None else False
        italic = bool(run.font.italic) if run is not None else False
        size_pt = run.font.size.pt if run is not None and run.font.size else 14
        font_size = max(8, int(size_pt * WIDTH / (13.333 * 72)))
        font = ImageFont.truetype(str(FONTS[(bold, italic)]), font_size)
        color = color_tuple(run.font.color, (24, 33, 47)) if run is not None else (24, 33, 47)
        for explicit_line in text.splitlines() or [""]:
            for line in wrap_line(draw, explicit_line, font, max(1, right - left)):
                line_height = int(font_size * 1.12)
                rendered.append((line, font, color, paragraph.alignment, line_height))
        rendered.append(("", font, color, paragraph.alignment, max(1, int(size_pt * 0.10))))
    total_height = sum(item[4] for item in rendered)
    if frame.vertical_anchor == MSO_ANCHOR.MIDDLE:
        cursor_y = top + max(0, (bottom - top - total_height) // 2)
    elif frame.vertical_anchor == MSO_ANCHOR.BOTTOM:
        cursor_y = max(top, bottom - total_height)
    else:
        cursor_y = top
    for text, font, color, alignment, line_height in rendered:
        if text:
            text_width = draw.textlength(text, font=font)
            if alignment == PP_ALIGN.CENTER:
                cursor_x = left + max(0, (right - left - text_width) / 2)
            elif alignment == PP_ALIGN.RIGHT:
                cursor_x = right - text_width
            else:
                cursor_x = left
            draw.text((cursor_x, cursor_y), text, font=font, fill=color)
        cursor_y += line_height


def render_shape(canvas, draw, shape, prs):
    box = shape_box(shape, prs)
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture = Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
        picture = picture.resize((max(1, box[2] - box[0]), max(1, box[3] - box[1])))
        canvas.paste(picture, (box[0], box[1]))
        return
    if shape.shape_type in {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX}:
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            fill = color_tuple(shape.fill.fore_color, (255, 255, 255))
            outline = color_tuple(shape.line.color, fill)
            draw.rounded_rectangle(box, radius=8, fill=fill, outline=outline, width=1)
        if shape.has_text_frame:
            render_text(draw, shape, box, prs)
        return
    if shape.shape_type in {MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.FREEFORM}:
        color = color_tuple(shape.line.color, (217, 222, 231))
        draw.line((box[0], box[1], box[2], box[3]), fill=color, width=2)
        return
    if shape.has_text_frame:
        render_text(draw, shape, box, prs)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation(DECK)
    previews = []
    for page, slide in enumerate(prs.slides, start=1):
        canvas = Image.new("RGB", (WIDTH, HEIGHT), "white")
        draw = ImageDraw.Draw(canvas)
        for shape in slide.shapes:
            render_shape(canvas, draw, shape, prs)
        output = OUT / f"slide_{page:02d}.png"
        canvas.save(output)
        previews.append(canvas.resize((800, 450)))
    contact = Image.new("RGB", (1600, 900), "white")
    for index, preview in enumerate(previews):
        contact.paste(preview, ((index % 2) * 800, (index // 2) * 450))
    contact.save(OUT / "contact_sheet.png")
    print(OUT / "contact_sheet.png")


if __name__ == "__main__":
    main()
