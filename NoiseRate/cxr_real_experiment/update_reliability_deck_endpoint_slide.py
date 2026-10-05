#!/usr/bin/env python3
"""Replace slide 6 of the reliability deck with reusable endpoint evidence."""

from __future__ import annotations

import argparse
import json
import math
import os
import posixpath
import re
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
REPORT_DIR = HERE / "reliability_validation_report_20260715"
DEFAULT_INPUT_DECK = REPORT_DIR / "reliability_validation_update_20260715.pptx"
DEFAULT_OUTPUT_DECK = (
    REPORT_DIR / "reliability_validation_update_20260717_exact_definitions.pptx"
)
DEFAULT_ENDPOINT_DIR = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "results_mobilenetv3_own_top20_binary_v5_llm_refine_5seed/20260714_065539/"
    "evaluation_interim_3seed_loop5_loop8_20260717"
)
DEFAULT_SAMPLE_CHART = (
    HERE
    / "evaluation_followup_20260713"
    / "real_data_quality_figures"
    / "sample_issue_free_rate_five_seed.png"
)
DEFAULT_SYNTHETIC_CHART = (
    HERE
    / "evaluation_followup_20260713"
    / "synthetic_dqs_validation"
    / "synthetic_dqs_known_truth_trajectories.png"
)

INK = RGBColor(0x1F, 0x24, 0x30)
MUTED = RGBColor(0x6F, 0x76, 0x8A)
BLUE = RGBColor(0x2E, 0x47, 0x80)
BLUE_MARK = "#2E77BB"
RED_MARK = "#C9362B"
GRID = RGBColor(0xE6, 0xE8, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


@dataclass(frozen=True)
class Contrast:
    name: str
    short_label: str
    direction_label: str
    mean_delta: float
    positive_seeds: int
    n_seeds: int
    exact_p: float
    holm_exact_p: float
    ci_low: float
    ci_high: float


@dataclass(frozen=True)
class EndpointEvidence:
    per_seed: pd.DataFrame
    summary: pd.DataFrame
    seeds: list[int]
    endpoint_loop: int
    primary: Contrast
    secondary: Contrast
    planned_contrasts: tuple[Contrast, ...]
    proper_score_note: str


def build_speaker_notes(evidence: EndpointEvidence) -> dict[int, str]:
    seed_count = len(evidence.seeds)
    if evidence.secondary.short_label in {"vs refinement Loop 5", "vs Loop 5"}:
        secondary_spoken = "Loop 5"
    elif evidence.secondary.short_label == "vs removal":
        secondary_spoken = "simple removal"
    else:
        secondary_spoken = evidence.secondary.short_label.removeprefix("vs ")

    direction_note = (
        f"和 baseline 相比，{evidence.primary.positive_seeds}/{evidence.primary.n_seeds} 个 seed 是正向的；"
        f"和 {secondary_spoken} 相比，{evidence.secondary.positive_seeds}/"
        f"{evidence.secondary.n_seeds} 个是正向的。"
    )
    primary_crosses_zero = evidence.primary.ci_low <= 0 <= evidence.primary.ci_high
    secondary_crosses_zero = evidence.secondary.ci_low <= 0 <= evidence.secondary.ci_high
    if primary_crosses_zero and secondary_crosses_zero:
        uncertainty_note = (
            "右边两个区间都跨过了 0。"
            "所以更准确的说法是：Loop 8 相对 baseline 的方向一致，"
            "但 Loop 5 之后的额外收益还不稳定。"
        )
    else:
        uncertainty_note = (
            "右边的区间用来说明这个结果还有多大不确定性。"
            "这里我会把平均差值、区间和每个 seed 的方向放在一起看。"
        )

    proper_score_note = (
        "Brier 和 NLL 也支持同一个方向。" if evidence.proper_score_note else ""
    )

    return {
        1: (
            "今天我主要讲三个问题。第一，coverage-adjusted DQS 到底有没有意义。"
            "第二，分数变好是不是只是因为删掉了更多数据。"
            "第三，换不同的 seed 后，实验结果还会不会一样。"
            "前五页讲指标和分母，第六页讲三个 seed 跑到 Loop 8 的模型结果。"
        ),
        2: (
            "这一页把几个分数的定义写清楚。E_t 是第 t 轮 action 之后仍然有效的 label entries。"
            "M_t 是在这些 entries 里面，根据 K-fold OOF predictions 和 confident learning 估计出的 issue 数。"
            "所以 raw entry DQS 是一减去 M_t 除以 E_t，它回答的是现在留下来的 entries 有多干净。"
            "Valid-entry coverage 是 E_t 除以最开始的 E_0。把这两个数相乘，"
            "就得到 coverage-adjusted entry DQS，也就是用原始分母看，还有多少 entry 留下来并且被估计为干净。"
            "Sample level 完全对应：S_t 只包括还有至少一个有效标签的 samples，"
            "Q_t 是其中至少有一个有效标签被 flag 的 samples。没有有效标签的 sample 不会被当成干净。"
        ),
        3: (
            "这里我们用合成数据，因为哪些标签是对的、哪些是错的，我们都知道。"
            "每组数据有 5,000 个二分类样本，每个样本只有一个标签，开始时加入 12% 的标签噪声。"
            "我们重复了 10 个 random seeds，每个 seed 做 5 个 cleaning loops；图里的 Loop 0 是操作前的 baseline。"
            "DQS 用 4-fold OOF predictions 来算，而且这些 predictions 在 baseline 生成后就固定下来，"
            "这样图里的变化只来自 correction 或 removal，而不是因为每一轮又换了一个 detector。"
            "图里的第一条 truth line 是剩余标签中的正确率；第二条是原始标签中现在仍保留且正确的比例。"
            "左上是把错误改对，两种分数都会上升。右上是把错误删掉，raw DQS 会接近 1，"
            "但 coverage-adjusted DQS 会保留数据损失的信息。左下是随机删除，"
            "raw DQS 变化不大，coverage-adjusted DQS 会下降。"
            "右下是故意把正确标签改错，两种分数都会下降。"
            "所以 coverage-adjusted DQS 的方向基本符合我们的预期。"
            "最后我们也检查了用来算 DQS 的预测：预测比较可靠时，误差是 0.019；"
            "预测变弱时，误差升到 0.055。"
        ),
        4: (
            "这一页回到真实数据。左边看的是现在留下来的标签，右边把分母固定在最开始的数据上。"
            "Simple removal 的 raw entry DQS 一直上升，但 coverage-adjusted entry DQS 降到了 0.810，"
            "因为它删掉了不少数据。LLM refinement 的两个 entry DQS 都比 baseline 稍高，"
            "而且还保留了大约 99.6% 的标签。"
        ),
        5: (
            "这一页把上一页的 DQS 从 entry level 换到 sample level。"
            "左边是 raw sample DQS：只要一个 sample 里有任何有效标签被标成可疑，"
            "整个 sample 就算有问题。右边是 coverage-adjusted sample DQS，"
            "它把 sample coverage 也算进去。Simple removal 的 raw sample DQS 升到 0.865，"
            "但 coverage-adjusted sample DQS 只有 0.698。LLM refinement 的 raw sample DQS 和 "
            "coverage-adjusted sample DQS 是 0.756 和 0.754，"
            "同时保留了大约 99.7% 的 samples。所以 sample level 和 entry level 得到的是同一个结论。"
        ),
        6: (
            f"这一页看模型结果。浅色线是每个 seed，粗线是 {seed_count} 个 seed 的平均值，"
            "误差条表示 seed 之间的差别。"
            f"在 Loop {evidence.endpoint_loop}，LLM refinement 比 baseline 平均高 "
            f"{evidence.primary.mean_delta:.4f}，比 {secondary_spoken} 平均高 "
            f"{evidence.secondary.mean_delta:.4f}。{direction_note}"
            f"{uncertainty_note}{proper_score_note}"
        ),
        7: (
            "这页是前面真实数据结果的详细数字，分别列出 entry 和 sample 的 raw DQS、"
            "coverage-adjusted DQS 以及 coverage。正常汇报时我不会展开。"
            "如果老师想看某个分数、覆盖率或者具体小数，我再回到这一页。"
        ),
        8: (
            "这页是合成实验的详细数字。它把每种操作的真实结果、raw DQS、"
            "coverage-adjusted DQS 和数据保留比例放在一起。表里的 retained and correct fraction，"
            "就是原始标签中现在仍保留而且正确的比例。老师如果想核对图里的具体数值，可以看这一页。"
        ),
        9: (
            "这页是四个预先定好的比较。最重要的是 Loop 8 对 baseline：平均高 0.0425，"
            "而且三个 seed 都是正向的。不过区间还是稍微跨过 0。"
            "Loop 8 对 Loop 5 只高 0.0068，而且只有两个 seed 是正向的。"
            "所以目前可以说 Loop 8 保持了相对 baseline 的改善，"
            "但不能说 Loop 5 之后每一轮都带来稳定的额外提升。"
        ),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-deck", type=Path, default=DEFAULT_INPUT_DECK)
    parser.add_argument("--endpoint-dir", type=Path, default=DEFAULT_ENDPOINT_DIR)
    parser.add_argument("--sample-chart", type=Path, default=DEFAULT_SAMPLE_CHART)
    parser.add_argument("--synthetic-chart", type=Path, default=DEFAULT_SYNTHETIC_CHART)
    parser.add_argument("--output-deck", type=Path, default=DEFAULT_OUTPUT_DECK)
    parser.add_argument(
        "--chart-output",
        type=Path,
        default=REPORT_DIR / "endpoint_performance_slide_3seed_loop8.png",
    )
    parser.add_argument(
        "--plan-output",
        type=Path,
        default=REPORT_DIR / "deck_plan_20260717_exact_definitions.json",
    )
    parser.add_argument(
        "--summary-output",
        type=Path,
        default=REPORT_DIR / "delivery_summary_20260717_exact_definitions.json",
    )
    return parser.parse_args()


def find_single(directory: Path, patterns: list[str]) -> Path:
    matches: list[Path] = []
    for pattern in patterns:
        matches.extend(directory.glob(pattern))
    matches = sorted(set(path for path in matches if path.is_file()))
    if len(matches) != 1:
        raise FileNotFoundError(
            f"Expected one match in {directory} for {patterns}, found {matches}"
        )
    return matches[0]


def hierarchical_ci(bootstrap: pd.DataFrame, comparison: str) -> tuple[float, float]:
    rows = bootstrap.loc[
        bootstrap["comparison"].eq(comparison)
        & bootstrap["scope"].eq("all_labels")
        & bootstrap["uncertainty"].eq("hierarchical_seed_and_study")
    ]
    if len(rows) != 1:
        raise ValueError(f"Expected one hierarchical interval for {comparison}, got {len(rows)}")
    row = rows.iloc[0]
    return float(row["ci_low_2p5"]), float(row["ci_high_97p5"])


def contrast_from_row(
    stats: pd.DataFrame,
    bootstrap: pd.DataFrame,
    name: str,
    short_label: str,
    direction_label: str,
) -> Contrast:
    rows = stats.loc[stats["comparison"].eq(name)]
    if len(rows) != 1:
        raise ValueError(f"Expected one paired-statistics row for {name}, got {len(rows)}")
    row = rows.iloc[0]
    ci_low, ci_high = hierarchical_ci(bootstrap, name)
    return Contrast(
        name=name,
        short_label=short_label,
        direction_label=direction_label,
        mean_delta=float(row["mean_delta"]),
        positive_seeds=int(row["positive_delta_seeds"]),
        n_seeds=int(row["n_seeds"]),
        exact_p=float(row["exact_signflip_p_two_sided"]),
        holm_exact_p=float(row["holm_family_exact_p"]),
        ci_low=ci_low,
        ci_high=ci_high,
    )


def load_endpoint_evidence(endpoint_dir: Path) -> EndpointEvidence:
    performance_dir = endpoint_dir / "performance"
    if not performance_dir.is_dir():
        performance_dir = endpoint_dir
    per_seed_path = find_single(performance_dir, ["per_seed_stage_metrics.csv"])
    summary_path = find_single(
        performance_dir,
        ["interim_seed_stage_summary.csv", "five_seed_stage_summary.csv"],
    )
    stats_path = find_single(
        performance_dir,
        ["interim_seed_paired_statistics.csv", "prelocked_seed_paired_statistics.csv"],
    )
    bootstrap_path = find_single(
        performance_dir,
        ["interim_hierarchical_bootstrap.csv", "prelocked_hierarchical_bootstrap.csv"],
    )

    per_seed = pd.read_csv(per_seed_path)
    summary = pd.read_csv(summary_path)
    stats = pd.read_csv(stats_path)
    bootstrap = pd.read_csv(bootstrap_path)
    required = {"seed", "method", "loop", "study_weighted_auroc"}
    if not required.issubset(per_seed.columns):
        raise ValueError(f"Missing per-seed columns: {sorted(required - set(per_seed.columns))}")
    seeds = sorted(int(value) for value in per_seed["seed"].unique())
    refine_loops = sorted(
        int(value) for value in per_seed.loc[per_seed["method"].eq("refine"), "loop"].unique()
    )
    remove_loops = sorted(
        int(value) for value in per_seed.loc[per_seed["method"].eq("remove"), "loop"].unique()
    )
    if not seeds or not refine_loops or not remove_loops:
        raise ValueError("Endpoint evidence does not contain seeds, refinement loops, and removal loops")
    endpoint_loop = max(refine_loops)

    primary_name = f"refine_L{endpoint_loop}_vs_baseline"
    primary = contrast_from_row(
        stats,
        bootstrap,
        primary_name,
        "vs baseline",
        f"refinement Loop {endpoint_loop} > baseline",
    )

    same_loop_removal = f"refine_L{endpoint_loop}_vs_remove_L{endpoint_loop}"
    checkpoint_name = f"refine_L{endpoint_loop}_vs_refine_L5"
    if stats["comparison"].eq(same_loop_removal).any():
        secondary = contrast_from_row(
            stats,
            bootstrap,
            same_loop_removal,
            "vs removal",
            f"refinement > simple removal at Loop {endpoint_loop}",
        )
        comparator_loop = endpoint_loop
    elif stats["comparison"].eq(checkpoint_name).any():
        secondary = contrast_from_row(
            stats,
            bootstrap,
            checkpoint_name,
            "vs Loop 5",
            f"refinement Loop {endpoint_loop} > refinement Loop 5",
        )
        comparator_loop = 5
    else:
        fallback_loop = max(loop for loop in remove_loops if loop <= endpoint_loop)
        fallback_name = f"refine_L{fallback_loop}_vs_remove_L{fallback_loop}"
        secondary = contrast_from_row(
            stats,
            bootstrap,
            fallback_name,
            f"Loop {fallback_loop} vs removal",
            f"refinement > simple removal at Loop {fallback_loop}",
        )
        comparator_loop = fallback_loop

    proper_score_note = ""
    if comparator_loop == endpoint_loop:
        refine = per_seed.loc[
            per_seed["method"].eq("refine") & per_seed["loop"].eq(endpoint_loop)
        ].set_index("seed")
        remove = per_seed.loc[
            per_seed["method"].eq("remove") & per_seed["loop"].eq(comparator_loop)
        ].set_index("seed")
        common = refine.index.intersection(remove.index)
        if len(common) == len(seeds):
            brier_better = int(
                (refine.loc[common, "micro_brier"] < remove.loc[common, "micro_brier"]).sum()
            )
            nll_better = int(
                (refine.loc[common, "micro_nll"] < remove.loc[common, "micro_nll"]).sum()
            )
            if brier_better == nll_better == len(common):
                proper_score_note = (
                    f"Brier and NLL also favor refinement over removal in all {len(common)} seeds."
                )
            else:
                proper_score_note = (
                    f"Brier and NLL favor refinement in {brier_better}/{len(common)} "
                    f"and {nll_better}/{len(common)} seeds."
                )

    if primary.n_seeds != len(seeds) or secondary.n_seeds != len(seeds):
        raise ValueError("Paired-statistics seed count does not match per-seed results")
    planned_specs = [
        ("refine_L5_vs_baseline", "Loop 5 refine vs baseline"),
        (f"refine_L{endpoint_loop}_vs_baseline", f"Loop {endpoint_loop} refine vs baseline"),
        (
            f"refine_L{endpoint_loop}_vs_refine_L5",
            f"Loop {endpoint_loop} refine vs Loop 5",
        ),
        ("refine_L5_vs_remove_L5", "Loop 5 refine vs removal"),
    ]
    planned_contrasts = tuple(
        contrast_from_row(stats, bootstrap, name, label, label)
        for name, label in planned_specs
    )
    return EndpointEvidence(
        per_seed=per_seed,
        summary=summary,
        seeds=seeds,
        endpoint_loop=endpoint_loop,
        primary=primary,
        secondary=secondary,
        planned_contrasts=planned_contrasts,
        proper_score_note=proper_score_note,
    )


def build_trajectory_chart(evidence: EndpointEvidence, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    frame = evidence.per_seed.copy()
    baseline = frame.loc[frame["method"].eq("baseline")]
    baseline_mean = float(baseline["study_weighted_auroc"].mean())
    baseline_sd = float(baseline["study_weighted_auroc"].std(ddof=1))

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 11,
            "axes.labelcolor": "#1F2430",
            "xtick.color": "#4F5668",
            "ytick.color": "#4F5668",
        }
    )
    fig, axis = plt.subplots(figsize=(11.4, 6.5), dpi=180)
    fig.patch.set_facecolor("white")
    axis.set_facecolor("white")

    axis.axhspan(
        baseline_mean - baseline_sd,
        baseline_mean + baseline_sd,
        color="#6F768A",
        alpha=0.09,
        linewidth=0,
        zorder=0,
    )
    axis.axhline(
        baseline_mean,
        color="#525A6B",
        linewidth=2.0,
        linestyle=(0, (5, 3)),
        zorder=2,
    )

    styles = {
        "remove": {"color": BLUE_MARK, "marker": "s", "linestyle": "--"},
        "refine": {"color": RED_MARK, "marker": "o", "linestyle": "-"},
    }
    for method, style in styles.items():
        method_frame = frame.loc[frame["method"].eq(method)]
        for seed, seed_frame in method_frame.groupby("seed"):
            seed_frame = seed_frame.sort_values("loop")
            axis.plot(
                seed_frame["loop"],
                seed_frame["study_weighted_auroc"],
                color=style["color"],
                linewidth=1.1,
                linestyle=style["linestyle"],
                alpha=0.22,
                zorder=1,
            )
        method_summary = (
            method_frame.groupby("loop", as_index=False)["study_weighted_auroc"]
            .agg(["mean", "std"])
            .reset_index()
            .sort_values("loop")
        )
        axis.errorbar(
            method_summary["loop"],
            method_summary["mean"],
            yerr=method_summary["std"],
            color=style["color"],
            linewidth=2.7,
            linestyle=style["linestyle"],
            marker=style["marker"],
            markersize=7,
            markerfacecolor="white" if method == "remove" else style["color"],
            markeredgewidth=1.8,
            capsize=4,
            capthick=1.4,
            zorder=4,
        )

    axis.axvline(
        evidence.endpoint_loop,
        color="#AEB4C1",
        linewidth=1.2,
        linestyle=(0, (2, 3)),
        zorder=0,
    )
    values = frame["study_weighted_auroc"].dropna().to_numpy(dtype=float)
    padding = max(0.008, (values.max() - values.min()) * 0.11)
    lower = math.floor((values.min() - padding) * 100) / 100
    upper = math.ceil((values.max() + padding) * 100) / 100
    axis.set_ylim(lower, upper)
    axis.set_xlim(0.72, evidence.endpoint_loop + 0.28)
    axis.set_xticks(range(1, evidence.endpoint_loop + 1))
    axis.set_xlabel("Cleaning loop", fontsize=12, labelpad=8)
    axis.set_ylabel("Study-weighted AUROC", fontsize=12, labelpad=8)
    axis.grid(axis="y", color="#DDE1E9", linewidth=0.9)
    axis.grid(axis="x", visible=False)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.spines["left"].set_color("#AEB4C1")
    axis.spines["bottom"].set_color("#AEB4C1")
    axis.text(
        0.0,
        1.025,
        f"Mean +/- 1 SD; faint traces are seeds {', '.join(map(str, evidence.seeds))}; focused y-axis",
        transform=axis.transAxes,
        fontsize=10,
        color="#6F768A",
        ha="left",
        va="bottom",
    )
    handles = [
        Line2D([0], [0], color="#525A6B", lw=2, linestyle=(0, (5, 3)), label=f"Baseline ({baseline_mean:.3f})"),
        Line2D([0], [0], color=BLUE_MARK, lw=2.7, linestyle="--", marker="s", markerfacecolor="white", label="Simple removal"),
        Line2D([0], [0], color=RED_MARK, lw=2.7, linestyle="-", marker="o", label="LLM refinement"),
    ]
    axis.legend(
        handles=handles,
        loc="upper left",
        frameon=False,
        ncol=3,
        bbox_to_anchor=(0.0, 0.99),
        borderaxespad=0,
        fontsize=10,
        handlelength=2.8,
        columnspacing=1.4,
    )
    fig.subplots_adjust(left=0.095, right=0.985, top=0.89, bottom=0.14)
    fig.savefig(output_path, facecolor="white", bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def set_shape_text(shape, text: str) -> None:
    paragraphs = shape.text_frame.paragraphs
    first = paragraphs[0]
    if first.runs:
        first.runs[0].text = text
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.text = text
    for paragraph in paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def add_text(
    slide,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    *,
    size: float,
    color: RGBColor = INK,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.name = name
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.vertical_anchor = valign
    text_frame.margin_left = Inches(margin)
    text_frame.margin_right = Inches(margin)
    text_frame.margin_top = Inches(margin)
    text_frame.margin_bottom = Inches(margin)
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    paragraph.space_before = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_multi_line_text(
    slide,
    name: str,
    x: float,
    y: float,
    width: float,
    height: float,
    lines: list[tuple[str, float, RGBColor, bool]],
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.name = name
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0
    for index, (text, size, color, bold) in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(3 if index < len(lines) - 1 else 0)
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def add_frame(slide, name: str, x: float, y: float, width: float, height: float):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = GRID
    shape.line.width = Pt(0.75)
    return shape


def add_separator(slide, name: str, x: float, y: float, width: float):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(0.012),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = GRID
    shape.line.color.rgb = GRID
    return shape


def add_picture_contain(slide, path: Path, x: float, y: float, width: float, height: float):
    with Image.open(path) as image:
        image_ratio = image.width / image.height
    box_ratio = width / height
    if image_ratio >= box_ratio:
        rendered_width = width
        rendered_height = width / image_ratio
        rendered_x = x
        rendered_y = y + (height - rendered_height) / 2
    else:
        rendered_height = height
        rendered_width = height * image_ratio
        rendered_x = x + (width - rendered_width) / 2
        rendered_y = y
    picture = slide.shapes.add_picture(
        str(path),
        Inches(rendered_x),
        Inches(rendered_y),
        width=Inches(rendered_width),
        height=Inches(rendered_height),
    )
    picture.name = "endpoint-chart"
    return picture


def replace_slide_six(slide, chart_path: Path, evidence: EndpointEvidence) -> str:
    for shape in list(slide.shapes):
        if shape.name != "slide-background":
            slide.shapes._spTree.remove(shape._element)

    primary_all_positive = evidence.primary.positive_seeds == evidence.primary.n_seeds
    secondary_all_positive = evidence.secondary.positive_seeds == evidence.secondary.n_seeds
    if primary_all_positive and evidence.endpoint_loop > 5:
        title = (
            f"Loop {evidence.endpoint_loop} refinement remains above baseline "
            f"in all {evidence.primary.n_seeds} seed-specific runs"
        )
    elif primary_all_positive and secondary_all_positive:
        title = (
            f"Loop {evidence.endpoint_loop} refinement improves held-out AUROC "
            f"in all {evidence.primary.n_seeds} seed-specific runs"
        )
    else:
        title = (
            f"Loop {evidence.endpoint_loop} refinement performance across "
            f"{evidence.primary.n_seeds} seed-specific runs"
        )
    add_text(slide, "slide-title", 0.62, 0.38, 12.0, 0.82, title, size=20, bold=True)

    add_frame(slide, "endpoint-chart-frame", 0.72, 1.28, 8.0, 5.05)
    add_picture_contain(slide, chart_path, 0.93, 1.53, 7.58, 4.55)
    add_frame(slide, "endpoint-callout-frame", 9.0, 1.28, 3.55, 5.05)

    add_text(
        slide,
        "endpoint-result-heading",
        9.25,
        1.52,
        3.05,
        0.3,
        f"RESULT AT LOOP {evidence.endpoint_loop}",
        size=9.5,
        color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        "endpoint-direction-value",
        9.25,
        1.86,
        3.05,
        0.48,
        f"{evidence.primary.positive_seeds} / {evidence.primary.n_seeds} seeds",
        size=21,
        color=BLUE,
        bold=True,
    )
    if primary_all_positive and secondary_all_positive and evidence.secondary.short_label == "vs removal":
        direction_text = (
            f"Refinement is higher than both baseline and simple removal at Loop {evidence.endpoint_loop}."
        )
    elif evidence.secondary.short_label in {"vs refinement Loop 5", "vs Loop 5"}:
        direction_text = (
            f"Above baseline in {evidence.primary.positive_seeds}/{evidence.primary.n_seeds} seeds; "
            f"above Loop 5 in {evidence.secondary.positive_seeds}/{evidence.secondary.n_seeds}."
        )
    else:
        direction_text = evidence.primary.direction_label.capitalize()
        if secondary_all_positive:
            direction_text += f"; {evidence.secondary.direction_label}."
        else:
            direction_text += "."
    add_text(
        slide,
        "endpoint-direction-note",
        9.25,
        2.36,
        3.05,
        0.48,
        direction_text,
        size=10.5,
        color=INK,
    )

    add_separator(slide, "endpoint-separator-1", 9.25, 2.91, 3.05)
    add_text(
        slide,
        "endpoint-effect-heading",
        9.25,
        3.08,
        3.05,
        0.28,
        "MEAN AUROC CHANGE",
        size=9.5,
        color=MUTED,
        bold=True,
    )
    add_multi_line_text(
        slide,
        "endpoint-effect-values",
        9.25,
        3.42,
        3.05,
        0.72,
        [
            (f"{evidence.primary.mean_delta:+.3f} {evidence.primary.short_label}", 15.5, RED_MARK_RGB, True),
            (f"{evidence.secondary.mean_delta:+.3f} {evidence.secondary.short_label}", 15.5, RED_MARK_RGB, True),
        ],
    )

    add_separator(slide, "endpoint-separator-2", 9.25, 4.27, 3.05)
    add_text(
        slide,
        "endpoint-uncertainty-heading",
        9.25,
        4.44,
        3.05,
        0.28,
        "95% INTERVALS (SEED + STUDY)",
        size=9.5,
        color=MUTED,
        bold=True,
    )
    p_text = (
        f"Exact paired p = {evidence.primary.exact_p:.2f} (both)"
        if math.isclose(evidence.primary.exact_p, evidence.secondary.exact_p)
        else f"Exact paired p = {evidence.primary.exact_p:.2f} / {evidence.secondary.exact_p:.2f}"
    )
    uncertainty_lines = [
        (
            f"{evidence.primary.short_label}: "
            f"[{evidence.primary.ci_low:+.3f}, {evidence.primary.ci_high:+.3f}]",
            9.5,
            INK,
            False,
        ),
        (
            f"{evidence.secondary.short_label}: "
            f"[{evidence.secondary.ci_low:+.3f}, {evidence.secondary.ci_high:+.3f}]",
            9.5,
            INK,
            False,
        ),
        (p_text, 9.5, MUTED, False),
    ]
    add_multi_line_text(
        slide,
        "endpoint-uncertainty-values",
        9.25,
        4.78,
        3.05,
        0.88,
        uncertainty_lines,
    )
    if evidence.proper_score_note:
        add_text(
            slide,
            "endpoint-support-note",
            9.25,
            5.78,
            3.05,
            0.38,
            evidence.proper_score_note,
            size=8.8,
            color=MUTED,
        )

    source_note = (
        "Source: paired MobileNetV3 runs on the same held-out test set; "
        f"seeds {', '.join(map(str, evidence.seeds))}; study-weighted AUROC; mean +/- 1 SD."
    )
    add_text(
        slide,
        "source-note",
        0.62,
        7.02,
        12.1,
        0.34,
        source_note,
        size=7.5,
        color=MUTED,
    )
    slide.notes_slide.notes_text_frame.text = (
        "这页回答老师最关心的 repeatability。左图的浅色轨迹是每个 seed，粗线和误差条是平均值和一个标准差。"
        f"在 Loop {evidence.endpoint_loop}，refinement 相对 baseline 和比较条件的平均 AUROC 变化分别是 "
        f"{evidence.primary.mean_delta:+.4f} 和 {evidence.secondary.mean_delta:+.4f}；方向分别为 "
        f"{evidence.primary.positive_seeds}/{evidence.primary.n_seeds} 和 "
        f"{evidence.secondary.positive_seeds}/{evidence.secondary.n_seeds}。"
        "同时要准确说明 hierarchical intervals 跨过 0，三个 seeds 的 exact paired p-value 分辨率也有限，"
        "所以口头结论是方向一致、结果有希望，而不是已经证明稳定优越。"
        "如果被问到标签支持度，补充说明全标签 AUROC 对极低支持标签敏感；Brier 和 NLL 相对 removal 的结果更一致。"
    )
    return title


RED_MARK_RGB = RGBColor(0xC9, 0x36, 0x2B)


REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def replace_xml_text(data: bytes, replacements: dict[str, str]) -> bytes:
    root = etree.fromstring(data)
    counts = {old: 0 for old in replacements}
    for node in root.xpath(".//a:t", namespaces={"a": DRAWING_NS}):
        if node.text in replacements:
            counts[node.text] += 1
            node.text = replacements[node.text]
    missing = [old for old, count in counts.items() if count != 1]
    if missing:
        raise ValueError(f"Expected exactly one XML text match for: {missing}; counts={counts}")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def statistics_backup_replacements(evidence: EndpointEvidence) -> dict[str, str]:
    if len(evidence.planned_contrasts) != 4:
        raise ValueError(
            f"Expected four planned endpoint contrasts, got {len(evidence.planned_contrasts)}"
        )
    rows = evidence.planned_contrasts
    evidence_text = [
        (
            f"{row.positive_seeds}/{row.n_seeds} positive; "
            f"CI [{row.ci_low:+.3f}, {row.ci_high:+.3f}]; exact p = {row.exact_p:.2f}"
        )
        for row in rows
    ]
    holm_values = {round(row.holm_exact_p, 8) for row in rows}
    if len(holm_values) == 1:
        holm_text = f"Holm-adjusted exact p = {rows[0].holm_exact_p:.2f} for all"
    else:
        holm_text = "No exact p-value survives Holm adjustment"
    return {
        "Backup: how to interpret the final statistics": (
            "Backup: exact three-seed Loop 5/8 endpoint statistics"
        ),
        "Quantity": "Contrast",
        "Plain meaning": "Mean AUROC change",
        "Role in the update": "Direction and uncertainty",
        "Effect size": rows[0].short_label,
        "The average paired change, such as refinement AUROC minus baseline AUROC": (
            f"{rows[0].mean_delta:+.4f}"
        ),
        "Primary: states how large the observed effect is": evidence_text[0],
        "95% confidence interval": rows[1].short_label,
        "A range of effect sizes compatible with seed and study uncertainty": (
            f"{rows[1].mean_delta:+.4f}"
        ),
        "Primary: shows precision and whether material alternatives remain plausible": (
            evidence_text[1]
        ),
        "Exact paired p-value": rows[2].short_label,
        "How surprising the paired seed differences would be under no systematic effect": (
            f"{rows[2].mean_delta:+.4f}"
        ),
        "Supporting: with five seeds, the smallest two-sided exact value is 0.0625": (
            evidence_text[2]
        ),
        "Holm-adjusted p-value": rows[3].short_label,
        "The exact p-value corrected for several pre-specified comparisons": (
            f"{rows[3].mean_delta:+.4f}"
        ),
        "Supporting: limits false positives from multiple testing": evidence_text[3],
        "Mean +/- SD error bar": "Multiplicity",
        "The observed spread across seeds around the mean trajectory": (
            "Four pre-specified contrasts"
        ),
        "Descriptive only: it is not a confidence interval or p-value": holm_text,
        "Effect size and uncertainty lead; p-values are limited-resolution supporting evidence with five seeds.": (
            "Interim n = 3; intervals include seed and shared-study uncertainty."
        ),
        "Source: Locked endpoint statistics plan · Table: read_text · Paired effect sizes, hierarchical confidence intervals, exact sign-flip tests, and Holm adjustment.": (
            "Source: paired MobileNetV3 endpoint evaluation; seeds 7, 13, 42; "
            "Loop 5/8 pre-specified contrasts."
        ),
    }


def replace_note_text(data: bytes, note: str) -> bytes:
    root = etree.fromstring(data)
    body_shapes = root.xpath(
        ".//p:sp[p:nvSpPr/p:nvPr/p:ph[@type='body']]",
        namespaces={
            "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        },
    )
    if len(body_shapes) != 1:
        raise ValueError(f"Expected one notes body shape, got {len(body_shapes)}")
    text_nodes = body_shapes[0].xpath(".//a:t", namespaces={"a": DRAWING_NS})
    if not text_nodes:
        raise ValueError("Notes body does not contain a text node")
    text_nodes[0].text = note
    for node in text_nodes[1:]:
        node.text = ""
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def relationship_part_target(rels_data: bytes, relation_type_suffix: str) -> str:
    root = etree.fromstring(rels_data)
    matches = [
        relation
        for relation in root
        if str(relation.get("Type", "")).endswith(relation_type_suffix)
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one {relation_type_suffix} relationship, got {len(matches)}"
        )
    target = str(matches[0].get("Target"))
    return posixpath.normpath(posixpath.join("ppt/slides", target))


def slide_part_at_position(source: zipfile.ZipFile, position: int) -> tuple[str, str]:
    presentation_root = etree.fromstring(source.read("ppt/presentation.xml"))
    presentation_rels = etree.fromstring(source.read("ppt/_rels/presentation.xml.rels"))
    rel_targets = {
        str(relation.get("Id")): str(relation.get("Target"))
        for relation in presentation_rels
    }
    slide_ids = presentation_root.xpath(
        ".//p:sldId",
        namespaces={
            "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
            "r": OFFICE_REL_NS,
        },
    )
    if not 1 <= position <= len(slide_ids):
        raise IndexError(f"Slide position {position} is outside 1..{len(slide_ids)}")
    relation_id = str(slide_ids[position - 1].get(f"{{{OFFICE_REL_NS}}}id"))
    target = rel_targets.get(relation_id)
    if target is None:
        raise ValueError(f"Missing presentation relationship for {relation_id}")
    slide_part = posixpath.normpath(posixpath.join("ppt", target))
    slide_name = posixpath.basename(slide_part)
    slide_rels = posixpath.join(
        posixpath.dirname(slide_part),
        "_rels",
        f"{slide_name}.rels",
    )
    return slide_part, slide_rels


def prepare_slide_rels_and_xml(
    base_rels_data: bytes,
    temp_slide_data: bytes,
    media_name: str,
) -> tuple[bytes, bytes, str]:
    rels_root = etree.fromstring(base_rels_data)
    image_relations = [
        relation
        for relation in rels_root
        if str(relation.get("Type", "")).endswith("/image")
    ]
    if image_relations:
        image_relation = image_relations[0]
    else:
        used_ids = {str(relation.get("Id")) for relation in rels_root}
        next_number = 1
        while f"rId{next_number}" in used_ids:
            next_number += 1
        image_relation = etree.SubElement(rels_root, f"{{{REL_NS}}}Relationship")
        image_relation.set("Id", f"rId{next_number}")
        image_relation.set("Type", f"{OFFICE_REL_NS}/image")
    image_relation.set("Target", f"../media/{media_name}")
    image_rid = str(image_relation.get("Id"))

    notes_targets = [
        posixpath.normpath(posixpath.join("ppt/slides", str(relation.get("Target"))))
        for relation in rels_root
        if str(relation.get("Type", "")).endswith("/notesSlide")
    ]
    if len(notes_targets) != 1:
        raise ValueError(f"Expected one slide-6 notes target, got {notes_targets}")

    slide_root = etree.fromstring(temp_slide_data)
    blips = slide_root.xpath(
        ".//a:blip",
        namespaces={"a": DRAWING_NS, "r": OFFICE_REL_NS},
    )
    if len(blips) != 1:
        raise ValueError(f"Expected one chart image in temporary slide, got {len(blips)}")
    blips[0].set(f"{{{OFFICE_REL_NS}}}embed", image_rid)
    return (
        etree.tostring(slide_root, xml_declaration=True, encoding="UTF-8", standalone=True),
        etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone=True),
        notes_targets[0],
    )


def ensure_png_content_type(data: bytes) -> bytes:
    root = etree.fromstring(data)
    defaults = root.findall(f"{{{CONTENT_TYPES_NS}}}Default")
    if any(str(node.get("Extension", "")).lower() == "png" for node in defaults):
        return data
    node = etree.SubElement(root, f"{{{CONTENT_TYPES_NS}}}Default")
    node.set("Extension", "png")
    node.set("ContentType", "image/png")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def build_temporary_slide(
    input_deck: Path,
    chart_path: Path,
    evidence: EndpointEvidence,
    temp_deck: Path,
) -> str:
    base = Presentation(input_deck)
    presentation = Presentation()
    presentation.slide_width = base.slide_width
    presentation.slide_height = base.slide_height
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        presentation.slide_width,
        presentation.slide_height,
    )
    background.name = "slide-background"
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(0xFC, 0xFC, 0xFD)
    background.line.color.rgb = RGBColor(0xFC, 0xFC, 0xFD)
    title = replace_slide_six(slide, chart_path, evidence)
    presentation.save(temp_deck)
    return title


def transplant_slide_and_cover(
    input_deck: Path,
    temp_deck: Path,
    chart_path: Path,
    sample_chart_path: Path,
    synthetic_chart_path: Path,
    output_deck: Path,
    evidence: EndpointEvidence,
    notes_by_position: dict[int, str],
) -> None:
    media_name = "endpoint_performance_slide.png"
    with zipfile.ZipFile(input_deck) as source, zipfile.ZipFile(temp_deck) as temp:
        slide6_part, slide6_rels_part = slide_part_at_position(source, 6)
        base_slide6_rels = source.read(slide6_rels_part)
        slide6_xml, slide6_rels, _slide6_notes_target = prepare_slide_rels_and_xml(
            base_slide6_rels,
            temp.read("ppt/slides/slide1.xml"),
            media_name,
        )
        slide2_part, _ = slide_part_at_position(source, 2)
        slide3_part, slide3_rels_part = slide_part_at_position(source, 3)
        synthetic_chart_part = relationship_part_target(
            source.read(slide3_rels_part),
            "/image",
        )
        slide4_part, _ = slide_part_at_position(source, 4)
        slide5_part, slide5_rels_part = slide_part_at_position(source, 5)
        sample_chart_part = relationship_part_target(
            source.read(slide5_rels_part),
            "/image",
        )
        slide7_part, _ = slide_part_at_position(source, 7)
        slide8_part, _ = slide_part_at_position(source, 8)
        slide9_part, _ = slide_part_at_position(source, 9)
        replacements: dict[str, bytes] = {
            "ppt/slides/slide1.xml": replace_xml_text(
                source.read("ppt/slides/slide1.xml"),
                {
                    "Current evaluation evidence | 15 July 2026": "Current evaluation evidence | 17 July 2026",
                    "Known-truth validation, denominator-aware quality measures, and sample-level evidence; final downstream Loop 5/8 endpoints remain pending.": (
                        "Known-truth validation, denominator-aware quality measures, sample-level evidence, "
                        "and repeated-run performance."
                    ),
                    "Real-data seeds": "Quality diagnostic seeds",
                    "Paired seeds in the current entry- and sample-level diagnostics.": (
                        "Five paired seeds in the matched Loop 5 entry- and sample-level diagnostics."
                    ),
                },
            ),
            slide2_part: replace_xml_text(
                source.read(slide2_part),
                {
                    "Raw and adjusted metrics answer different denominator questions": (
                        "Exact definitions of raw and coverage-adjusted DQS"
                    ),
                    "Metric": "Term",
                    "Plain definition": "Exact definition",
                    "Question answered": "Interpretation",
                    "Estimated clean fraction among entries that remain valid now": (
                        "1 - M_t / |E_t|"
                    ),
                    "How clean are the surviving entries?": (
                        "Of the entries still valid, what fraction is estimated clean?"
                    ),
                    "Current valid entries divided by baseline valid entries": (
                        "|E_t| / |E_0|"
                    ),
                    "How much of the original entry denominator remains?": (
                        "What fraction of the original valid entries remains?"
                    ),
                    "Adjusted entry DQS": "Coverage-adjusted entry DQS",
                    "Raw entry DQS multiplied by valid-entry coverage": (
                        "Raw entry DQS × coverage = (|E_t| - M_t) / |E_0|"
                    ),
                    "How much estimated healthy entry mass remains versus baseline?": (
                        "Of the original valid entries, what fraction remains and is estimated clean?"
                    ),
                    "Sample issue-free rate": "Raw sample DQS",
                    "Fraction of evaluable samples with no valid entry flagged by confident learning": (
                        "1 - Q_t / |S_t|"
                    ),
                    "How clean are the currently evaluable samples?": (
                        "Of the samples still evaluable, what fraction has no flagged entry?"
                    ),
                    "Adjusted sample health": "Coverage-adjusted sample DQS",
                    "Sample issue-free rate multiplied by evaluable-sample coverage": (
                        "Raw sample DQS × (|S_t| / |S_0|) = (|S_t| - Q_t) / |S_0|"
                    ),
                    "How much healthy sample mass remains versus baseline?": (
                        "Of the original evaluable samples, what fraction remains with no flagged entry?"
                    ),
                    "Raw metrics describe current survivors; adjusted metrics retain the baseline denominator.": (
                        "Notation: E_t = valid entries after Loop t; M_t = issues estimated in E_t by "
                        "confident learning from K-fold OOF predictions; S_t = samples with >=1 valid "
                        "entry; Q_t = samples in S_t with >=1 valid entry flagged."
                    ),
                    "Source: Quality metric definitions · Table: read_json_auto · Definitions, denominator identities, seed set, and error-bar convention.": (
                        "Metric protocol: issue counts are estimated by confident learning from K-fold "
                        "out-of-fold predictions."
                    ),
                },
            ),
            slide3_part: replace_xml_text(
                source.read(slide3_part),
                {
                    "Raw DQS follows survivor cleanliness; adjusted DQS follows healthy mass against the original denominator. Matching-target MAE is 0.019 with informative OOF evidence and 0.055 with weak evidence.": (
                        "Raw DQS follows accuracy among labels that remain; coverage-adjusted DQS follows "
                        "the fraction of original labels that remain and are correct. "
                        "Matching-target MAE is 0.019 with informative OOF evidence and 0.055 "
                        "with weak evidence."
                    ),
                },
            ),
            synthetic_chart_part: synthetic_chart_path.read_bytes(),
            slide4_part: replace_xml_text(
                source.read(slide4_part),
                {
                    "Removal cleans survivors but loses healthy entry mass": (
                        "Removal raises raw DQS but lowers coverage-adjusted DQS"
                    ),
                    "Refinement reaches raw/adjusted DQS 0.937/0.933 while retaining 99.6% entry coverage.": (
                        "Refinement reaches raw / coverage-adjusted entry DQS 0.937 / 0.933 "
                        "while retaining 99.6% entry coverage."
                    ),
                    "Five-seed entry-level raw and adjusted DQS with individual seeds and standard deviations": (
                        "Five-seed raw and coverage-adjusted entry DQS with individual seeds "
                        "and standard deviations"
                    ),
                },
            ),
            slide5_part: replace_xml_text(
                source.read(slide5_part),
                {
                    "The same denominator effect appears at sample level": (
                        "Sample DQS shows the same denominator effect"
                    ),
                    "Context": "SAMPLE-LEVEL DQS",
                    "Refinement reaches 0.756 raw and 0.754 adjusted health while retaining 99.7% sample coverage.": (
                        "Refinement reaches 0.756 raw sample DQS and 0.754 coverage-adjusted "
                        "sample DQS while retaining 99.7% sample coverage."
                    ),
                    "Five-seed sample-level issue-free rate and adjusted health with individual seeds and standard deviations": (
                        "Five-seed raw and coverage-adjusted sample DQS with individual seeds "
                        "and standard deviations"
                    ),
                },
            ),
            sample_chart_part: sample_chart_path.read_bytes(),
            slide6_part: slide6_xml,
            slide6_rels_part: slide6_rels,
            slide7_part: replace_xml_text(
                source.read(slide7_part),
                {
                    "Entry raw": "Raw entry DQS",
                    "Entry adjusted": "Coverage-adjusted entry DQS",
                    "Sample raw": "Raw sample DQS",
                    "Sample adjusted": "Coverage-adjusted sample DQS",
                    "Loop 5 means; adjusted quantities use the corresponding baseline denominator.": (
                        "Loop 5 means; coverage-adjusted DQS uses the corresponding baseline denominator."
                    ),
                },
            ),
            slide8_part: replace_xml_text(
                source.read(slide8_part),
                {
                    "Truth: covered": "Accuracy among retained labels",
                    "Truth: original N": "Fraction retained and correct",
                    "Adjusted DQS": "Coverage-adjusted DQS",
                    "Informative OOF evidence; means across 10 synthetic seeds.": (
                        "Truth columns are fractions; retained-and-correct uses the original denominator."
                    ),
                },
            ),
            slide9_part: replace_xml_text(
                source.read(slide9_part),
                statistics_backup_replacements(evidence),
            ),
            "[Content_Types].xml": ensure_png_content_type(source.read("[Content_Types].xml")),
        }
        expected_positions = set(range(1, 10))
        if set(notes_by_position) != expected_positions:
            raise ValueError(
                f"Expected notes for slide positions 1..9, got {sorted(notes_by_position)}"
            )
        for position, note in notes_by_position.items():
            _, slide_rels_part = slide_part_at_position(source, position)
            notes_target = relationship_part_target(
                source.read(slide_rels_part),
                "/notesSlide",
            )
            replacements[notes_target] = replace_note_text(
                source.read(notes_target),
                note,
            )

        output_deck.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            prefix=output_deck.stem + "_",
            suffix=".pptx",
            dir=output_deck.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        try:
            with zipfile.ZipFile(
                temporary_path,
                mode="w",
                compression=zipfile.ZIP_DEFLATED,
            ) as target:
                for info in source.infolist():
                    data = replacements.get(info.filename, source.read(info.filename))
                    target.writestr(info, data)
                target.writestr(f"ppt/media/{media_name}", chart_path.read_bytes())
            os.replace(temporary_path, output_deck)
            os.chmod(output_deck, 0o644)
        finally:
            if temporary_path.exists():
                temporary_path.unlink()


def update_plan(input_plan: Path, output_plan: Path, title: str, evidence: EndpointEvidence) -> None:
    plan = json.loads(input_plan.read_text(encoding="utf-8"))
    if len(plan) != 9:
        raise ValueError(f"Expected 9 plan entries, got {len(plan)}")
    plan[1]["title"] = "Exact definitions of raw and coverage-adjusted DQS"
    plan[2]["source"] = "Known-truth synthetic DQS with plain-language truth labels"
    plan[3]["title"] = "Removal raises raw DQS but lowers coverage-adjusted DQS"
    plan[4]["title"] = "Sample DQS shows the same denominator effect"
    plan[4]["source"] = "Five-seed raw and coverage-adjusted sample DQS"
    plan[5] = {
        "kind": "chart_evidence",
        "title": title,
        "chart_index": 5,
        "source_kind": "embedded_png",
        "source": "Three-seed paired endpoint evaluation",
        "elements": [
            "slide-title",
            "endpoint-chart",
            "direction",
            "mean-effect",
            "uncertainty",
            "source-note",
        ],
        "slide_number": 6,
        "section": "main",
        "endpoint_loop": evidence.endpoint_loop,
        "seeds": evidence.seeds,
    }
    plan[8] = {
        "kind": "table",
        "title": "Backup: exact three-seed Loop 5/8 endpoint statistics",
        "table_index": 4,
        "source": "Three-seed paired endpoint evaluation with hierarchical intervals",
        "elements": [
            "slide-title",
            "table-main",
            "table-note",
            "source-note",
        ],
        "slide_number": 9,
        "section": "backup",
        "endpoint_loop": evidence.endpoint_loop,
        "seeds": evidence.seeds,
        "comparisons": [row.name for row in evidence.planned_contrasts],
    }
    output_plan.write_text(json.dumps(plan, indent=2) + "\n", encoding="utf-8")


def validate_deck(
    path: Path,
    title: str,
    evidence: EndpointEvidence,
    notes_by_position: dict[int, str],
) -> dict[str, object]:
    presentation = Presentation(path)
    if len(presentation.slides) != 9:
        raise ValueError(f"Expected 9 slides, got {len(presentation.slides)}")
    slide_six = presentation.slides[5]
    texts = [shape.text for shape in slide_six.shapes if getattr(shape, "has_text_frame", False)]
    if title not in texts:
        raise ValueError("Updated slide 6 title is missing")
    if not any(shape.shape_type == 13 for shape in slide_six.shapes):
        raise ValueError("Updated slide 6 chart image is missing")
    cover_text = " ".join(
        shape.text
        for shape in presentation.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    if "Current evaluation evidence | 17 July 2026" not in cover_text:
        raise ValueError("Updated cover date is missing")
    slide_two_text = " ".join(
        cell.text
        for shape in presentation.slides[1].shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    )
    for required in [
        "Term",
        "Exact definition",
        "Interpretation",
        "1 - M_t / |E_t|",
        "|E_t| / |E_0|",
        "Coverage-adjusted entry DQS",
        "1 - Q_t / |S_t|",
        "Raw sample DQS",
        "Coverage-adjusted sample DQS",
    ]:
        if required not in slide_two_text:
            raise ValueError(f"Slide 2 is missing plain DQS wording: {required}")
    slide_three_text = " ".join(
        shape.text
        for shape in presentation.slides[2].shapes
        if getattr(shape, "has_text_frame", False)
    )
    if "fraction of original labels that remain and are correct" not in slide_three_text:
        raise ValueError("Slide 3 is missing the plain original-denominator definition")
    slide_four_text = " ".join(
        shape.text
        for shape in presentation.slides[3].shapes
        if getattr(shape, "has_text_frame", False)
    )
    if "Removal raises raw DQS but lowers coverage-adjusted DQS" not in slide_four_text:
        raise ValueError("Slide 4 is missing the direct denominator title")
    slide_five = presentation.slides[4]
    slide_five_text = " ".join(
        shape.text
        for shape in slide_five.shapes
        if getattr(shape, "has_text_frame", False)
    )
    for required in [
        "Sample DQS shows the same denominator effect",
        "0.756 raw sample DQS",
        "0.754 coverage-adjusted sample DQS",
    ]:
        if required not in slide_five_text:
            raise ValueError(f"Slide 5 is missing sample-DQS wording: {required}")
    slide_seven = presentation.slides[6]
    slide_seven_text = " ".join(
        cell.text
        for shape in slide_seven.shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    )
    for required in ["Raw sample DQS", "Coverage-adjusted sample DQS"]:
        if required not in slide_seven_text:
            raise ValueError(f"Slide 7 is missing sample-DQS wording: {required}")
    slide_eight_text = " ".join(
        cell.text
        for shape in presentation.slides[7].shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    )
    for required in ["Fraction retained and correct", "Coverage-adjusted DQS"]:
        if required not in slide_eight_text:
            raise ValueError(f"Slide 8 is missing aligned terminology: {required}")
    slide_nine = presentation.slides[8]
    slide_nine_text = " ".join(
        cell.text
        for shape in slide_nine.shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    )
    for required in [
        "Loop 5 refine vs baseline",
        f"Loop {evidence.endpoint_loop} refine vs baseline",
        f"Loop {evidence.endpoint_loop} refine vs Loop 5",
        "Loop 5 refine vs removal",
    ]:
        if required not in slide_nine_text:
            raise ValueError(f"Slide 9 is missing contrast: {required}")
    stale = " ".join(
        shape.text
        for slide in list(presentation.slides)[:6]
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    if "final downstream Loop 5/8 endpoints remain pending" in stale:
        raise ValueError("Stale pending-endpoint text remains in the main deck")
    visible_and_notes = "\n".join(
        [
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
        ]
        + [slide.notes_slide.notes_text_frame.text for slide in presentation.slides]
    )
    stale_patterns = {
        "sample issue terminology": r"sample (?:issue rate|issue-free rate|health)",
        "unqualified adjusted DQS": r"(?<!coverage-)adjusted DQS",
        "unqualified adjusted sample DQS": r"(?<!coverage-)adjusted sample DQS",
        "raw/adjusted shorthand": r"raw\s*(?:/|and)\s*adjusted",
    }
    for label, pattern in stale_patterns.items():
        if re.search(pattern, visible_and_notes, flags=re.IGNORECASE):
            raise ValueError(f"Deck contains stale {label}")
    for slide_index, slide in enumerate(presentation.slides, start=1):
        actual_note = slide.notes_slide.notes_text_frame.text.strip()
        if actual_note != notes_by_position[slide_index]:
            raise ValueError(f"Speaker note mismatch on slide {slide_index}")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                raise ValueError(f"Negative shape position on slide {slide_index}: {shape.name}")
            if shape.left + shape.width > presentation.slide_width + 2:
                raise ValueError(f"Shape exceeds slide width on slide {slide_index}: {shape.name}")
            if shape.top + shape.height > presentation.slide_height + 2:
                raise ValueError(f"Shape exceeds slide height on slide {slide_index}: {shape.name}")
    with zipfile.ZipFile(path) as archive:
        corrupt_member = archive.testzip()
    if corrupt_member is not None:
        raise ValueError(f"Corrupt PPTX member: {corrupt_member}")
    backup_titles = []
    for slide in list(presentation.slides)[6:]:
        slide_texts = [
            shape.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        ]
        backup_titles.append(slide_texts[0] if slide_texts else "")
    return {
        "slides": len(presentation.slides),
        "main_slides": 6,
        "backup_slides": 3,
        "backup_titles": backup_titles,
        "pptx_zip_test": "passed",
        "shape_bounds": "passed",
        "slide_6_picture_count": sum(shape.shape_type == 13 for shape in slide_six.shapes),
        "slide_3_plain_truth_labels": "passed",
        "slide_5_sample_dqs_labels": "passed",
        "slide_9_contrast_rows": 4,
        "speaker_notes": "present" if slide_six.notes_slide.notes_text_frame.text.strip() else "missing",
    }


def main() -> int:
    args = parse_args()
    for path in [
        args.input_deck,
        args.endpoint_dir,
        args.sample_chart,
        args.synthetic_chart,
        REPORT_DIR / "deck_plan.json",
    ]:
        if not path.exists():
            raise FileNotFoundError(path)
    evidence = load_endpoint_evidence(args.endpoint_dir)
    notes_by_position = build_speaker_notes(evidence)
    build_trajectory_chart(evidence, args.chart_output)

    with tempfile.TemporaryDirectory(prefix="endpoint-slide-") as temp_dir:
        temp_deck = Path(temp_dir) / "endpoint_slide.pptx"
        title = build_temporary_slide(
            args.input_deck,
            args.chart_output,
            evidence,
            temp_deck,
        )
        transplant_slide_and_cover(
            args.input_deck,
            temp_deck,
            args.chart_output,
            args.sample_chart,
            args.synthetic_chart,
            args.output_deck,
            evidence,
            notes_by_position,
        )

    update_plan(REPORT_DIR / "deck_plan.json", args.plan_output, title, evidence)
    validation = validate_deck(args.output_deck, title, evidence, notes_by_position)
    summary = {
        "status": "passed",
        "input_deck": str(args.input_deck),
        "endpoint_dir": str(args.endpoint_dir),
        "endpoint_loop": evidence.endpoint_loop,
        "seeds": evidence.seeds,
        "chart": str(args.chart_output),
        "sample_chart": str(args.sample_chart),
        "synthetic_chart": str(args.synthetic_chart),
        "output_deck": str(args.output_deck),
        "plan": str(args.plan_output),
        "slide_6_title": title,
        "primary_mean_delta": evidence.primary.mean_delta,
        "secondary_mean_delta": evidence.secondary.mean_delta,
        "planned_contrasts": [
            {
                "name": row.name,
                "mean_delta": row.mean_delta,
                "positive_seeds": row.positive_seeds,
                "n_seeds": row.n_seeds,
                "ci": [row.ci_low, row.ci_high],
                "exact_p": row.exact_p,
                "holm_exact_p": row.holm_exact_p,
            }
            for row in evidence.planned_contrasts
        ],
        "validation": validation,
    }
    args.summary_output.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
