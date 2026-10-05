#!/usr/bin/env python3
"""Build the four-slide 2026-08-07 meeting update from the locked outline."""

from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
SOURCE_MD = HERE / "next_meeting_slides_20260807_structure_中文.md"
OUTPUT_DIR = HERE / "next_meeting_slides_20260807"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_PPTX = OUTPUT_DIR / "next_meeting_update_20260807.pptx"
MANIFEST = OUTPUT_DIR / "delivery_manifest.json"
FLOWCHART_IMAGE = HERE / "slide_assets" / "pipeline_flowchart.png"

MIMIC_IMAGE = Path(
    "/vol/gpudata/yz3522-llmtest/MResProject/MedSoul/datasets/"
    "mimic-cxr-jpg-224/p12/p12489885/s57337262/"
    "b4a3902d-f2651574-6f46993b-7f21d467-6b820335.jpg"
)
MEDPALM_IMAGE = Path(
    "/vol/gpudata/yz3522-llmtest/MResProject/MedSoul/datasets/"
    "mimic-cxr-jpg-224/p11/p11569093/s50008596/"
    "2f108c10-c8669b9a-f7f02e0f-272d2904-dd0b345e.jpg"
)
VINDR_IMAGE = Path(
    "/vol/bitbucket/yz3522/datasets/vindr-cxr/1.0.0/derived/"
    "png224_20260804_v1/images/004f33259ee4aef671c2b95d54e4be68.png"
)

VINDR_ONE_SHOT = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "vindr_global_symmetric_noise_mobilenet/20260805_v1/"
    "evaluation_formal_v1/budget_metrics.csv"
)
MEDPALM_ONE_SHOT = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "medpalm_cl_detection_benchmark/20260803_frozen4seed/"
    "medpalm_cl_detection_metrics.csv"
)
VINDR_ITERATIVE = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "vindr_iterative_oracle_cleaning/20260806_combined_8seed_v1/"
    "evaluation_formal_v1/dynamic_loop_summary.csv"
)
MEDPALM_ITERATIVE = Path(
    "/vol/bitbucket/yz3522/NoiseRate_results_archive/cxr_real_experiment/"
    "medpalm_cl_iterative_oracle_benchmark/20260803_seed13_dynamic5loop/"
    "iterative_detection_trajectory.csv"
)
VINDR_TAU035_DQS = (
    HERE
    / "next_meeting_slides_20260807"
    / "extra_figures"
    / "vindr_entry_dqs_tau035_curve.csv"
)

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Times New Roman"


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.lstrip("#")
    return RGBColor(int(value[:2], 16), int(value[2:4], 16), int(value[4:], 16))


INK = rgb("#18212F")
MUTED = rgb("#667085")
LIGHT_MUTED = rgb("#98A2B3")
BLUE = rgb("#0077BB")
BLUE_DARK = rgb("#24557A")
BLUE_LIGHT = rgb("#E8F3F9")
RED = rgb("#CC3311")
RED_LIGHT = rgb("#FAECE8")
TEAL = rgb("#009988")
TEAL_LIGHT = rgb("#E5F4F1")
GREY = rgb("#6B7280")
GREY_LIGHT = rgb("#EEF1F5")
GRID = rgb("#D9DEE7")
PANEL = rgb("#F7F8FA")
WHITE = rgb("#FFFFFF")

HEX_INK = "#18212F"
HEX_MUTED = "#667085"
HEX_BLUE = "#0077BB"
HEX_RED = "#CC3311"
HEX_TEAL = "#009988"
HEX_GREY = "#6B7280"
HEX_GRID = "#D9DEE7"

TITLES = [
    "From Scalable Labelling to Evidence-Based Dataset Refinement",
    "How Can We Know Whether Confident Learning Finds Real Label Errors?",
    "CL-Guided Review Finds More True Errors Than Random Review",
    "Iterative CL-Guided Review Improves Quality With Less Review",
]


def configure_plot_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": [FONT, "Liberation Serif", "DejaVu Serif"],
            "font.size": 11,
            "axes.titlesize": 13,
            "axes.labelsize": 11,
            "xtick.labelsize": 9.5,
            "ytick.labelsize": 9.5,
            "legend.fontsize": 9.5,
            "figure.dpi": 180,
            "savefig.dpi": 220,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.labelcolor": HEX_MUTED,
            "xtick.color": HEX_MUTED,
            "ytick.color": HEX_MUTED,
        }
    )


def add_box(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    fill: RGBColor = WHITE,
    line: RGBColor = GRID,
    line_width: float = 0.8,
    radius: bool = False,
    name: str | None = None,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def add_text(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    *,
    size: float = 14,
    color: RGBColor = INK,
    bold: bool = False,
    italic: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.0,
    name: str | None = None,
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return shape


def add_lines(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    lines: list[dict[str, Any]],
    *,
    margin: float = 0.08,
    name: str | None = None,
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    for idx, item in enumerate(lines):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.alignment = item.get("align", PP_ALIGN.LEFT)
        paragraph.space_before = Pt(item.get("space_before", 0))
        paragraph.space_after = Pt(item.get("space_after", 3))
        run = paragraph.add_run()
        run.text = item["text"]
        run.font.name = FONT
        run.font.size = Pt(item.get("size", 13))
        run.font.bold = item.get("bold", False)
        run.font.italic = item.get("italic", False)
        run.font.color.rgb = item.get("color", INK)
    return shape


def add_rule(
    slide,
    x: float,
    y: float,
    width: float,
    *,
    color: RGBColor = GRID,
    line_width: float = 0.8,
):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x),
        Inches(y),
        Inches(x + width),
        Inches(y),
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(line_width)
    return shape


def add_picture_contain(
    slide,
    path: Path,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    name: str | None = None,
):
    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = width / height
    if ratio >= box_ratio:
        render_w = width
        render_h = width / ratio
        render_x = x
        render_y = y + (height - render_h) / 2
    else:
        render_h = height
        render_w = height * ratio
        render_x = x + (width - render_w) / 2
        render_y = y
    picture = slide.shapes.add_picture(
        str(path),
        Inches(render_x),
        Inches(render_y),
        width=Inches(render_w),
        height=Inches(render_h),
    )
    if name:
        picture.name = name
    return picture


def new_slide(prs: Presentation, page: int, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_text(
        slide,
        0.55,
        0.25,
        11.85,
        0.55,
        title,
        size=24.5,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
        name="slide-title",
    )
    add_text(
        slide,
        12.42,
        0.34,
        0.34,
        0.22,
        f"{page:02d}",
        size=9.5,
        color=LIGHT_MUTED,
        align=PP_ALIGN.RIGHT,
        name="page-number",
    )
    add_rule(slide, 0.55, 0.91, 12.2)
    return slide


def add_takeaway(slide, text: str, *, color: RGBColor = BLUE_DARK) -> None:
    add_rule(slide, 0.62, 6.62, 12.08, color=color, line_width=1.6)
    add_text(
        slide,
        0.72,
        6.72,
        11.9,
        0.48,
        text,
        size=13.2,
        color=color,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
        name="takeaway",
    )


def parse_notes() -> dict[int, str]:
    source = SOURCE_MD.read_text(encoding="utf-8")
    sections = re.findall(
        r"^## Slide (\d+) - .*?\n(.*?)(?=^---\n\n## Slide |^---\n\n## \u5236作约束|\Z)",
        source,
        flags=re.MULTILINE | re.DOTALL,
    )
    notes: dict[int, str] = {}
    for page_text, section in sections:
        match = re.search(
            r"^### Speaker notes.*?\n\n(.*?)(?=^### |\Z)",
            section,
            flags=re.MULTILINE | re.DOTALL,
        )
        if not match:
            raise ValueError(f"Missing speaker notes for slide {page_text}")
        lines = []
        for line in match.group(1).strip().splitlines():
            if line == ">":
                lines.append("")
            elif line.startswith("> "):
                lines.append(line[2:])
            else:
                lines.append(line)
        notes[int(page_text)] = "\n".join(lines).strip()
    if set(notes) != {1, 2, 3, 4}:
        raise ValueError(f"Unexpected note pages: {sorted(notes)}")
    return notes


def set_notes(slide, note: str) -> None:
    frame = slide.notes_slide.notes_text_frame
    frame.text = note
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(12)


def load_one_shot() -> tuple[pd.DataFrame, dict[str, float]]:
    vindr = pd.read_csv(VINDR_ONE_SHOT)
    selected = vindr[
        (vindr["method"] == "cl_first")
        & (vindr["budget_name"] == "true_error_count")
        & (vindr["rate_percent"].isin([10, 20, 30]))
    ].copy()
    if selected["seed"].nunique() != 6:
        raise AssertionError("VinDr one-shot results must contain six seeds")
    grouped = (
        selected.groupby("rate_percent")
        .agg(
            injected_errors=("budget", "first"),
            tp_mean=("tp", "mean"),
            tp_sd=("tp", "std"),
            recall_mean=("recall", "mean"),
            recall_sd=("recall", "std"),
            random_mean=("random_expected_tp", "first"),
        )
        .reset_index()
        .sort_values("rate_percent")
    )
    med = pd.read_csv(MEDPALM_ONE_SHOT)
    row = med[
        (med["cohort"] == "all_entries")
        & (med["estimator"] == "ensemble")
        & (med["method"] == "self_percentile")
    ].iloc[0]
    med_values = {
        "n_entries": float(row["n_entries"]),
        "true_issues": float(row["true_issues"]),
        "review_budget": float(row["top20_n"]),
        "issues_found": float(row["top20_true_issues"]),
        "issue_recall": float(row["top20_recall"]),
        "random_expected": float(row["true_issues"] * row["top20_n"] / row["n_entries"]),
        "random_recall": float(row["top20_n"] / row["n_entries"]),
    }
    return grouped, med_values


def plot_vindr_one_shot(path: Path, data: pd.DataFrame) -> None:
    labels = [
        f"{int(row.rate_percent)}% noise\n{int(row.injected_errors):,} reviews"
        for row in data.itertuples()
    ]
    cl = data["recall_mean"].to_numpy() * 100
    sd = data["recall_sd"].to_numpy() * 100
    random = data["rate_percent"].to_numpy(dtype=float)
    x = np.arange(len(labels))
    width = 0.34
    fig, ax = plt.subplots(figsize=(8.6, 4.8))
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    cl_bars = ax.bar(
        x - width / 2,
        cl,
        width,
        yerr=sd,
        capsize=4,
        color=HEX_RED,
        edgecolor=HEX_INK,
        linewidth=0.7,
        label="CL-guided review",
    )
    random_bars = ax.bar(
        x + width / 2,
        random,
        width,
        color="#D7DCE3",
        edgecolor=HEX_GREY,
        linewidth=0.7,
        hatch="///",
        label="Random expectation",
    )
    for idx, (cl_bar, random_bar, row) in enumerate(
        zip(cl_bars, random_bars, data.itertuples())
    ):
        ax.text(
            cl_bar.get_x() + cl_bar.get_width() / 2,
            cl_bar.get_height() + sd[idx] + 2.1,
            f"{row.tp_mean:,.0f}/{int(row.injected_errors):,}\n{row.recall_mean * 100:.1f}%",
            ha="center",
            va="bottom",
            color=HEX_RED,
            fontsize=9.2,
            fontweight="bold",
        )
        ax.text(
            random_bar.get_x() + random_bar.get_width() / 2,
            random_bar.get_height() + 1.2,
            f"{row.random_mean:,.0f}\n{row.rate_percent:.0f}%",
            ha="center",
            va="bottom",
            color=HEX_GREY,
            fontsize=8.7,
        )
    ax.set_title("VinDr-CXR: Known Injected Errors", loc="left", fontweight="bold")
    ax.text(
        1.0,
        1.02,
        "Mean +/- SD across 6 seeds",
        transform=ax.transAxes,
        ha="right",
        va="bottom",
        color=HEX_MUTED,
        fontsize=9,
    )
    ax.set_ylabel("True errors found (%)")
    ax.set_xticks(x, labels)
    ax.set_ylim(0, 104)
    ax.set_yticks(np.arange(0, 101, 20))
    ax.grid(axis="y", color=HEX_GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.legend(loc="upper left", frameon=False, ncol=2)
    fig.tight_layout(pad=0.8)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def plot_medpalm_one_shot(path: Path, data: dict[str, float]) -> None:
    values = [data["issue_recall"] * 100, data["random_recall"] * 100]
    colors = [HEX_RED, "#D7DCE3"]
    fig, ax = plt.subplots(figsize=(4.4, 4.8))
    fig.patch.set_facecolor("white")
    bars = ax.bar(
        [0, 1],
        values,
        width=0.56,
        color=colors,
        edgecolor=[HEX_INK, HEX_GREY],
        linewidth=0.8,
        hatch=[None, "///"],
    )
    labels = [
        f"{data['issues_found']:.0f}/{data['true_issues']:.0f}\n{values[0]:.1f}%",
        f"{data['random_expected']:.1f}/{data['true_issues']:.0f}\n{values[1]:.1f}%",
    ]
    for idx, bar in enumerate(bars):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + 2.2,
            labels[idx],
            ha="center",
            va="bottom",
            fontsize=10,
            fontweight="bold" if idx == 0 else "normal",
            color=HEX_RED if idx == 0 else HEX_GREY,
        )
    ax.set_title("Med-PaLM: 100/498 Entries Reviewed", loc="left", fontweight="bold")
    ax.set_ylabel("Expert-confirmed issues found (%)")
    ax.set_xticks([0, 1], ["CL-guided", "Random expected"])
    ax.set_ylim(0, 70)
    ax.set_yticks(np.arange(0, 71, 10))
    ax.grid(axis="y", color=HEX_GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    fig.tight_layout(pad=0.8)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def plot_vindr_quality(
    path: Path,
    data: pd.DataFrame,
    dqs_data: pd.DataFrame,
) -> None:
    loops = data["loop"].to_numpy()
    fig, ax = plt.subplots(figsize=(8.7, 3.15))
    fig.patch.set_facecolor("white")
    ax.errorbar(
        loops,
        data["true_quality_mean"],
        yerr=data["true_quality_sd"],
        color=HEX_BLUE,
        marker="o",
        linewidth=2.2,
        markersize=4.5,
        capsize=3,
        label="Known true quality",
    )
    ax.errorbar(
        loops,
        dqs_data["thresholded_dqs_mean"],
        yerr=dqs_data["thresholded_dqs_sd"],
        color=HEX_RED,
        marker="s",
        linewidth=2.1,
        markersize=4.3,
        capsize=3,
        linestyle="--",
        label="Raw DQS",
    )
    ax.text(
        8.08,
        data.iloc[-1]["true_quality_mean"] - 0.004,
        f"True quality {data.iloc[-1]['true_quality_mean']:.3f}",
        color=HEX_BLUE,
        fontsize=9,
        fontweight="bold",
        va="top",
    )
    ax.text(
        8.08,
        dqs_data.iloc[-1]["thresholded_dqs_mean"] + 0.004,
        f"Raw DQS {dqs_data.iloc[-1]['thresholded_dqs_mean']:.3f}",
        color=HEX_RED,
        fontsize=9,
        fontweight="bold",
        va="bottom",
    )
    ax.set_title("VinDr: Complete Dataset Quality", loc="left", fontweight="bold")
    ax.text(
        1,
        1.02,
        "Mean +/- SD across 8 seeds",
        transform=ax.transAxes,
        ha="right",
        color=HEX_MUTED,
        fontsize=8.8,
    )
    ax.set_xlim(-0.15, 9.2)
    ax.set_ylim(0.76, 0.98)
    ax.set_xticks(loops)
    ax.set_ylabel("Quality score")
    ax.grid(axis="y", color=HEX_GRID, linewidth=0.8)
    ax.legend(loc="lower right", frameon=False, ncol=2)
    fig.tight_layout(pad=0.65)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def plot_vindr_efficiency(path: Path, data: pd.DataFrame) -> None:
    loops = data["loop"].to_numpy()
    reviewed = data["reviews_mean"].to_numpy() / 18000 * 100
    recall = data["recall_mean"].to_numpy() * 100
    recall_sd = data["recall_sd"].to_numpy() * 100
    fig, ax = plt.subplots(figsize=(8.7, 3.15))
    fig.patch.set_facecolor("white")
    ax.errorbar(
        reviewed,
        recall,
        yerr=recall_sd,
        color=HEX_RED,
        marker="s",
        linewidth=2.4,
        markersize=4.8,
        capsize=3,
        label="Dynamic CL",
    )
    baseline_limit = 16.0
    ax.plot(
        [0, baseline_limit],
        [0, baseline_limit],
        color=HEX_GREY,
        linestyle="--",
        linewidth=1.8,
        label="Random expectation",
    )
    ax.annotate(
        f"Loop 8: {reviewed[-1]:.1f}% reviewed\n{recall[-1]:.1f}% corrected",
        xy=(reviewed[-1], recall[-1]),
        xytext=(9.4, 63.5),
        textcoords="data",
        arrowprops={"arrowstyle": "-", "color": HEX_RED, "linewidth": 1},
        color=HEX_RED,
        fontsize=9.2,
        fontweight="bold",
        ha="left",
    )
    ax.text(
        14.15,
        12.8,
        "Random: 14.0%",
        color=HEX_GREY,
        fontsize=8.8,
        ha="right",
    )
    ax.set_title("VinDr: Review Efficiency", loc="left", fontweight="bold")
    ax.text(
        1,
        1.02,
        "Mean across 8 seeds",
        transform=ax.transAxes,
        ha="right",
        color=HEX_MUTED,
        fontsize=8.8,
    )
    ax.text(
        0.01,
        0.93,
        "Each red marker = one loop",
        transform=ax.transAxes,
        color=HEX_MUTED,
        fontsize=8.6,
        va="top",
    )
    ax.set_xlim(0, 16.5)
    ax.set_ylim(0, 70)
    ax.set_xticks(np.arange(0, 17, 2))
    ax.set_xlabel("Entries reviewed (%)")
    ax.set_ylabel("Known errors corrected (%)")
    ax.grid(color=HEX_GRID, linewidth=0.8)
    ax.legend(loc="lower right", frameon=False, ncol=2)
    fig.tight_layout(pad=0.65)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def plot_medpalm_iterative(path: Path, data: pd.DataFrame) -> None:
    dynamic = data[data["method"] == "dynamic_cl"].sort_values("checkpoint")
    x = np.concatenate([[0], dynamic["review_fraction"].to_numpy() * 100])
    y = np.concatenate([[0], dynamic["cumulative_recall"].to_numpy() * 100])
    fig, ax = plt.subplots(figsize=(5.0, 5.15))
    fig.patch.set_facecolor("white")
    ax.plot(
        x,
        y,
        color=HEX_RED,
        marker="s",
        linewidth=2.4,
        markersize=5,
        label="Dynamic CL",
    )
    ax.plot(
        [0, 100],
        [0, 100],
        color=HEX_GREY,
        linestyle="--",
        linewidth=1.8,
        label="Random expectation",
    )
    ax.scatter([x[1]], [y[1]], s=70, facecolor="white", edgecolor=HEX_RED, linewidth=1.8, zorder=4)
    ax.annotate(
        "100/498 reviewed\n68/127 issues found",
        xy=(x[1], y[1]),
        xytext=(34, 59),
        textcoords="data",
        arrowprops={"arrowstyle": "-", "color": HEX_RED, "linewidth": 1},
        color=HEX_RED,
        fontsize=9.4,
        fontweight="bold",
        ha="left",
    )
    ax.set_title("Med-PaLM: Expert-Confirmed Issues", loc="left", fontweight="bold")
    ax.text(
        0.04,
        0.94,
        "Single-seed iterative pilot",
        transform=ax.transAxes,
        ha="left",
        va="top",
        color=HEX_MUTED,
        fontsize=8.8,
    )
    ax.set_xlim(0, 102)
    ax.set_ylim(0, 105)
    ax.set_xlabel("Entries reviewed (%)")
    ax.set_ylabel("Expert-confirmed issues found (%)")
    ax.set_xticks(np.arange(0, 101, 20))
    ax.set_yticks(np.arange(0, 101, 20))
    ax.grid(color=HEX_GRID, linewidth=0.8)
    ax.legend(loc="lower right", frameon=False)
    ax.set_aspect("equal", adjustable="box")
    fig.tight_layout(pad=0.75)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def slide_1(prs: Presentation, note: str) -> None:
    slide = new_slide(prs, 1, TITLES[0])
    # The supplied figure is a complete slide, including its own title.
    # Cover the standard title chrome, then place the figure without cropping.
    add_box(
        slide,
        0,
        0,
        SLIDE_W,
        SLIDE_H,
        fill=WHITE,
        line=WHITE,
        line_width=0,
        name="flowchart-background",
    )
    add_picture_contain(
        slide,
        FLOWCHART_IMAGE,
        0.08,
        0.08,
        13.173,
        7.34,
        name="pipeline-flowchart",
    )
    set_notes(slide, note)


def dataset_panel(
    slide,
    *,
    x: float,
    title: str,
    image: Path,
    lines: list[str],
    accent: RGBColor,
    name: str,
) -> None:
    add_text(slide, x, 1.48, 3.72, 0.32, title, size=15.5, color=accent, bold=True)
    add_box(slide, x, 1.86, 3.72, 2.38, fill=PANEL, line=GRID, name=f"{name}-image-frame")
    add_picture_contain(slide, image, x + 0.05, 1.91, 3.62, 2.28, name=f"{name}-image")
    add_lines(
        slide,
        x,
        4.38,
        3.72,
        1.42,
        [
            {"text": lines[0], "size": 13.2, "bold": True, "color": INK, "space_after": 5},
            {"text": lines[1], "size": 12.5, "color": MUTED, "space_after": 4},
            {"text": lines[2], "size": 12.5, "color": MUTED},
        ],
        margin=0,
        name=f"{name}-text",
    )


def slide_2(prs: Presentation, note: str) -> None:
    slide = new_slide(prs, 2, TITLES[1])
    add_box(slide, 0.66, 1.04, 12.0, 0.34, fill=BLUE_LIGHT, line=BLUE_LIGHT)
    add_text(
        slide,
        0.82,
        1.085,
        11.65,
        0.24,
        "MIMIC-CXR has no exhaustive ground truth for its report-derived training labels.",
        size=13.2,
        color=BLUE_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    dataset_panel(
        slide,
        x=0.55,
        title="MIMIC-CXR: Main Dataset",
        image=MIMIC_IMAGE,
        lines=[
            "Images and radiology reports",
            "CheXpert-derived binary labels",
            "No exhaustive training-entry ground truth",
        ],
        accent=BLUE,
        name="mimic",
    )
    dataset_panel(
        slide,
        x=4.80,
        title="Med-PaLM: Real Expert Reference",
        image=MEDPALM_IMAGE,
        lines=[
            "Expert-reviewed MIMIC-CXR subset",
            "498 compatible binary entries",
            "127 confirmed label issues",
        ],
        accent=TEAL,
        name="medpalm",
    )
    dataset_panel(
        slide,
        x=9.05,
        title="VinDr-CXR: Controlled Ground Truth",
        image=VINDR_IMAGE,
        lines=[
            "3,000 radiologist-labelled images",
            "6 findings = 18,000 binary entries",
            "Known 10%, 20%, and 30% injected errors",
        ],
        accent=RED,
        name="vindr",
    )
    arrow = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(4.28),
        Inches(2.95),
        Inches(4.72),
        Inches(2.95),
    )
    arrow.line.color.rgb = TEAL
    arrow.line.width = Pt(1.7)
    arrow.line.end_arrowhead = True
    add_text(
        slide,
        3.70,
        2.48,
        1.63,
        0.35,
        "Expert-reviewed\ntest subset",
        size=9.2,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_takeaway(
        slide,
        "Two complementary tests: real expert-confirmed MIMIC-CXR errors and complete controlled VinDr ground truth.",
    )
    set_notes(slide, note)


def slide_3(
    prs: Presentation,
    note: str,
    vindr_chart: Path,
    medpalm_chart: Path,
) -> None:
    slide = new_slide(prs, 3, TITLES[2])
    add_picture_contain(slide, vindr_chart, 0.55, 1.08, 7.82, 4.72, name="vindr-one-shot")
    divider = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(8.45),
        Inches(1.2),
        Inches(8.45),
        Inches(5.72),
    )
    divider.line.color.rgb = GRID
    divider.line.width = Pt(0.8)
    add_picture_contain(slide, medpalm_chart, 8.65, 1.08, 4.12, 4.72, name="medpalm-one-shot")
    add_box(slide, 0.82, 5.88, 11.72, 0.46, fill=PANEL, line=GRID, radius=True)
    add_text(
        slide,
        1.00,
        5.985,
        11.36,
        0.24,
        "Random expected errors = total errors x reviewed entries / total entries     |     Example: 3,600 x 3,600 / 18,000 = 720",
        size=11.3,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_takeaway(
        slide,
        "At the same tested review budgets, CL finds 2.5-6.9x more true label errors than uniform random review.",
        color=RED,
    )
    set_notes(slide, note)


def slide_4(
    prs: Presentation,
    note: str,
    quality_chart: Path,
    efficiency_chart: Path,
    medpalm_chart: Path,
) -> None:
    slide = new_slide(prs, 4, TITLES[3])
    add_box(slide, 0.66, 1.02, 12.0, 0.56, fill=BLUE_LIGHT, line=BLUE_LIGHT, radius=True)
    add_text(
        slide,
        0.82,
        1.08,
        11.7,
        0.24,
        "Top 20% of unreviewed CL pool  ->  Oracle correction*  ->  Retrain OOF  ->  Rerank",
        size=12.6,
        color=BLUE_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        0.82,
        1.34,
        11.7,
        0.16,
        "Previously reviewed entries are skipped.  *The real pipeline uses LLM correction, evaluated separately last week.",
        size=8.8,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_picture_contain(slide, quality_chart, 0.55, 1.68, 7.58, 2.13, name="vindr-quality")
    add_picture_contain(slide, efficiency_chart, 0.55, 3.93, 7.58, 2.13, name="vindr-efficiency")
    divider = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(8.25),
        Inches(1.72),
        Inches(8.25),
        Inches(6.05),
    )
    divider.line.color.rgb = GRID
    divider.line.width = Pt(0.8)
    add_picture_contain(slide, medpalm_chart, 8.40, 1.68, 4.38, 4.38, name="medpalm-iterative")
    add_takeaway(
        slide,
        "CL concentrates errors early: 14.0% VinDr review corrects 61.1% of errors; 20.1% Med-PaLM review finds 53.5% of issues.",
        color=RED,
    )
    set_notes(slide, note)


def ensure_inputs() -> None:
    required = [
        SOURCE_MD,
        FLOWCHART_IMAGE,
        MIMIC_IMAGE,
        MEDPALM_IMAGE,
        VINDR_IMAGE,
        VINDR_ONE_SHOT,
        MEDPALM_ONE_SHOT,
        VINDR_ITERATIVE,
        MEDPALM_ITERATIVE,
        VINDR_TAU035_DQS,
    ]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing required inputs:\n" + "\n".join(missing))


def validate_deck(path: Path, notes: dict[int, str]) -> dict[str, Any]:
    prs = Presentation(path)
    if len(prs.slides) != 4:
        raise AssertionError(f"Expected 4 slides, found {len(prs.slides)}")
    titles = []
    out_of_bounds = []
    unexpected_fonts = []
    too_small = []
    note_checks = {}
    for page, slide in enumerate(prs.slides, start=1):
        title = next((s for s in slide.shapes if s.name == "slide-title"), None)
        if title is None:
            raise AssertionError(f"Slide {page} has no title")
        titles.append(title.text.strip())
        actual_note = slide.notes_slide.notes_text_frame.text.strip()
        if actual_note != notes[page]:
            raise AssertionError(f"Slide {page} speaker notes do not match source")
        note_checks[str(page)] = len(actual_note)
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if not run.text.strip():
                            continue
                        if run.font.name != FONT:
                            unexpected_fonts.append(
                                {"slide": page, "shape": shape.name, "font": run.font.name}
                            )
                        if run.font.size is not None and run.font.size.pt < 8:
                            too_small.append(
                                {
                                    "slide": page,
                                    "shape": shape.name,
                                    "size": run.font.size.pt,
                                    "text": run.text[:60],
                                }
                            )
            if (
                shape.left < -2
                or shape.top < -2
                or shape.left + shape.width > prs.slide_width + 2
                or shape.top + shape.height > prs.slide_height + 2
            ):
                out_of_bounds.append({"slide": page, "shape": shape.name})
    if titles != TITLES:
        raise AssertionError({"expected": TITLES, "observed": titles})
    if out_of_bounds:
        raise AssertionError(f"Out-of-bounds shapes: {out_of_bounds}")
    if unexpected_fonts:
        raise AssertionError(f"Unexpected fonts: {unexpected_fonts}")
    if too_small:
        raise AssertionError(f"Text below 8 pt: {too_small}")
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise AssertionError(f"Corrupt PPTX member: {bad_member}")
        note_parts = [
            name
            for name in archive.namelist()
            if name.startswith("ppt/notesSlides/notesSlide") and name.endswith(".xml")
        ]
        if len(note_parts) != 4:
            raise AssertionError(f"Expected 4 notes slides, found {len(note_parts)}")
    return {
        "status": "passed",
        "slides": 4,
        "speaker_note_characters": note_checks,
        "visible_font": FONT,
        "minimum_visible_text_size_pt": 8,
        "shape_bounds": "passed",
        "pptx_integrity": "passed",
    }


def main() -> int:
    ensure_inputs()
    configure_plot_style()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    notes = parse_notes()

    one_shot, med_one_shot = load_one_shot()
    vindr_iter = pd.read_csv(VINDR_ITERATIVE).sort_values("loop")
    vindr_tau035_dqs = pd.read_csv(VINDR_TAU035_DQS).sort_values("loop")
    med_iter = pd.read_csv(MEDPALM_ITERATIVE)

    vindr_one_shot_chart = ASSET_DIR / "slide3_vindr_one_shot.png"
    medpalm_one_shot_chart = ASSET_DIR / "slide3_medpalm_one_shot.png"
    quality_chart = ASSET_DIR / "slide4_vindr_quality.png"
    efficiency_chart = ASSET_DIR / "slide4_vindr_efficiency.png"
    medpalm_iter_chart = ASSET_DIR / "slide4_medpalm_iterative.png"
    plot_vindr_one_shot(vindr_one_shot_chart, one_shot)
    plot_medpalm_one_shot(medpalm_one_shot_chart, med_one_shot)
    plot_vindr_quality(quality_chart, vindr_iter, vindr_tau035_dqs)
    plot_vindr_efficiency(efficiency_chart, vindr_iter)
    plot_medpalm_iterative(medpalm_iter_chart, med_iter)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "Validated confident-learning detection"
    prs.core_properties.subject = "VinDr and Med-PaLM detection validation"
    prs.core_properties.author = "Yihang"
    prs.core_properties.keywords = (
        "confident learning, dataset quality, VinDr-CXR, Med-PaLM, DQS"
    )
    slide_1(prs, notes[1])
    slide_2(prs, notes[2])
    slide_3(prs, notes[3], vindr_one_shot_chart, medpalm_one_shot_chart)
    slide_4(prs, notes[4], quality_chart, efficiency_chart, medpalm_iter_chart)
    prs.save(OUTPUT_PPTX)

    preflight = validate_deck(OUTPUT_PPTX, notes)
    manifest = {
        "status": "passed",
        "source_outline": str(SOURCE_MD),
        "deck": str(OUTPUT_PPTX),
        "font": FONT,
        "slides": 4,
        "speaker_notes": "embedded from source Markdown",
        "data_sources": {
            "vindr_one_shot": str(VINDR_ONE_SHOT),
            "medpalm_one_shot": str(MEDPALM_ONE_SHOT),
            "vindr_iterative": str(VINDR_ITERATIVE),
            "vindr_tau035_dqs": str(VINDR_TAU035_DQS),
            "medpalm_iterative": str(MEDPALM_ITERATIVE),
        },
        "computed_one_shot": one_shot.to_dict(orient="records"),
        "medpalm_one_shot": med_one_shot,
        "preflight": preflight,
    }
    MANIFEST.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"deck": str(OUTPUT_PPTX), "preflight": preflight}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
